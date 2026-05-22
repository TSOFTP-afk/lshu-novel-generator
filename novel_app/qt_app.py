from __future__ import annotations

import json
import os
import re
import shutil
import sys
import traceback
import uuid
from collections.abc import Callable
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from threading import Event
from typing import Any

BASE_DIR = Path(__file__).resolve().parent.parent

from PyQt6.QtCore import QByteArray, QEasingCurve, QPropertyAnimation, QRectF, QSize, QSizeF, QTimer, Qt, QThread, QUrl, pyqtSignal
from PyQt6.QtGui import QAction, QBrush, QColor, QCloseEvent, QFont, QIcon, QImage, QKeySequence, QMovie, QPainter, QPixmap, QShortcut, QTextCharFormat, QTextCursor, QStandardItem, QStandardItemModel
from PyQt6.QtMultimedia import QMediaPlayer
from PyQt6.QtMultimediaWidgets import QGraphicsVideoItem
from PyQt6.QtWidgets import (
    QApplication,
    QCheckBox,
    QColorDialog,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QFrame,
    QGraphicsScene,
    QGraphicsView,
    QGridLayout,
    QHBoxLayout,
    QInputDialog,
    QLabel,
    QLineEdit,
    QLayout,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMenu,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QSlider,
    QSplitter,
    QSplitterHandle,
    QStackedWidget,
    QTabWidget,
    QTextEdit,
    QTreeView,
    QVBoxLayout,
    QWidget,
)

from novel_app.ai_service import DEFAULT_BASE_URL, SimpleAIService
from novel_app.database import ALL_VOLUMES, UNASSIGNED_VOLUMES, Database
from novel_app.exporter import BookExporter
from novel_app.truth import TruthManager
from novel_app.rag import ChapterRetriever
from novel_app.model_router import ModelRouter
from novel_app.qt.dialogs import AiSettingsDialog, PromptSettingsDialog, ask_multiline, ask_text, confirm, error, info, ask_unsaved
from novel_app.qt.drawer import DRAWER_TAB_SEQUENCE, DrawerController, pages as drawer_pages
from novel_app.qt.character_dialog import CharacterDetailDialog
from novel_app.qt.forms import BookOutlineFormWidget, CharacterProfileFormWidget, WorldFormWidget
from novel_app.qt.helpers import (
    DEFAULT_AI_PROBABILITY_META,
    compute_template_stats,
    count_text_characters,
    format_chapter_tree_display_title,
    format_chapter_tree_title,
    get_ai_probability_meta,
    normalize_ai_color_config,
    normalize_ai_probability_pair,
    previous_chapter_ids_in_order,
    row_to_dict,
    set_ai_probability_meta,
)
from novel_app.qt.chat_dialog import ChatDialog
from novel_app.qt.icons import make_tree_icon
from novel_app.qt.image_utils import set_adaptive_image
from novel_app.qt.state import QtAppState
from novel_app.qt.star_graph import ask_relationship_type
from novel_app.qt.theme import PRESET_THEMES, ThemeTokens, build_stylesheet, get_theme
from novel_app.qt.workers import AiWorker, FunctionWorker
from novel_app.secure_storage import DPAPI_PREFIX, protect_secret, unprotect_secret
from novel_app.prompt_settings import (
    AGENT_PROMPT_ENTRIES,
    PROMPT_SETTINGS_VERSION,
    normalize_prompt_overrides,
)
from novel_app.prompt_utils import prompt_value, render_prompt
from novel_app.structured_templates import (
    BOOK_OUTLINE_FORM_SCHEMA,
    BOOK_OUTLINE_VOLUME_PLAN_FIELDS,
    WORLD_FORM_SCHEMA,
    book_outline_template,
    book_outline_form_defaults,
    book_outline_form_to_markdown,
    book_outline_markdown_to_form,
    character_profile_form_defaults,
    character_profile_form_to_markdown,
    character_profile_markdown_to_form,
    character_profile_schema_for_role,
    coerce_book_outline_markdown,
    coerce_world_markdown,
    merge_book_outline_form,
    merge_book_outline_supplement,
    merge_character_analysis_payload,
    merge_character_profile_form,
    merge_character_profile_supplement,
    merge_world_analysis_payload,
    merge_world_supplement,
    normalize_book_outline_form,
    normalize_character_profile_form,
    normalize_world_form,
    world_form_defaults,
    world_form_to_markdown,
    world_markdown_to_form,
    world_entry_template,
)
from novel_app.text_importer import ParsedBook, parse_long_text

try:
    from docx import Document as _DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


KIND_ROLE = Qt.ItemDataRole.UserRole + 1
ID_ROLE = Qt.ItemDataRole.UserRole + 2
GROUP_BOOK_ROLE = Qt.ItemDataRole.UserRole + 3
AI_LEVEL_ROLE = Qt.ItemDataRole.UserRole + 4

DEEPSEEK_CHAT_SYSTEM_PROMPT = """#### 定位
- 智能助手名称：灵枢小说创作工作台 AI 对话助手
- 主要任务：结合当前书籍数据，为作者提供剧情讨论、设定问答、人物推演、正文修改建议和创作决策支持。

#### 能力
- 上下文整合：优先读取 Truth File、RAG 检索结果、书籍级设定、世界观、人物卡、当前章节大纲和正文。
- 写作协作：可以做剧情推演、结构建议、文本润色、人物关系分析和问题排查。
- 冲突识别：当资料互相矛盾时，指出冲突来源，并以 Truth File 与用户当前明确指令为优先。

#### 使用说明
- 回答要直接、自然、便于作者继续操作。
- 需要改写时，给出可复制的片段或清晰修改方案；不要擅自覆盖用户正文。
- 不确定的信息要标注为推断，不要把猜测写成既定事实。
- 用户要求角色代入时，遵守所选人物卡、关系状态和剧情进度。"""

DEEPSEEK_CHARACTER_CHAT_PROMPT_TEMPLATE = """#### 定位
- 当前模式：角色对话
- 你正在代入书中人物：{name}
- 人物定位：{role}

#### 人物资料
{profile}
{self_line}

#### 使用说明
- 使用该角色的第一人称口吻回答作者。
- 只知道人物在当前剧情和设定中合理知道的内容。
- 结合上方 Truth File、RAG、章节上下文和人物关系说话。
- 不要跳出角色解释自己是 AI；除非作者明确要求写作建议。
- 资料不足时，可以用角色视角表达迟疑，不要硬编书中没有的事实。"""

BUILTIN_CHAT_SYSTEM_PROMPT = DEEPSEEK_CHAT_SYSTEM_PROMPT
CHARACTER_CHAT_PROMPT_TEMPLATE = DEEPSEEK_CHARACTER_CHAT_PROMPT_TEMPLATE

PROMPT_SETTING_ENTRIES = [
    {
        "key": "chat.system",
        "category": "AI 对话",
        "label": "AI 对话内置系统提示词",
        "default": BUILTIN_CHAT_SYSTEM_PROMPT,
        "placeholders": "",
    },
    {
        "key": "chat.character_template",
        "category": "AI 对话",
        "label": "角色对话提示词模板",
        "default": CHARACTER_CHAT_PROMPT_TEMPLATE,
        "placeholders": "name, role, profile, self_name, self_line",
    },
    *AGENT_PROMPT_ENTRIES,
]

THEME_PRESET_LABELS = [
    ("white", "白色亮色"),
    ("light_blue", "浅蓝亮色"),
    ("night", "暗色夜间"),
    ("sepia", "米黄护眼"),
    ("custom", "自定义"),
]

THEME_COLOR_FIELDS = [
    ("bg", "底层背景"),
    ("surface", "主面板"),
    ("surface_alt", "次级面板/按钮"),
    ("editor_bg", "正文编辑区"),
    ("text", "正文文字"),
    ("muted", "弱提示文字"),
    ("primary", "强调色"),
    ("primary_soft", "强调浅色"),
    ("border", "边框"),
    ("focus", "聚焦光圈"),
    ("success", "成功"),
    ("warning", "警告"),
    ("danger", "危险"),
]

BACKGROUND_FILTERS = [
    ("none", "无滤镜"),
    ("soft_dark", "柔和暗化（推荐）"),
    ("soft_light", "柔和亮化"),
    ("warm", "护眼暖化"),
    ("blur_dark", "模糊暗化"),
]

AI_COLOR_FIELDS = [("fg", "文字"), ("bg", "徽章"), ("tree", "目录")]

AI_PURPOSES = [
    ("writing", "写作生成"),
    ("outline", "同步大纲"),
    ("detector", "AI概率检测"),
    ("correction", "AI纠错"),
    ("import", "智能导入分类"),
    ("skills", "Skills提炼"),
    ("book_analysis", "全书分析"),
    ("review", "长篇生成/校验"),
    ("chat", "AI 对话"),
]

APP_DISPLAY_NAME = "灵枢v1.3"
DEFAULT_AI_REQUEST_RETRIES = 2
DEFAULT_BOOK_ANALYSIS_CONCURRENCY = 3
DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON = '{"thinking":{"type":"disabled"}}'
DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON = '{"thinking":{"type":"enabled"},"reasoning_effort":"high"}'
BOOK_ANALYSIS_SOURCE_REQUEST_LIMIT = 12
BOOK_ANALYSIS_SOURCE_SNIPPETS_PER_REQUEST = 3
BOOK_ANALYSIS_SOURCE_SNIPPET_CHARS = 1200
BOOK_ANALYSIS_SOURCE_CONTEXT_CHARS = 18000
BOOK_ANALYSIS_MIN_FIELD_CONFIDENCE = 0.65
AI_SETTINGS_VERSION = "ai_profiles_v5_deepseek_prompts"

AI_MODE_PURPOSE = {
    "draft": "writing",
    "continue": "writing",
    "polish": "writing",
    "summary": "outline",
}


class ViewSettingsDialog(QDialog):
    def __init__(self, parent: QWidget, settings: dict[str, Any]) -> None:
        super().__init__(parent)
        self.setObjectName("ViewSettingsDialog")
        self.setWindowTitle("视图配置")
        self.setModal(True)
        self.resize(760, 560)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(12)

        preset_row = QHBoxLayout()
        preset_row.addWidget(QLabel("系统预设"))
        self.preset_combo = QComboBox()
        for code, label in THEME_PRESET_LABELS:
            self.preset_combo.addItem(label, code)
        preset = str(settings.get("theme_preset", "light_blue"))
        index = self.preset_combo.findData(preset)
        self.preset_combo.setCurrentIndex(index if index >= 0 else 1)
        preset_row.addWidget(self.preset_combo, 1)
        layout.addLayout(preset_row)

        self.color_toggle = QPushButton("展开自定义颜色")
        self.color_toggle.setCheckable(True)
        self.color_toggle.setObjectName("DrawerToggle")
        self.color_toggle.toggled.connect(self._toggle_color_drawer)
        layout.addWidget(self.color_toggle)

        self.color_drawer = QFrame()
        self.color_drawer.setObjectName("Card")
        drawer_layout = QVBoxLayout(self.color_drawer)
        drawer_layout.setContentsMargins(12, 12, 12, 12)
        drawer_layout.setSpacing(8)

        grid = QGridLayout()
        grid.setHorizontalSpacing(10)
        grid.setVerticalSpacing(8)
        self.color_fields: dict[str, QLineEdit] = {}
        custom_colors = settings.get("custom_theme_colors")
        colors = custom_colors if isinstance(custom_colors, dict) and custom_colors else self._preset_colors(preset)
        for row, (key, label) in enumerate(THEME_COLOR_FIELDS):
            field = QLineEdit(str(colors.get(key, "")))
            button = QPushButton("选择")
            button.clicked.connect(lambda _checked=False, color_key=key: self._choose_color(color_key))
            self.color_fields[key] = field
            grid.addWidget(QLabel(label), row, 0)
            grid.addWidget(field, row, 1)
            grid.addWidget(button, row, 2)
        drawer_layout.addLayout(grid)
        self.color_drawer.setVisible(False)
        layout.addWidget(self.color_drawer)

        self.ai_color_values = normalize_ai_color_config(settings.get("ai_colors", {}))
        self.ai_color_buttons: dict[tuple[str, str], QPushButton] = {}
        self.ai_color_toggle = QPushButton("展开 AI 标识颜色")
        self.ai_color_toggle.setCheckable(True)
        self.ai_color_toggle.setObjectName("DrawerToggle")
        self.ai_color_toggle.toggled.connect(self._toggle_ai_color_drawer)
        layout.addWidget(self.ai_color_toggle)

        self.ai_color_drawer = QFrame()
        self.ai_color_drawer.setObjectName("Card")
        ai_drawer_layout = QVBoxLayout(self.ai_color_drawer)
        ai_drawer_layout.setContentsMargins(12, 12, 12, 12)
        ai_drawer_layout.setSpacing(8)
        ai_grid = QGridLayout()
        ai_grid.setHorizontalSpacing(8)
        ai_grid.setVerticalSpacing(8)
        ai_grid.addWidget(QLabel("级别"), 0, 0)
        for col, (_field, label) in enumerate(AI_COLOR_FIELDS, start=1):
            ai_grid.addWidget(QLabel(label), 0, col)
        for row, (level, meta) in enumerate(DEFAULT_AI_PROBABILITY_META.items(), start=1):
            ai_grid.addWidget(QLabel(str(meta["label"])), row, 0)
            for col, (field, _label) in enumerate(AI_COLOR_FIELDS, start=1):
                button = QPushButton()
                button.clicked.connect(lambda _checked=False, lv=level, fd=field: self._choose_ai_color(lv, fd))
                self.ai_color_buttons[(level, field)] = button
                ai_grid.addWidget(button, row, col)
                self._sync_ai_color_button(level, field)
        reset_ai_colors = QPushButton("重置默认")
        reset_ai_colors.clicked.connect(self._reset_ai_colors)
        ai_drawer_layout.addLayout(ai_grid)
        ai_drawer_layout.addWidget(reset_ai_colors)
        self.ai_color_drawer.setVisible(False)
        layout.addWidget(self.ai_color_drawer)

        image_row = QHBoxLayout()
        image_row.addWidget(QLabel("底层图片"))
        self.background_field = QLineEdit(str(settings.get("background_image", "")))
        choose_image = QPushButton("选择图片")
        clear_image = QPushButton("清除")
        choose_image.clicked.connect(self._choose_image)
        clear_image.clicked.connect(lambda: self.background_field.clear())
        image_row.addWidget(self.background_field, 1)
        image_row.addWidget(choose_image)
        image_row.addWidget(clear_image)
        layout.addLayout(image_row)

        video_row = QHBoxLayout()
        video_row.addWidget(QLabel("底层视频"))
        self.background_video_field = QLineEdit(str(settings.get("background_video", "")))
        choose_video = QPushButton("选择视频/GIF")
        clear_video = QPushButton("清除")
        choose_video.clicked.connect(self._choose_video)
        clear_video.clicked.connect(lambda: self.background_video_field.clear())
        video_row.addWidget(self.background_video_field, 1)
        video_row.addWidget(choose_video)
        video_row.addWidget(clear_video)
        layout.addLayout(video_row)

        filter_row = QHBoxLayout()
        filter_row.addWidget(QLabel("图片滤镜"))
        self.filter_combo = QComboBox()
        for code, label in BACKGROUND_FILTERS:
            self.filter_combo.addItem(label, code)
        filter_code = str(settings.get("background_filter", "soft_dark"))
        filter_index = self.filter_combo.findData(filter_code)
        self.filter_combo.setCurrentIndex(filter_index if filter_index >= 0 else 1)
        filter_row.addWidget(self.filter_combo, 1)
        filter_row.addWidget(QLabel("强度"))
        self.filter_strength = QSlider(Qt.Orientation.Horizontal)
        self.filter_strength.setRange(0, 100)
        self.filter_strength.setValue(int(settings.get("background_filter_strength", 35) or 35))
        filter_row.addWidget(self.filter_strength, 1)
        layout.addLayout(filter_row)

        font_row = QHBoxLayout()
        font_row.addWidget(QLabel("正文字号"))
        self.font_size_slider = QSlider(Qt.Orientation.Horizontal)
        self.font_size_slider.setRange(10, 24)
        self.font_size_slider.setValue(int(settings.get("editor_font_size", 14)))
        self.font_size_label = QLabel(str(self.font_size_slider.value()))
        self.font_size_slider.valueChanged.connect(lambda v: self.font_size_label.setText(str(v)))
        font_row.addWidget(self.font_size_slider)
        font_row.addWidget(self.font_size_label)
        layout.addLayout(font_row)

        line_row = QHBoxLayout()
        line_row.addWidget(QLabel("行高"))
        self.line_height_slider = QSlider(Qt.Orientation.Horizontal)
        self.line_height_slider.setRange(100, 300)
        self.line_height_slider.setValue(int(settings.get("editor_line_height", 145)))
        self.line_height_label = QLabel(f"{self.line_height_slider.value() / 100:.2f}")
        self.line_height_slider.valueChanged.connect(lambda v: self.line_height_label.setText(f"{v / 100:.2f}"))
        line_row.addWidget(self.line_height_slider)
        line_row.addWidget(self.line_height_label)
        layout.addLayout(line_row)

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        save_button = buttons.button(QDialogButtonBox.StandardButton.Save)
        if save_button:
            save_button.setText("应用并保存")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

        self.preset_combo.currentIndexChanged.connect(self._load_selected_preset)
        if preset == "custom":
            self.color_toggle.setChecked(True)

    def _preset_colors(self, preset: str) -> dict[str, str]:
        tokens = PRESET_THEMES.get(preset, PRESET_THEMES["light_blue"])
        return {key: getattr(tokens, key) for key, _label in THEME_COLOR_FIELDS}

    def _load_selected_preset(self) -> None:
        preset = str(self.preset_combo.currentData() or "light_blue")
        if preset == "custom":
            self.color_toggle.setChecked(True)
            return
        for key, value in self._preset_colors(preset).items():
            self.color_fields[key].setText(value)

    def _toggle_color_drawer(self, checked: bool) -> None:
        self.color_drawer.setVisible(checked)
        self.color_toggle.setText("收起自定义颜色" if checked else "展开自定义颜色")
        self._resize_for_drawers()

    def _toggle_ai_color_drawer(self, checked: bool) -> None:
        self.ai_color_drawer.setVisible(checked)
        self.ai_color_toggle.setText("收起 AI 标识颜色" if checked else "展开 AI 标识颜色")
        self._resize_for_drawers()

    def _resize_for_drawers(self) -> None:
        extra = 0
        if self.color_drawer.isVisible():
            extra += 160
        if self.ai_color_drawer.isVisible():
            extra += 180
        self.resize(760, 560 + extra)

    def _choose_color(self, key: str) -> None:
        current = QColor(self.color_fields[key].text().strip())
        color = QColorDialog.getColor(current if current.isValid() else QColor("#FFFFFF"), self, "选择颜色")
        if color.isValid():
            self.color_fields[key].setText(color.name().upper())
            custom_index = self.preset_combo.findData("custom")
            if custom_index >= 0:
                self.preset_combo.setCurrentIndex(custom_index)

    def _sync_ai_color_button(self, level: str, field: str) -> None:
        button = self.ai_color_buttons.get((level, field))
        if not button:
            return
        color = self.ai_color_values.get(level, {}).get(field, "#FFFFFF")
        button.setText(color)
        text_color = "#FFFFFF" if field in {"bg", "tree"} and QColor(color).lightness() < 120 else "#1F2A36"
        button.setStyleSheet(f"background: {color}; color: {text_color}; border-radius: 8px; padding: 5px 8px;")

    def _choose_ai_color(self, level: str, field: str) -> None:
        current = QColor(self.ai_color_values.get(level, {}).get(field, "#FFFFFF"))
        color = QColorDialog.getColor(current if current.isValid() else QColor("#FFFFFF"), self, "选择 AI 标识颜色")
        if not color.isValid():
            return
        self.ai_color_values.setdefault(level, {})[field] = color.name().upper()
        self._sync_ai_color_button(level, field)

    def _reset_ai_colors(self) -> None:
        self.ai_color_values = normalize_ai_color_config({})
        for level in DEFAULT_AI_PROBABILITY_META:
            for field, _label in AI_COLOR_FIELDS:
                self._sync_ai_color_button(level, field)

    def _choose_image(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择底层图片",
            "",
            "Image Files (*.png *.jpg *.jpeg *.bmp *.webp);;All Files (*.*)",
        )
        if path:
            self.background_field.setText(path)

    def _choose_video(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择底层视频/GIF",
            "",
            "Video & GIF Files (*.mp4 *.avi *.mkv *.webm *.mov *.gif *.webp *.apng);;All Files (*.*)",
        )
        if path:
            self.background_video_field.setText(path)

    def values(self) -> dict[str, Any]:
        return {
            "theme_preset": str(self.preset_combo.currentData() or "light_blue"),
            "custom_theme_colors": {key: field.text().strip() for key, field in self.color_fields.items()},
            "ai_colors": normalize_ai_color_config(self.ai_color_values),
            "background_image": self.background_field.text().strip(),
            "background_video": self.background_video_field.text().strip(),
            "background_filter": str(self.filter_combo.currentData() or "soft_dark"),
            "background_filter_strength": int(self.filter_strength.value()),
            "editor_font_size": int(self.font_size_slider.value()),
            "editor_line_height": int(self.line_height_slider.value()),
        }


class SoftSplitterHandle(QSplitterHandle):
    def __init__(self, orientation: Qt.Orientation, parent: QSplitter) -> None:
        super().__init__(orientation, parent)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground, True)

    def sizeHint(self) -> QSize:  # noqa: N802
        if self.orientation() == Qt.Orientation.Horizontal:
            return QSize(10, 10)
        return QSize(10, 10)

    def paintEvent(self, _event) -> None:  # noqa: N802
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        color = QColor(96, 120, 145, 58)
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(color)
        if self.orientation() == Qt.Orientation.Horizontal:
            handle_height = min(118, max(44, int(self.height() * 0.14)))
            rect_x = int((self.width() - 4) / 2)
            rect_y = int((self.height() - handle_height) / 2)
            painter.drawRoundedRect(rect_x, rect_y, 4, handle_height, 2, 2)
        else:
            handle_width = min(118, max(44, int(self.width() * 0.14)))
            rect_x = int((self.width() - handle_width) / 2)
            rect_y = int((self.height() - 4) / 2)
            painter.drawRoundedRect(rect_x, rect_y, handle_width, 4, 2, 2)
        painter.end()


class SoftSplitter(QSplitter):
    def createHandle(self) -> QSplitterHandle:  # noqa: N802
        return SoftSplitterHandle(self.orientation(), self)


def _find_plain_text_matches(
    text: str,
    query: str,
    *,
    case_sensitive: bool = False,
    whole_word: bool = False,
) -> list[tuple[int, int]]:
    if not query:
        return []
    flags = 0 if case_sensitive else re.IGNORECASE
    if whole_word:
        pattern = re.compile(rf"(?<![\w\u4e00-\u9fff]){re.escape(query)}(?![\w\u4e00-\u9fff])", flags)
        return [(match.start(), match.end()) for match in pattern.finditer(text)]

    haystack = text if case_sensitive else text.casefold()
    needle = query if case_sensitive else query.casefold()
    matches: list[tuple[int, int]] = []
    index = 0
    while True:
        index = haystack.find(needle, index)
        if index < 0:
            break
        end = index + len(query)
        matches.append((index, end))
        index += max(1, len(query))
    return matches


def _search_snippet(text: str, start: int, end: int, radius: int = 28) -> str:
    left = max(0, start - radius)
    right = min(len(text), end + radius)
    prefix = "…" if left > 0 else ""
    suffix = "…" if right < len(text) else ""
    snippet = text[left:right].replace("\n", " ").strip()
    return f"{prefix}{snippet}{suffix}"


class SearchDialog(QDialog):
    def __init__(self, parent: Any) -> None:
        super().__init__(parent)
        self._main_window = parent
        self.setWindowTitle("查找")
        self.setModal(False)
        self.resize(560, 440)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(14, 14, 14, 14)
        layout.setSpacing(10)

        query_row = QHBoxLayout()
        query_row.addWidget(QLabel("查找内容"))
        self.query_field = QLineEdit()
        self.query_field.setClearButtonEnabled(True)
        self.query_field.returnPressed.connect(self._find_next)
        query_row.addWidget(self.query_field, 1)
        layout.addLayout(query_row)

        scope_row = QHBoxLayout()
        scope_row.addWidget(QLabel("搜索范围"))
        self.scope_combo = QComboBox()
        self.scope_combo.currentIndexChanged.connect(lambda _index: self._run_search())
        scope_row.addWidget(self.scope_combo, 1)
        layout.addLayout(scope_row)

        options_row = QHBoxLayout()
        self.case_box = QCheckBox("区分大小写")
        self.whole_word_box = QCheckBox("全词匹配")
        self.case_box.toggled.connect(lambda _checked: self._run_search())
        self.whole_word_box.toggled.connect(lambda _checked: self._run_search())
        options_row.addWidget(self.case_box)
        options_row.addWidget(self.whole_word_box)
        options_row.addStretch(1)
        layout.addLayout(options_row)

        action_row = QHBoxLayout()
        self.prev_button = QPushButton("查找上一个")
        self.next_button = QPushButton("查找下一个")
        close_button = QPushButton("关闭")
        self.prev_button.clicked.connect(self._find_previous)
        self.next_button.clicked.connect(self._find_next)
        close_button.clicked.connect(self.close)
        action_row.addStretch(1)
        action_row.addWidget(self.prev_button)
        action_row.addWidget(self.next_button)
        action_row.addWidget(close_button)
        layout.addLayout(action_row)

        self.result_list = QListWidget()
        self.result_list.itemActivated.connect(self._activate_result)
        self.result_list.currentItemChanged.connect(lambda current, _previous: self._activate_current_result(current))
        layout.addWidget(self.result_list, 1)

        self.status_label = QLabel("输入关键词后按 Enter 或“查找下一个”。")
        self.status_label.setObjectName("MetaLabel")
        layout.addWidget(self.status_label)

    def set_scope_options(self, options: list[tuple[str, str]], current_scope: str) -> None:
        current = str(self.scope_combo.currentData() or current_scope)
        self.scope_combo.blockSignals(True)
        self.scope_combo.clear()
        for scope, label in options:
            self.scope_combo.addItem(label, scope)
        target = current if any(scope == current for scope, _label in options) else current_scope
        index = self.scope_combo.findData(target)
        self.scope_combo.setCurrentIndex(index if index >= 0 else 0)
        self.scope_combo.blockSignals(False)

    def focus_query(self) -> None:
        self.query_field.setFocus()
        self.query_field.selectAll()

    def set_results(self, results: list[dict[str, Any]], current_index: int) -> None:
        self.result_list.blockSignals(True)
        self.result_list.clear()
        for index, result in enumerate(results):
            label = f"{index + 1}. {result.get('title', '')} · {result.get('field_label', '')}：{result.get('snippet', '')}"
            item = QListWidgetItem(label)
            item.setData(Qt.ItemDataRole.UserRole, index)
            self.result_list.addItem(item)
        if results and 0 <= current_index < len(results):
            self.result_list.setCurrentRow(current_index)
        self.result_list.blockSignals(False)
        self._sync_status(len(results), current_index)

    def _scope(self) -> str:
        return str(self.scope_combo.currentData() or "")

    def _run_search(self) -> None:
        query = self.query_field.text()
        if not query.strip():
            self._main_window._clear_search_results()
            return
        self._main_window._run_search_from_dialog(
            query=query,
            scope=self._scope(),
            case_sensitive=self.case_box.isChecked(),
            whole_word=self.whole_word_box.isChecked(),
        )

    def _find_next(self) -> None:
        if not self._main_window._search_matches or self._params_changed():
            self._run_search()
            return
        self._main_window._search_next_match()

    def _find_previous(self) -> None:
        if not self._main_window._search_matches or self._params_changed():
            self._run_search()
            return
        self._main_window._search_previous_match()

    def _params_changed(self) -> bool:
        return (
            self.query_field.text().strip() != self._main_window._search_query
            or self._scope() != self._main_window._search_scope
            or self.case_box.isChecked() != self._main_window._search_case_sensitive
            or self.whole_word_box.isChecked() != self._main_window._search_whole_word
        )

    def _activate_result(self, item: QListWidgetItem) -> None:
        self._activate_current_result(item)

    def _activate_current_result(self, item: QListWidgetItem | None) -> None:
        if item is None:
            return
        index = int(item.data(Qt.ItemDataRole.UserRole) or 0)
        self._main_window._activate_search_result(index)

    def _sync_status(self, count: int, current_index: int) -> None:
        if count <= 0:
            self.status_label.setText("未找到匹配项。")
            return
        self.status_label.setText(f"找到 {count} 处匹配，当前第 {current_index + 1} 处。")


class NovelQtMainWindow(QMainWindow):
    def __init__(self, database: Database, ai_service: SimpleAIService) -> None:
        super().__init__()
        self.setObjectName("MainWindow")
        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        self.database = database
        self.ai_service = ai_service
        self.exporter = BookExporter(database)
        self.truth_manager = TruthManager(self.database.data_dir, self.database)
        self.rag_retriever = ChapterRetriever(
            str(self.database.data_dir / "rag_vectors"),
        )
        self.model_router = ModelRouter(default_model=self.ai_service.model or "deepseek-chat", preset="budget")
        self.ai_service.set_model_router(self.model_router)
        self.state = QtAppState()

        self.ui_config_path = self.database.data_dir / "ui_settings.json"
        self.ai_config_path = self.database.data_dir / "ai_settings.json"
        self.ui_settings: dict[str, Any] = {}
        self.theme_mode = "light"
        self.theme_preset = "light_blue"
        self.custom_theme_colors: dict[str, str] = {}
        self.background_image = ""
        self.background_video = ""
        self.background_filter = "soft_dark"
        self.background_filter_strength = 35
        self.background_processed_image = ""
        self.ai_colors = normalize_ai_color_config({})
        self.library_collapsed = False
        self.library_width = 260
        self.library_auto_collapse_width = 132
        self.qt_drawer_width = 420
        self.auto_save_enabled = False
        self.auto_save_interval_seconds = 30
        self.focus_mode = False
        self._applying_library_layout = False
        self._tree_icon_cache: dict[tuple[str, str, str, str, str], Any] = {}

        self.loading_editor = False
        self.suppress_tree_selection = False
        self.current_tree_key: tuple[str, int | None, int | None] | None = None
        self.global_skills: list[dict[str, Any]] = []
        self.book_skills: list[dict[str, Any]] = []
        self.bound_skills: list[dict[str, Any]] = []
        self.characters: list[dict[str, Any]] = []
        self.world_entries: list[dict[str, Any]] = []
        self.selected_character_id: int | None = None
        self._world_form_book_id: int | None = None
        self.selected_review_run_id: int | None = None
        self.chat_dialog: ChatDialog | None = None
        self.character_detail_dialog: CharacterDetailDialog | None = None
        self.search_dialog: SearchDialog | None = None
        self._search_matches: list[dict[str, Any]] = []
        self._search_match_index = -1
        self._search_query = ""
        self._search_scope = ""
        self._search_case_sensitive = False
        self._search_whole_word = False
        self.drawer_tab_indexes: dict[str, int] = {}
        self.drawer_controller = DrawerController(self)
        self.current_book_cover_path = ""
        self.character_image_path = ""
        self.ai_thread: QThread | None = None
        self.ai_worker: AiWorker | None = None
        self.ai_jobs: dict[str, dict[str, Any]] = {}
        self.detached_cancelled_jobs: dict[str, dict[str, Any]] = {}
        self.close_after_task_stop = False
        self.current_ai_spans: list[dict[str, Any]] = []
        self.suppress_ai_span_clear = False
        self._ai_task_queue: list[dict[str, Any]] = []
        self.cancel_event: Event | None = None
        self.task_target = "log"
        self.task_stream_mode = "replace"
        self.task_prefix_pending = ""
        self.input_events: list[tuple[datetime, str, int]] = []
        self.editor_text_lengths: dict[int, int] = {}
        self.ai_writing = False
        self.ai_total_chars: int = 0
        self.ai_total_tokens: int = 0
        self.task_target_chapter_id: int | None = None
        self.task_result_text = ""
        self.task_base_content = ""
        self.task_base_outline = ""
        self.task_buffer_prefix = ""
        self.task_log_streaming = False
        self.last_task_log_line = ""
        self.task_summary_expanded = False
        self.task_summary_animation: QPropertyAnimation | None = None
        self.loaded_chapter_content = ""
        self.loaded_chapter_summary_text = ""

        self.word_timer = QTimer(self)
        self.word_timer.setSingleShot(True)
        self.word_timer.timeout.connect(self._update_word_count)

        self._status_fade_timer = QTimer(self)
        self._status_fade_timer.setSingleShot(True)
        self._status_fade_timer.timeout.connect(lambda: self.status_label.setText(""))

        self.cancel_grace_timer = QTimer(self)
        self.cancel_grace_timer.setSingleShot(True)
        self.cancel_grace_timer.timeout.connect(self._finalize_stale_cancelled_jobs)

        self.close_poll_timer = QTimer(self)
        self.close_poll_timer.setInterval(250)
        self.close_poll_timer.timeout.connect(self._complete_close_when_tasks_stop)

        self.stats_timer = QTimer(self)
        self.stats_timer.setInterval(2000)
        self.stats_timer.timeout.connect(self._refresh_input_stats_table)
        self.stats_timer.start()

        self._load_ui_settings()
        self._load_ai_settings()
        self.auto_save_timer = QTimer(self)
        self.auto_save_timer.timeout.connect(self._auto_save_current)
        self._build_ui()
        self._apply_theme()
        self._apply_background_video()
        self._apply_editor_font_settings(self.ui_settings)
        self._restore_window_state()
        self._load_library_tree()
        self._setup_shortcuts()
        self._sync_auto_save_timer()

    def _load_ui_settings(self) -> None:
        payload: dict[str, Any] = {}
        if self.ui_config_path.exists():
            try:
                payload = json.loads(self.ui_config_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                payload = {}
        self.ui_settings = payload
        legacy_mode = str(payload.get("theme_mode", "light")).lower()
        self.theme_preset = str(payload.get("theme_preset", "night" if legacy_mode == "dark" else "light_blue"))
        if self.theme_preset not in PRESET_THEMES and self.theme_preset != "custom":
            self.theme_preset = "light_blue"
        custom_colors = payload.get("custom_theme_colors", {})
        self.custom_theme_colors = dict(custom_colors) if isinstance(custom_colors, dict) else {}
        self.background_image = str(payload.get("background_image", ""))
        self.background_video = str(payload.get("background_video", ""))
        self.background_filter = str(payload.get("background_filter", "soft_dark"))
        if self.background_filter not in {code for code, _label in BACKGROUND_FILTERS}:
            self.background_filter = "soft_dark"
        self.background_filter_strength = self._safe_int(payload.get("background_filter_strength", 35), 35, 0, 100)
        self.ai_colors = normalize_ai_color_config(payload.get("ai_colors", {}))
        set_ai_probability_meta(self.ai_colors)
        self.theme_mode = "dark" if self.theme_preset == "night" else "light"
        self.state.theme_mode = self.theme_mode
        self.library_collapsed = bool(payload.get("library_collapsed", False))
        self.library_width = self._safe_int(payload.get("library_width", 260), 260, 180, 520)
        self.qt_drawer_width = self._safe_int(payload.get("qt_drawer_width", 420), 420, 320, 620)
        self.auto_save_enabled = bool(payload.get("auto_save_enabled", False))
        self.auto_save_interval_seconds = self._safe_int(payload.get("auto_save_interval_seconds", 30), 30, 5, 3600)

    def _save_ui_settings(self) -> None:
        self.ui_settings.update(
            {
                "theme_mode": self.theme_mode,
                "theme_preset": self.theme_preset,
                "custom_theme_colors": dict(self.custom_theme_colors),
                "background_image": self.background_image,
                "background_video": self.background_video,
                "background_filter": self.background_filter,
                "background_filter_strength": int(self.background_filter_strength),
                "ai_colors": normalize_ai_color_config(self.ai_colors),
                "library_collapsed": bool(self.library_collapsed),
                "library_width": int(self.library_width),
                "qt_drawer_width": int(self.qt_drawer_width),
                "auto_save_enabled": bool(self.auto_save_enabled),
                "auto_save_interval_seconds": int(self.auto_save_interval_seconds),
            }
        )
        try:
            if hasattr(self, "main_splitter"):
                self.ui_settings["qt_splitter_state"] = bytes(self.main_splitter.saveState().toBase64()).decode("ascii")
            self.ui_settings["qt_window_geometry"] = bytes(self.saveGeometry().toBase64()).decode("ascii")
            self._write_json_file(self.ui_config_path, self.ui_settings)
        except OSError:
            return

    def _write_json_file(self, path: Path, payload: dict[str, Any]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = path.with_name(f"{path.name}.tmp")
        temp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        temp_path.replace(path)

    def _safe_int(self, value: Any, default: int, minimum: int, maximum: int) -> int:
        try:
            result = int(value)
        except (TypeError, ValueError):
            result = default
        return max(minimum, min(maximum, result))

    @staticmethod
    def _compact_json_object_text(payload: dict[str, Any]) -> str:
        return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))

    def _normalize_json_object_text(self, value: Any, default_text: str) -> str:
        try:
            default_payload = json.loads(default_text)
        except json.JSONDecodeError:
            default_payload = {}
        if isinstance(value, dict):
            return self._compact_json_object_text(value)
        text = str(value or "").strip() or default_text
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            return self._compact_json_object_text(default_payload if isinstance(default_payload, dict) else {})
        if not isinstance(parsed, dict):
            return self._compact_json_object_text(default_payload if isinstance(default_payload, dict) else {})
        return self._compact_json_object_text(parsed)

    def _parse_json_object_text(self, value: Any, label: str) -> dict[str, Any]:
        text = str(value or "").strip()
        if not text:
            return {}
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ValueError(f"{label}必须是 JSON 对象：{exc}") from exc
        if not isinstance(parsed, dict):
            raise ValueError(f"{label}必须是 JSON 对象。")
        return parsed

    def _load_ai_settings(self) -> None:
        config: dict[str, Any] = {}
        if self.ai_config_path.exists():
            try:
                config = json.loads(self.ai_config_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                config = {}
        profiles_raw = config.get("profiles")
        profiles: dict[str, dict[str, str | bool]] = {}
        if isinstance(profiles_raw, dict):
            for purpose, _label in AI_PURPOSES:
                raw_profile = profiles_raw.get(purpose, {})
                profiles[purpose] = self._load_ai_profile(raw_profile if isinstance(raw_profile, dict) else {})
            if config.get("version") != AI_SETTINGS_VERSION:
                writing_profile = profiles.get("writing", {})
                if self._ai_profile_has_credentials(writing_profile):
                    for purpose, _label in AI_PURPOSES:
                        if purpose == "writing":
                            continue
                        profile = profiles.get(purpose, {})
                        if not self._ai_profile_has_credentials(profile):
                            profiles[purpose] = dict(writing_profile)
        else:
            legacy_profile = self._load_ai_profile(
                {
                    "remote_enabled": config.get("remote_enabled", self.ai_service.is_remote_configured()),
                    "api_key": config.get("api_key", self.ai_service.api_key or ""),
                    "base_url": config.get("base_url", self.ai_service.base_url or DEFAULT_BASE_URL),
                    "model": config.get("model", self.ai_service.model or ""),
                }
            )
            for purpose, _label in AI_PURPOSES:
                profiles[purpose] = (
                    dict(legacy_profile)
                    if self._ai_profile_has_credentials(legacy_profile)
                    else {
                        "remote_enabled": False,
                        "api_key": "",
                        "base_url": DEFAULT_BASE_URL,
                        "model": "",
                    }
                )
        request_retries = self._safe_int(
            config.get("request_retries", DEFAULT_AI_REQUEST_RETRIES),
            DEFAULT_AI_REQUEST_RETRIES,
            0,
            5,
        )
        book_analysis_concurrency = self._safe_int(
            config.get("book_analysis_concurrency", DEFAULT_BOOK_ANALYSIS_CONCURRENCY),
            DEFAULT_BOOK_ANALYSIS_CONCURRENCY,
            1,
            8,
        )
        book_analysis_fast_extra_body = self._normalize_json_object_text(
            config.get("book_analysis_fast_extra_body", DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON),
            DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON,
        )
        book_analysis_final_extra_body = self._normalize_json_object_text(
            config.get("book_analysis_final_extra_body", DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON),
            DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON,
        )
        stored_prompt_version = str(config.get("prompt_settings_version", ""))
        prompt_source = (
            config.get("prompt_overrides", {})
            if stored_prompt_version == PROMPT_SETTINGS_VERSION
            else {}
        )
        prompt_overrides = normalize_prompt_overrides(prompt_source, PROMPT_SETTING_ENTRIES)
        self.ai_request_retries = request_retries
        self.book_analysis_concurrency = book_analysis_concurrency
        self.book_analysis_fast_extra_body = book_analysis_fast_extra_body
        self.book_analysis_final_extra_body = book_analysis_final_extra_body
        self.prompt_overrides = prompt_overrides
        self.ai_settings = {
            "profiles": profiles,
            "request_retries": request_retries,
            "book_analysis_concurrency": book_analysis_concurrency,
            "book_analysis_fast_extra_body": book_analysis_fast_extra_body,
            "book_analysis_final_extra_body": book_analysis_final_extra_body,
            "prompt_settings_version": PROMPT_SETTINGS_VERSION,
            "prompt_overrides": prompt_overrides,
        }
        writing_service = self._service_from_profile(profiles.get("writing", {}))
        self.ai_service.configure(
            api_key=writing_service.api_key,
            base_url=writing_service.base_url,
            model=writing_service.model,
        )
        self.ai_service.set_prompt_overrides(prompt_overrides)
        if config and (not isinstance(profiles_raw, dict) or self._ai_config_needs_migration(config)):
            try:
                self._write_json_file(
                    self.ai_config_path,
                    self._build_stored_ai_settings(
                        profiles,
                        request_retries,
                        book_analysis_concurrency,
                        book_analysis_fast_extra_body,
                        book_analysis_final_extra_body,
                        prompt_overrides,
                    ),
                )
            except (OSError, ValueError):
                return

    def _load_ai_profile(self, profile: dict[str, Any]) -> dict[str, str | bool]:
        enabled = bool(profile.get("remote_enabled", False))
        raw_api_key = str(profile.get("api_key", ""))
        try:
            api_key = unprotect_secret(raw_api_key)
        except Exception:
            api_key = ""
        base_url = str(profile.get("base_url", DEFAULT_BASE_URL) or DEFAULT_BASE_URL)
        model = str(profile.get("model", ""))
        try:
            service = SimpleAIService()
            if enabled and api_key and model:
                service.configure(api_key=api_key, base_url=base_url, model=model)
                base_url = service.base_url or DEFAULT_BASE_URL
            else:
                service.configure(api_key=None, base_url=base_url, model=None)
                base_url = service.base_url or DEFAULT_BASE_URL
                enabled = False
        except ValueError:
            enabled = False
            api_key = ""
            base_url = DEFAULT_BASE_URL
            model = ""
        return {
            "remote_enabled": enabled,
            "api_key": api_key,
            "base_url": base_url,
            "model": model,
        }

    def _ai_config_needs_migration(self, config: dict[str, Any]) -> bool:
        if config.get("version") != AI_SETTINGS_VERSION:
            return True
        if config.get("prompt_settings_version") != PROMPT_SETTINGS_VERSION:
            return True
        profiles = config.get("profiles")
        if not isinstance(profiles, dict):
            return True
        if "request_retries" not in config:
            return True
        for key in ("book_analysis_concurrency", "book_analysis_fast_extra_body", "book_analysis_final_extra_body", "prompt_overrides"):
            if key not in config:
                return True
        prompt_overrides = config.get("prompt_overrides")
        if not isinstance(prompt_overrides, dict):
            return True
        for item in PROMPT_SETTING_ENTRIES:
            if str(item.get("key", "")) not in prompt_overrides:
                return True
        for purpose, _label in AI_PURPOSES:
            if purpose not in profiles:
                return True
            profile = profiles.get(purpose, {})
            if not isinstance(profile, dict):
                return True
            if str(profile.get("api_key", "")).strip() and not str(profile.get("api_key", "")).startswith(DPAPI_PREFIX):
                return True
        return False

    def _build_stored_ai_settings(
        self,
        profiles: dict[str, dict[str, str | bool]],
        request_retries: int | None = None,
        book_analysis_concurrency: int | None = None,
        book_analysis_fast_extra_body: str | dict[str, Any] | None = None,
        book_analysis_final_extra_body: str | dict[str, Any] | None = None,
        prompt_overrides: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        stored_profiles: dict[str, dict[str, str | bool]] = {}
        for purpose, _label in AI_PURPOSES:
            profile = profiles.get(purpose, {})
            plain_api_key = str(profile.get("api_key", ""))
            stored_profiles[purpose] = {
                "remote_enabled": bool(profile.get("remote_enabled", False)),
                "api_key": protect_secret(plain_api_key) if plain_api_key else "",
                "base_url": str(profile.get("base_url", DEFAULT_BASE_URL) or DEFAULT_BASE_URL),
                "model": str(profile.get("model", "")),
            }
        retries = self._safe_int(
            request_retries if request_retries is not None else getattr(self, "ai_request_retries", DEFAULT_AI_REQUEST_RETRIES),
            DEFAULT_AI_REQUEST_RETRIES,
            0,
            5,
        )
        concurrency = self._safe_int(
            book_analysis_concurrency
            if book_analysis_concurrency is not None
            else getattr(self, "book_analysis_concurrency", DEFAULT_BOOK_ANALYSIS_CONCURRENCY),
            DEFAULT_BOOK_ANALYSIS_CONCURRENCY,
            1,
            8,
        )
        fast_extra = self._normalize_json_object_text(
            book_analysis_fast_extra_body
            if book_analysis_fast_extra_body is not None
            else getattr(self, "book_analysis_fast_extra_body", DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON),
            DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON,
        )
        final_extra = self._normalize_json_object_text(
            book_analysis_final_extra_body
            if book_analysis_final_extra_body is not None
            else getattr(self, "book_analysis_final_extra_body", DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON),
            DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON,
        )
        prompts = normalize_prompt_overrides(
            prompt_overrides if prompt_overrides is not None else getattr(self, "prompt_overrides", {}),
            PROMPT_SETTING_ENTRIES,
        )
        return {
            "version": AI_SETTINGS_VERSION,
            "prompt_settings_version": PROMPT_SETTINGS_VERSION,
            "profiles": stored_profiles,
            "request_retries": retries,
            "book_analysis_concurrency": concurrency,
            "book_analysis_fast_extra_body": fast_extra,
            "book_analysis_final_extra_body": final_extra,
            "prompt_overrides": prompts,
        }

    def _ai_profile_has_credentials(self, profile: dict[str, Any] | dict[str, str | bool]) -> bool:
        return bool(
            isinstance(profile, dict)
            and profile.get("remote_enabled", False)
            and str(profile.get("api_key", "")).strip()
            and str(profile.get("model", "")).strip()
        )

    def _service_from_profile(self, profile: dict[str, str | bool]) -> SimpleAIService:
        service = SimpleAIService()
        if bool(profile.get("remote_enabled", False)) and str(profile.get("api_key", "")) and str(profile.get("model", "")):
            service.configure(
                api_key=str(profile.get("api_key", "")),
                base_url=str(profile.get("base_url", DEFAULT_BASE_URL) or DEFAULT_BASE_URL),
                model=str(profile.get("model", "")),
            )
        else:
            service.configure(
                api_key=None,
                base_url=str(profile.get("base_url", DEFAULT_BASE_URL) or DEFAULT_BASE_URL),
                model=None,
            )
        service.set_prompt_overrides(getattr(self, "prompt_overrides", {}))
        return service

    def _ai_profiles(self) -> dict[str, dict[str, str | bool]]:
        profiles = self.ai_settings.get("profiles", {}) if isinstance(self.ai_settings, dict) else {}
        return profiles if isinstance(profiles, dict) else {}

    def _ai_service_for_purpose(self, purpose: str) -> SimpleAIService | None:
        profile = self._ai_profiles().get(purpose, {})
        if not isinstance(profile, dict):
            return None
        try:
            service = self._service_from_profile(profile)
        except ValueError:
            return None
        return service if service.is_remote_configured() else None

    def _require_ai_service(self, purpose: str, label: str) -> SimpleAIService | None:
        service = self._ai_service_for_purpose(purpose)
        if service is None:
            info(self, "AI 未配置", f"{label}需要在 AI 设置中配置“{dict(AI_PURPOSES).get(purpose, purpose)}”用途。")
            return None
        return service

    def _save_ai_settings(self, values: dict[str, Any]) -> None:
        profiles_raw = values.get("profiles", {})
        if not isinstance(profiles_raw, dict):
            profiles_raw = {}
        request_retries = self._safe_int(
            values.get("request_retries", DEFAULT_AI_REQUEST_RETRIES),
            DEFAULT_AI_REQUEST_RETRIES,
            0,
            5,
        )
        book_analysis_concurrency = self._safe_int(
            values.get("book_analysis_concurrency", DEFAULT_BOOK_ANALYSIS_CONCURRENCY),
            DEFAULT_BOOK_ANALYSIS_CONCURRENCY,
            1,
            8,
        )
        profiles: dict[str, dict[str, str | bool]] = {}
        try:
            book_analysis_fast_extra = self._parse_json_object_text(
                values.get("book_analysis_fast_extra_body", DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON),
                "一级/二级额外请求参数 JSON",
            )
            book_analysis_final_extra = self._parse_json_object_text(
                values.get("book_analysis_final_extra_body", DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON),
                "三级额外请求参数 JSON",
            )
            book_analysis_fast_extra_text = self._compact_json_object_text(book_analysis_fast_extra)
            book_analysis_final_extra_text = self._compact_json_object_text(book_analysis_final_extra)
            for purpose, _label in AI_PURPOSES:
                profile = profiles_raw.get(purpose, {})
                profile = profile if isinstance(profile, dict) else {}
                loaded = self._load_ai_profile(
                    {
                        "remote_enabled": bool(profile.get("remote_enabled", False)),
                        "api_key": str(profile.get("api_key", "")),
                        "base_url": str(profile.get("base_url", DEFAULT_BASE_URL) or DEFAULT_BASE_URL),
                        "model": str(profile.get("model", "")),
                    }
                )
                if bool(profile.get("remote_enabled", False)) and (not loaded["api_key"] or not loaded["model"]):
                    raise ValueError(f"{dict(AI_PURPOSES).get(purpose, purpose)}已启用，但 API Key 或模型为空。")
                profiles[purpose] = loaded
            self._write_json_file(
                self.ai_config_path,
                self._build_stored_ai_settings(
                    profiles,
                    request_retries,
                    book_analysis_concurrency,
                    book_analysis_fast_extra_text,
                    book_analysis_final_extra_text,
                    getattr(self, "prompt_overrides", {}),
                ),
            )
        except (OSError, ValueError) as exc:
            error(self, "AI 设置无效", str(exc))
            return
        self.ai_request_retries = request_retries
        self.book_analysis_concurrency = book_analysis_concurrency
        self.book_analysis_fast_extra_body = book_analysis_fast_extra_text
        self.book_analysis_final_extra_body = book_analysis_final_extra_text
        self.ai_settings = {
            "profiles": profiles,
            "request_retries": request_retries,
            "book_analysis_concurrency": book_analysis_concurrency,
            "book_analysis_fast_extra_body": book_analysis_fast_extra_text,
            "book_analysis_final_extra_body": book_analysis_final_extra_text,
            "prompt_settings_version": PROMPT_SETTINGS_VERSION,
            "prompt_overrides": getattr(self, "prompt_overrides", normalize_prompt_overrides({}, PROMPT_SETTING_ENTRIES)),
        }
        writing_service = self._service_from_profile(profiles.get("writing", {}))
        self.ai_service.configure(
            api_key=writing_service.api_key,
            base_url=writing_service.base_url,
            model=writing_service.model,
        )
        self.ai_service.set_prompt_overrides(getattr(self, "prompt_overrides", {}))
        self._set_status("")

    def _build_ui(self) -> None:
        self.setWindowTitle(APP_DISPLAY_NAME)
        self.resize(1540, 940)
        self.setMinimumSize(1024, 720)

        self.central_shell = QWidget()
        self.central_shell.setObjectName("CentralShell")
        shell_layout = QGridLayout(self.central_shell)
        shell_layout.setContentsMargins(0, 0, 0, 0)
        shell_layout.setSpacing(0)

        self.background_video_label = QLabel(self.central_shell)
        self.background_video_label.setScaledContents(True)
        self.background_video_label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents, True)
        self.background_video_label.hide()
        self.background_video_scene = QGraphicsScene(self.central_shell)
        self.background_video_item = QGraphicsVideoItem()
        self.background_video_item.setAspectRatioMode(Qt.AspectRatioMode.KeepAspectRatioByExpanding)
        self.background_video_scene.addItem(self.background_video_item)
        self.background_video_view = QGraphicsView(self.background_video_scene, self.central_shell)
        self.background_video_view.setFrameShape(QFrame.Shape.NoFrame)
        self.background_video_view.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.background_video_view.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.background_video_view.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents, True)
        self.background_video_view.setStyleSheet("background: transparent; border: none;")
        self.background_video_view.hide()
        self.background_movie: QMovie | None = None
        self.background_media_player: QMediaPlayer | None = None

        root = QWidget(self.central_shell)
        root.setObjectName("AppRoot")
        root.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        self.root_widget = root
        root_layout = QVBoxLayout(root)
        root_layout.setContentsMargins(16, 14, 16, 12)
        root_layout.setSpacing(10)

        shell_layout.addWidget(self.background_video_view, 0, 0)
        shell_layout.addWidget(self.background_video_label, 0, 0)
        shell_layout.addWidget(root, 0, 0)
        self.setCentralWidget(self.central_shell)
        self.background_video_view.lower()
        self.background_video_label.lower()
        self.root_widget.raise_()

        self.header = QFrame()
        self.header.setObjectName("HeaderBar")
        self.header.setMinimumHeight(58)
        self.header.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        header_layout = QHBoxLayout(self.header)
        header_layout.setContentsMargins(10, 5, 10, 5)
        header_layout.setSpacing(12)
        title_box = QVBoxLayout()
        self.title_label = QLabel("小说写作台")
        self.title_label.setObjectName("TitleLabel")
        self.title_label.setVisible(False)
        title_box.addWidget(self.title_label)
        header_layout.addLayout(title_box, 0)
        self.toolbar_scroll = QScrollArea()
        self.toolbar_scroll.setObjectName("ToolbarScroll")
        self.toolbar_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.toolbar_scroll.setWidgetResizable(False)
        self.toolbar_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.toolbar_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.toolbar_scroll.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.toolbar_scroll.setFixedHeight(50)
        self.toolbar_scroll.setStyleSheet("QScrollArea { background: transparent; border: none; }")
        self.toolbar_scroll_content = QWidget()
        self.toolbar_scroll_content.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        self.toolbar_box = QHBoxLayout(self.toolbar_scroll_content)
        self.toolbar_box.setContentsMargins(0, 1, 0, 1)
        self.toolbar_box.setSpacing(8)
        self.toolbar_buttons: list[QPushButton] = []
        self.toolbar_scroll.setWidget(self.toolbar_scroll_content)
        header_layout.addWidget(self.toolbar_scroll)
        header_layout.addStretch(1)
        root_layout.addWidget(self.header)
        self._build_toolbar()
        self._sync_toolbar_scroll_size()

        body_wrapper = QWidget()
        body_wrapper.setObjectName("BodyWrapper")
        body_layout = QHBoxLayout(body_wrapper)
        body_layout.setContentsMargins(0, 0, 0, 0)
        body_layout.setSpacing(0)

        self.main_splitter = SoftSplitter(Qt.Orientation.Horizontal)
        self.main_splitter.setObjectName("MainSplitter")
        self.main_splitter.setChildrenCollapsible(False)
        self.main_splitter.setHandleWidth(10)

        self.library_panel = self._build_library_panel()
        self.workspace_panel = self._build_workspace_panel()
        self.main_splitter.addWidget(self.library_panel)
        self.main_splitter.addWidget(self.workspace_panel)
        self.main_splitter.setStretchFactor(0, 0)
        self.main_splitter.setStretchFactor(1, 1)
        self.main_splitter.setSizes([96 if self.library_collapsed else self.library_width, 1100])
        self.main_splitter.splitterMoved.connect(self._on_splitter_moved)

        body_layout.addWidget(self.main_splitter, 1)

        self.drawer = QFrame()
        self.drawer.setObjectName("SideDrawer")
        self.drawer.setFixedWidth(0)
        self.drawer.setVisible(False)
        drawer_layout = QVBoxLayout(self.drawer)
        drawer_layout.setContentsMargins(0, 0, 0, 0)
        drawer_layout.setSpacing(0)
        drawer_layout.addWidget(self._build_drawer_title_bar())
        self.drawer_tabs = self._build_drawer_tabs()
        self.drawer_tabs.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground, True)
        drawer_layout.addWidget(self.drawer_tabs, 1)

        body_layout.addWidget(self.drawer, 0)

        root_layout.addWidget(body_wrapper, 1)

        self.status_bar_frame = QFrame()
        self.status_bar_frame.setObjectName("StatusBar")
        self.status_bar_frame.setMinimumHeight(34)
        self.status_bar_frame.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        status_layout = QHBoxLayout(self.status_bar_frame)
        status_layout.setContentsMargins(10, 5, 10, 5)
        status_layout.setSpacing(10)
        self.status_label = QLabel("")
        self.status_label.setObjectName("StatusText")
        self.status_label.setMinimumWidth(0)
        self.status_label.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
        self.status_label.setAlignment(Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft)
        self.save_label = QLabel("已保存")
        self.save_label.setObjectName("StatusMetric")
        self.save_label.setMinimumWidth(66)
        self.selection_label = QLabel("选中 0 字")
        self.selection_label.setObjectName("StatusMetric")
        self.selection_label.setMinimumWidth(88)
        self.ai_ratio_label = QLabel("AI标记 0%")
        self.ai_ratio_label.setObjectName("StatusMetric")
        self.ai_ratio_label.setMinimumWidth(94)
        self.word_label = QLabel("字数 0")
        self.word_label.setObjectName("StatusMetric")
        self.word_label.setMinimumWidth(78)
        for metric_label in (self.save_label, self.selection_label, self.ai_ratio_label, self.word_label):
            metric_label.setMinimumHeight(24)
            metric_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            metric_label.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        status_layout.addWidget(self.status_label, 1)
        status_layout.addWidget(self.save_label)
        status_layout.addWidget(self.selection_label)
        status_layout.addWidget(self.ai_ratio_label)
        status_layout.addWidget(self.word_label)
        root_layout.addWidget(self.status_bar_frame)

        self._apply_responsive_chrome()

    def _build_toolbar(self) -> None:
        edit_menu = QMenu(self)
        for label, callback in [
            ("历史", self._open_history_dialog),
            ("导入到当前", self._import_for_current_context),
            ("智能导入", self._smart_import),
        ]:
            action = QAction(label, self)
            action.triggered.connect(callback)
            edit_menu.addAction(action)
        edit_menu.addSeparator()
        export_sub = edit_menu.addMenu("导出")
        for label, callback in [
            ("导出 TXT", lambda: self._export_current("txt")),
            ("导出 Markdown", lambda: self._export_current("md")),
            ("导出 DOCX", lambda: self._export_current("docx")),
        ]:
            action = QAction(label, self)
            action.triggered.connect(callback)
            export_sub.addAction(action)
        edit_button = QPushButton("编辑 ▾")
        edit_button.setMenu(edit_menu)
        self.toolbar_box.addWidget(edit_button)
        self.toolbar_buttons.append(edit_button)

        ai_menu = QMenu(self)
        ai_menu.addAction(QAction("AI 设置", self, triggered=self._open_ai_settings))
        ai_menu.addSeparator()
        for label, callback in [
            ("AI 对话", self._open_ai_chat),
            ("长篇生成", self._start_multi_agent_generation),
            ("润色", lambda: self._start_generation("polish")),
            ("同步大纲", self._refresh_summary),
            ("检测AI概率", self._detect_current_chapter_ai_probability),
            ("分析全书", self._analyze_book_full),
            ("清除本章AI标记", self._clear_current_ai_spans),
        ]:
            action = QAction(label, self)
            action.triggered.connect(callback)
            ai_menu.addAction(action)
        ai_button = QPushButton("AI ▾")
        ai_button.setMenu(ai_menu)
        self.toolbar_box.addWidget(ai_button)
        self.toolbar_buttons.append(ai_button)

        view_menu = QMenu(self)
        view_menu.addAction(QAction("视图配置", self, triggered=self._open_view_settings))
        view_menu.addSeparator()
        for label, callback in [
            ("折叠导航", self._toggle_library_collapsed),
            ("详情抽屉", lambda: self._open_drawer("details")),
            ("生成记录", lambda: self._open_drawer("review")),
            ("任务抽屉", lambda: self._open_drawer("tasks")),
            ("专注写作", self._toggle_focus_mode),
        ]:
            action = QAction(label, self)
            action.triggered.connect(callback)
            view_menu.addAction(action)
        view_button = QPushButton("视图 ▾")
        view_button.setMenu(view_menu)
        self.toolbar_box.addWidget(view_button)
        self.toolbar_buttons.append(view_button)

        settings_menu = QMenu(self)
        self.auto_save_action = QAction("启用定时保存", self)
        self.auto_save_action.setCheckable(True)
        self.auto_save_action.setChecked(self.auto_save_enabled)
        self.auto_save_action.triggered.connect(self._toggle_auto_save)
        self.auto_save_interval_action = QAction(self._auto_save_interval_action_text(), self)
        self.auto_save_interval_action.triggered.connect(self._configure_auto_save_interval)
        settings_menu.addAction(self.auto_save_action)
        settings_menu.addAction(self.auto_save_interval_action)
        settings_menu.addSeparator()
        settings_menu.addAction(QAction("Agent 与提示词配置", self, triggered=self._open_prompt_settings))
        settings_button = QPushButton("设定 ▾")
        settings_button.setMenu(settings_menu)
        self.toolbar_box.addWidget(settings_button)
        self.toolbar_buttons.append(settings_button)
        for button in self.toolbar_buttons:
            self._prepare_toolbar_button(button)

    def _prepare_toolbar_button(self, button: QPushButton) -> None:
        button.setObjectName("ToolbarButton")
        button.setMinimumHeight(34)
        button.setMinimumWidth(max(86, button.sizeHint().width()))
        button.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)

    def _sync_toolbar_scroll_size(self) -> None:
        if not hasattr(self, "toolbar_scroll_content"):
            return
        self.toolbar_scroll_content.adjustSize()
        hint = self.toolbar_scroll_content.sizeHint()
        self.toolbar_scroll_content.setMinimumSize(hint.width(), max(44, hint.height()))

    def _build_library_panel(self) -> QFrame:
        panel = QFrame()
        panel.setObjectName("Panel")
        panel.setMinimumWidth(72)
        panel.setMaximumWidth(560)
        panel.setSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Expanding)
        layout = QVBoxLayout(panel)
        self.library_layout = layout
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(10)

        top = QHBoxLayout()
        self.library_title = QLabel("作品轨道")
        self.library_title.setObjectName("SectionTitle")
        self.library_create_button = QPushButton("新建书籍")
        self.library_create_button.clicked.connect(self._create_contextual_item)
        self.collapse_button = QPushButton("收起")
        self.collapse_button.clicked.connect(self._toggle_library_collapsed)
        top.addWidget(self.library_title, 1)
        top.addWidget(self.library_create_button)
        top.addWidget(self.collapse_button)
        layout.addLayout(top)

        action_row = QHBoxLayout()
        action_row.setSpacing(6)
        self.library_generate_button = QPushButton("长篇生成")
        self.library_generate_button.setObjectName("TinyButton")
        self.library_generate_button.clicked.connect(self._start_multi_agent_generation)
        self.library_skills_button = QPushButton("Skills 分层管理")
        self.library_skills_button.setObjectName("TinyButton")
        self.library_skills_button.clicked.connect(self._open_book_settings)
        action_row.addWidget(self.library_generate_button, 1)
        action_row.addWidget(self.library_skills_button)
        layout.addLayout(action_row)

        self.tree_model = QStandardItemModel(self)
        self.tree_model.setHorizontalHeaderLabels(["书 / 卷 / 章"])
        self.library_tree = QTreeView()
        self.library_tree.setModel(self.tree_model)
        self.library_tree.setHeaderHidden(True)
        self.library_tree.setEditTriggers(QTreeView.EditTrigger.NoEditTriggers)
        self.library_tree.setIndentation(14)
        self.library_tree.setUniformRowHeights(True)
        self.library_tree.setDragDropMode(QTreeView.DragDropMode.InternalMove)
        self.library_tree.setDragEnabled(True)
        self.library_tree.setAcceptDrops(True)
        self.library_tree.setDropIndicatorShown(True)
        self.library_tree.setDefaultDropAction(Qt.DropAction.MoveAction)
        self.tree_model.supportedDropActions = lambda: Qt.DropAction.MoveAction
        self._patch_tree_model_drop()
        self.library_tree.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.library_tree.customContextMenuRequested.connect(self._show_library_context_menu)
        self.library_tree.selectionModel().currentChanged.connect(self._on_tree_current_changed)
        self.library_tree.doubleClicked.connect(self._on_library_double_clicked)
        layout.addWidget(self.library_tree, 1)
        self.library_compact_list = QListWidget()
        self.library_compact_list.setObjectName("LibraryCompactList")
        self.library_compact_list.setVisible(False)
        self.library_compact_list.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.library_compact_list.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.library_compact_list.setIconSize(QSize(24, 24))
        self.library_compact_list.setSpacing(6)
        self.library_compact_list.setUniformItemSizes(True)
        self.library_compact_list.currentItemChanged.connect(self._on_compact_book_selected)
        layout.addWidget(self.library_compact_list, 1)
        self._apply_library_collapsed_state()
        return panel

    def _build_workspace_panel(self) -> QFrame:
        panel = QFrame()
        panel.setObjectName("Panel")
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        hero = QHBoxLayout()
        title_box = QVBoxLayout()
        title_line = QHBoxLayout()
        self.editor_title = QLabel("")
        self.editor_title.setObjectName("TitleLabel")
        self.editor_title.setMinimumWidth(0)
        self.editor_title.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
        self.ai_probability_badge = QLabel("")
        self.ai_probability_badge.setObjectName("AiProbabilityBadge")
        self.ai_probability_badge.setVisible(False)
        self.ai_probability_button = QPushButton("AI检测")
        self.ai_probability_button.setObjectName("TinyButton")
        self.ai_probability_button.clicked.connect(self._detect_current_chapter_ai_probability)
        self.ai_probability_button.setVisible(False)
        self.editor_path = QLabel("")
        self.editor_path.setObjectName("MetaLabel")
        self.editor_path.setMinimumWidth(0)
        self.editor_path.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
        title_line.addWidget(self.editor_title)
        title_line.addWidget(self.ai_probability_badge)
        title_line.addWidget(self.ai_probability_button)
        title_line.addStretch(1)
        title_box.addLayout(title_line)
        title_box.addWidget(self.editor_path)
        hero.addLayout(title_box, 1)
        self.input_stats_card = QFrame()
        self.input_stats_card.setObjectName("Card")
        self.input_stats_card.setMinimumWidth(320)
        self.input_stats_card.setSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Fixed)
        self.input_stats_card.setMaximumHeight(70)
        stats_layout = QHBoxLayout(self.input_stats_card)
        stats_layout.setContentsMargins(10, 8, 10, 8)
        stats_layout.setSpacing(8)
        self.input_stat_value_labels: dict[str, QLabel] = {}
        for key, label in (("user", "用户"), ("ai", "AI"), ("total", "合计")):
            cell = QFrame()
            cell.setObjectName("StatCell")
            cell_layout = QVBoxLayout(cell)
            cell_layout.setContentsMargins(8, 4, 8, 4)
            cell_layout.setSpacing(1)
            name_label = QLabel(label)
            name_label.setObjectName("MetaLabel")
            value_label = QLabel("0 字/小时")
            value_label.setObjectName("SectionTitle")
            value_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            name_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            cell_layout.addWidget(name_label)
            cell_layout.addWidget(value_label)
            stats_layout.addWidget(cell, 1)
            self.input_stat_value_labels[key] = value_label
        hero.addWidget(self.input_stats_card)
        self.inline_details_button = QPushButton("详情")
        self.inline_details_button.clicked.connect(lambda: self._open_drawer("details"))
        self.inline_tasks_button = QPushButton("任务")
        self.inline_tasks_button.clicked.connect(lambda: self._open_drawer("tasks"))
        hero.addWidget(self.inline_details_button)
        hero.addWidget(self.inline_tasks_button)
        layout.addLayout(hero)
        self._refresh_input_stats_table()

        self.ai_quick_bar = QFrame()
        self.ai_quick_bar.setObjectName("Card")
        self.ai_quick_bar.setMinimumHeight(68)
        self.ai_quick_bar.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        quick_layout = QHBoxLayout(self.ai_quick_bar)
        quick_layout.setContentsMargins(10, 7, 10, 7)
        quick_layout.setSpacing(0)
        self.ai_quick_scroll = QScrollArea()
        self.ai_quick_scroll.setObjectName("QuickActionScroll")
        self.ai_quick_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.ai_quick_scroll.setWidgetResizable(False)
        self.ai_quick_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.ai_quick_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.ai_quick_scroll.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.ai_quick_scroll.setFixedHeight(54)
        self.ai_quick_scroll.setStyleSheet("QScrollArea { background: transparent; border: none; }")
        quick_content = QWidget()
        self.ai_quick_content = quick_content
        quick_content.setObjectName("QuickActionContent")
        quick_content.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        quick_buttons_layout = QHBoxLayout(quick_content)
        quick_buttons_layout.setContentsMargins(0, 2, 0, 2)
        quick_buttons_layout.setSpacing(6)
        self.ai_quick_buttons: list[QPushButton] = []
        for label, callback in [
            ("✨ 润色", lambda: self._start_generation("polish")),
            ("🔍 AI检测", self._detect_current_chapter_ai_probability),
            ("长篇生成", self._start_multi_agent_generation),
            ("💾 保存", self._save_current),
            ("📋 全章复制", self._copy_full_chapter),
            ("📐 一键排版", self._auto_format_content),
            ("🔎 搜索", self._show_search_bar),
            ("🔍 纠错", self._auto_correct_chapter),
            ("👁 隐藏正文", self._toggle_content_visibility),
        ]:
            btn = QPushButton(label)
            btn.setObjectName("TinyButton")
            btn.setMinimumHeight(32)
            btn.setMinimumWidth(max(72, btn.sizeHint().width()))
            btn.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
            btn.clicked.connect(callback)
            quick_buttons_layout.addWidget(btn)
            self.ai_quick_buttons.append(btn)
        quick_buttons_layout.addStretch(1)
        quick_content.adjustSize()
        quick_content.setMinimumSize(quick_content.sizeHint().width(), 44)
        self.ai_quick_scroll.setWidget(quick_content)
        quick_layout.addWidget(self.ai_quick_scroll, 1)
        layout.addWidget(self.ai_quick_bar)

        self.stack = QStackedWidget()
        self.empty_page = self._build_empty_page()
        self.outline_page = self._build_outline_page()
        self.chapter_page = self._build_chapter_page()
        self.stack.addWidget(self.empty_page)
        self.stack.addWidget(self.outline_page)
        self.stack.addWidget(self.chapter_page)
        layout.addWidget(self.stack, 1)
        return panel

    def _build_empty_page(self) -> QWidget:
        page = QFrame()
        page.setObjectName("GhostCard")
        layout = QVBoxLayout(page)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.addStretch(1)
        return page

    def _build_outline_page(self) -> QWidget:
        page = QFrame()
        page.setObjectName("Card")
        layout = QVBoxLayout(page)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(10)

        content_row = QHBoxLayout()
        content_row.setSpacing(14)
        self.book_cover_panel = QFrame()
        self.book_cover_panel.setObjectName("Card")
        cover_layout = QVBoxLayout(self.book_cover_panel)
        cover_layout.setContentsMargins(10, 10, 10, 10)
        cover_layout.setSpacing(8)
        self.book_cover_preview = QLabel("封面")
        self.book_cover_preview.setObjectName("ImagePreview")
        self.book_cover_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.book_cover_preview.setMinimumSize(120, 168)
        self.book_cover_preview.setMaximumSize(170, 240)
        cover_buttons = QHBoxLayout()
        choose_cover = QPushButton("选择")
        clear_cover = QPushButton("清除")
        choose_cover.clicked.connect(self._choose_book_cover)
        clear_cover.clicked.connect(self._clear_book_cover)
        cover_buttons.addWidget(choose_cover)
        cover_buttons.addWidget(clear_cover)
        cover_layout.addWidget(self.book_cover_preview)
        cover_layout.addLayout(cover_buttons)

        outline_box = QVBoxLayout()
        outline_header = QHBoxLayout()
        label = QLabel("大纲")
        label.setObjectName("SectionTitle")
        self.outline_template_btn = QPushButton("重置书籍模板")
        self.outline_template_btn.clicked.connect(self._fill_book_outline_template)
        self.outline_markdown_toggle_btn = QPushButton("Markdown")
        self.outline_markdown_toggle_btn.setObjectName("TinyButton")
        self.outline_markdown_toggle_btn.clicked.connect(lambda: self._toggle_markdown_view("outline"))
        self.outline_checkbox_btn = QPushButton("勾选")
        self.outline_checkbox_btn.setObjectName("TinyButton")
        self.outline_checkbox_btn.clicked.connect(lambda: self._toggle_markdown_checkbox("outline"))
        outline_header.addWidget(label, 1)
        outline_header.addWidget(self.outline_markdown_toggle_btn)
        outline_header.addWidget(self.outline_checkbox_btn)
        outline_header.addWidget(self.outline_template_btn)
        self.outline_editor = QPlainTextEdit()
        self.outline_editor.textChanged.connect(self._on_editor_changed)
        self.outline_editor.textChanged.connect(self._refresh_outline_markdown_preview)
        self.outline_editor.selectionChanged.connect(self._update_selection_count)
        self.book_outline_form = BookOutlineFormWidget()
        self.book_outline_form.changed.connect(self._on_book_outline_form_changed)
        self.outline_markdown_preview = QTextEdit()
        self.outline_markdown_preview.setReadOnly(True)
        self.outline_markdown_preview.setPlaceholderText("Markdown 预览")
        self.outline_markdown_stack = QStackedWidget()
        self.outline_markdown_stack.addWidget(self.book_outline_form)
        self.outline_markdown_stack.addWidget(self.outline_editor)
        self.outline_markdown_stack.addWidget(self.outline_markdown_preview)
        outline_box.addLayout(outline_header)
        outline_box.addWidget(self.outline_markdown_stack, 1)
        content_row.addWidget(self.book_cover_panel)
        content_row.addLayout(outline_box, 1)
        layout.addLayout(content_row, 1)
        return page

    def _build_chapter_page(self) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(10)

        editor_card = QFrame()
        editor_card.setObjectName("Card")
        editor_layout = QVBoxLayout(editor_card)
        editor_layout.setContentsMargins(12, 10, 12, 12)
        editor_label = QLabel("正文")
        editor_label.setObjectName("SectionTitle")
        self.content_label = editor_label
        self.content_editor = QPlainTextEdit()
        self.content_editor.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.content_editor.textChanged.connect(self._on_editor_changed)
        self.content_editor.selectionChanged.connect(self._update_selection_count)
        editor_layout.addWidget(editor_label)
        editor_layout.addWidget(self.content_editor, 1)
        layout.addWidget(editor_card, 1)
        return page

    def _build_drawer_tabs(self) -> QTabWidget:
        return self.drawer_controller.build_tabs()

    def _build_drawer_title_bar(self) -> QFrame:
        bar = QFrame()
        bar.setObjectName("DrawerTitleBar")
        bar.setFixedHeight(38)
        layout = QHBoxLayout(bar)
        layout.setContentsMargins(12, 4, 8, 4)
        layout.setSpacing(8)
        self.drawer_title_label = QLabel("详情抽屉")
        self.drawer_title_label.setObjectName("DrawerTitleLabel")
        close_button = QPushButton("×")
        close_button.setObjectName("TinyButton")
        close_button.setFixedWidth(32)
        close_button.clicked.connect(self._hide_drawer)
        layout.addWidget(self.drawer_title_label, 1)
        layout.addWidget(close_button)
        return bar

    def _make_drawer_page(self) -> QWidget:
        return self.drawer_controller.make_page()

    def _build_outline_tab(self) -> QWidget:
        return drawer_pages.build_outline_tab(self)

    def _build_overview_tab(self) -> QWidget:
        return drawer_pages.build_overview_tab(self)

    def _build_shortcuts_tab(self) -> QWidget:
        return drawer_pages.build_shortcuts_tab(self)

    def _build_characters_tab(self) -> QWidget:
        return drawer_pages.build_characters_tab(self)

    def _ensure_character_detail_dialog(self) -> CharacterDetailDialog:
        return drawer_pages.ensure_character_detail_dialog(self)

    def _build_world_tab(self) -> QWidget:
        return drawer_pages.build_world_tab(self)

    def _build_skills_tab(self) -> QWidget:
        return drawer_pages.build_skills_tab(self)

    def _build_review_tab(self) -> QWidget:
        return drawer_pages.build_review_tab(self)

    def _build_task_tab(self) -> QWidget:
        return drawer_pages.build_task_tab(self)

    def _apply_theme(self) -> None:
        self.tokens = self._build_theme_tokens()
        self.theme_mode = self.tokens.mode
        self.state.theme_mode = self.theme_mode
        video_active = bool(self.background_video and Path(self.background_video).exists())
        if video_active:
            self.background_processed_image = self._prepare_transparent_placeholder()
        else:
            self.background_processed_image = self._prepare_background_image()
        self._tree_icon_cache.clear()
        self.setStyleSheet(build_stylesheet(self.tokens, self.background_processed_image))
        if hasattr(self, "star_graph"):
            self.star_graph.apply_theme(self.tokens.surface, self.tokens.border)

    def _build_theme_tokens(self) -> ThemeTokens:
        base = get_theme(self.theme_preset)
        values = {key: getattr(base, key) for key in base.__dataclass_fields__ if key != "mode"}
        for key, _label in THEME_COLOR_FIELDS:
            color = self._normalize_hex_color(str(self.custom_theme_colors.get(key, "")))
            if color:
                values[key] = color
        mode = "dark" if self.theme_preset == "night" else "light"
        return ThemeTokens(mode=mode, **values)

    def _normalize_hex_color(self, value: str) -> str:
        color = QColor(value.strip())
        return color.name().upper() if color.isValid() else ""

    def _prepare_transparent_placeholder(self) -> str:
        output_dir = self.database.data_dir / "ui"
        output_dir.mkdir(parents=True, exist_ok=True)
        output = output_dir / "video_bg_placeholder.png"
        if not output.exists():
            pixmap = QPixmap(16, 16)
            pixmap.fill(Qt.GlobalColor.transparent)
            pixmap.save(str(output), "PNG")
        return str(output)

    def _prepare_background_image(self) -> str:
        if not self.background_image:
            return ""
        source = Path(self.background_image)
        if not source.exists():
            return ""
        image = QImage(str(source))
        if image.isNull():
            return ""
        image = image.convertToFormat(QImage.Format.Format_ARGB32)
        filter_code = self.background_filter
        strength = max(0, min(100, int(self.background_filter_strength)))
        if filter_code == "blur_dark":
            small_w = max(16, image.width() // 10)
            small_h = max(16, image.height() // 10)
            image = image.scaled(small_w, small_h, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
            image = image.scaled(max(1, small_w * 10), max(1, small_h * 10), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
            overlay = QColor(0, 0, 0, int(160 * strength / 100))
        elif filter_code == "soft_dark":
            overlay = QColor(0, 0, 0, int(150 * strength / 100))
        elif filter_code == "soft_light":
            overlay = QColor(255, 255, 255, int(170 * strength / 100))
        elif filter_code == "warm":
            overlay = QColor(246, 220, 160, int(130 * strength / 100))
        else:
            overlay = QColor(0, 0, 0, 0)
        if overlay.alpha() > 0:
            painter = QPainter(image)
            painter.fillRect(image.rect(), overlay)
            painter.end()
        output_dir = self.database.data_dir / "ui"
        output_dir.mkdir(parents=True, exist_ok=True)
        output = output_dir / "background_processed.png"
        image.save(str(output), "PNG")
        return str(output)

    def _store_background_source(self, source_value: str) -> str:
        if not source_value:
            return ""
        source = Path(source_value)
        if not source.exists():
            raise ValueError("背景图片不存在。")
        target_dir = self.database.data_dir / "ui"
        target_dir.mkdir(parents=True, exist_ok=True)
        suffix = source.suffix.lower() if source.suffix else ".png"
        target = target_dir / f"background_source{suffix}"
        if source.resolve() != target.resolve():
            shutil.copy2(source, target)
        return str(target)

    def _store_background_video_source(self, source_value: str) -> str:
        if not source_value:
            return ""
        source = Path(source_value)
        if not source.exists():
            return ""
        target_dir = self.database.data_dir / "ui"
        target_dir.mkdir(parents=True, exist_ok=True)
        suffix = source.suffix.lower() or ".mp4"
        target = target_dir / f"background_video{suffix}"
        if source.resolve() != target.resolve():
            shutil.copy2(source, target)
        return str(target)

    def _apply_background_video(self) -> None:
        if not hasattr(self, "background_video_label"):
            return
        self._stop_background_video()
        if not self.background_video:
            self.background_video_label.hide()
            if hasattr(self, "background_video_view"):
                self.background_video_view.hide()
            return
        source = Path(self.background_video)
        if not source.exists():
            self.background_video_label.hide()
            if hasattr(self, "background_video_view"):
                self.background_video_view.hide()
            return
        suffix = source.suffix.lower()
        if suffix in {".gif", ".webp", ".apng"}:
            if hasattr(self, "background_video_view"):
                self.background_video_view.hide()
            self.background_movie = QMovie(str(source))
            self.background_movie.setCacheMode(QMovie.CacheMode.CacheAll)
            self.background_video_label.setScaledContents(True)
            self.background_video_label.setMovie(self.background_movie)
            self.background_movie.start()
            self.background_video_label.show()
            self.background_video_label.lower()
            self.root_widget.raise_()
        elif suffix in {".mp4", ".avi", ".mkv", ".webm", ".mov", ".wmv", ".flv"}:
            self.background_video_label.hide()
            if self.background_movie is not None:
                self.background_movie.stop()
                self.background_movie = None
            self.background_media_player = QMediaPlayer()
            self.background_media_player.setVideoOutput(self.background_video_item)
            self.background_media_player.setSource(QUrl.fromLocalFile(str(source)))
            self.background_media_player.mediaStatusChanged.connect(self._on_video_media_status)
            self._resize_background_video_item()
            self.background_video_view.show()
            self.background_video_view.lower()
            self.root_widget.raise_()
            self.background_media_player.play()
        else:
            self.background_video_label.hide()
            if hasattr(self, "background_video_view"):
                self.background_video_view.hide()
            self.background_movie = None
            return

    def _on_video_media_status(self, status: QMediaPlayer.MediaStatus) -> None:
        if status == QMediaPlayer.MediaStatus.EndOfMedia and self.background_media_player:
            self.background_media_player.setPosition(0)
            self.background_media_player.play()
        elif status == QMediaPlayer.MediaStatus.InvalidMedia:
            self._disable_background_video("视频文件无效或解码失败，已自动清除视频背景。")

    def _disable_background_video(self, message: str) -> None:
        self._stop_background_video()
        self.background_video = ""
        self.background_processed_image = ""
        self._apply_theme()
        self._save_ui_settings()
        self._set_status(message)

    def _stop_background_video(self) -> None:
        if self.background_movie is not None:
            self.background_movie.stop()
            self.background_movie = None
        if self.background_media_player is not None:
            self.background_media_player.stop()
            try:
                self.background_media_player.setVideoOutput(None)
                self.background_media_player.setSource(QUrl())
            except TypeError:
                pass
            self.background_media_player = None
        if hasattr(self, "background_video_label") and self.background_video_label:
            self.background_video_label.clear()
            self.background_video_label.hide()
        if hasattr(self, "background_video_view") and self.background_video_view:
            self.background_video_view.hide()

    def _resize_background_video_item(self) -> None:
        if not hasattr(self, "background_video_item"):
            return
        size = self.central_shell.size() if hasattr(self, "central_shell") else self.size()
        width = max(1, size.width())
        height = max(1, size.height())
        rect = QRectF(0, 0, width, height)
        self.background_video_scene.setSceneRect(rect)
        self.background_video_item.setPos(0, 0)
        self.background_video_item.setSize(QSizeF(width, height))

    def _open_view_settings(self) -> None:
        dialog = ViewSettingsDialog(
            self,
            {
                "theme_preset": self.theme_preset,
                "custom_theme_colors": self.custom_theme_colors,
                "ai_colors": self.ai_colors,
                "background_image": self.background_image,
                "background_video": self.background_video,
                "background_filter": self.background_filter,
                "background_filter_strength": self.background_filter_strength,
            },
        )
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        values = dialog.values()
        try:
            self._apply_view_settings(values)
        except Exception as exc:  # noqa: BLE001
            error(self, "视图配置失败", str(exc))

    def _apply_view_settings(self, values: dict[str, Any]) -> None:
        preset = str(values.get("theme_preset", "light_blue"))
        if preset not in PRESET_THEMES and preset != "custom":
            preset = "light_blue"
        colors_raw = values.get("custom_theme_colors", {})
        if not isinstance(colors_raw, dict):
            colors_raw = {}
        normalized_colors: dict[str, str] = {}
        for key, _label in THEME_COLOR_FIELDS:
            color = self._normalize_hex_color(str(colors_raw.get(key, "")))
            if not color:
                raise ValueError(f"颜色值无效：{key}")
            normalized_colors[key] = color
        self.theme_preset = preset
        self.custom_theme_colors = normalized_colors
        self.ai_colors = normalize_ai_color_config(values.get("ai_colors", {}))
        set_ai_probability_meta(self.ai_colors)
        self.background_image = self._store_background_source(str(values.get("background_image", "")))
        self.background_video = self._store_background_video_source(str(values.get("background_video", "")))
        self.background_filter = str(values.get("background_filter", "soft_dark"))
        if self.background_filter not in {code for code, _label in BACKGROUND_FILTERS}:
            self.background_filter = "soft_dark"
        self.background_filter_strength = self._safe_int(values.get("background_filter_strength", 35), 35, 0, 100)
        self._apply_theme()
        self._apply_background_video()
        self._apply_editor_font_settings(values)
        self._refresh_ai_color_dependents()
        self._save_ui_settings()
        self._set_status("视图配置已更新")

    def _apply_editor_font_settings(self, values: dict[str, Any]) -> None:
        font_size = self._safe_int(values.get("editor_font_size", 14), 14, 10, 24)
        line_height = self._safe_int(values.get("editor_line_height", 145), 145, 100, 300) / 100.0
        new_qss = f"font-size: {font_size}pt; line-height: {line_height:.2f};"
        for attr in ("content_editor", "outline_editor", "chapter_outline_editor"):
            if hasattr(self, attr):
                widget = getattr(self, attr)
                existing = widget.styleSheet() or ""
                cleaned = re.sub(r"font-size:\s*[\d.]+pt\s*;\s*line-height:\s*[\d.]+\s*;", "", existing)
                widget.setStyleSheet(f"QPlainTextEdit, QTextEdit {{ {new_qss} }}\n{cleaned}".strip())

    def _refresh_ai_color_dependents(self) -> None:
        if hasattr(self, "library_tree"):
            self._load_library_tree(self.current_tree_key)
        if self.state.selected_chapter_id:
            chapter = self.database.get_chapter(self.state.selected_chapter_id)
            if chapter:
                self._update_ai_probability_badge(
                    str(chapter["ai_probability_level"] or "none"),
                    int(chapter["ai_probability"] or 0),
                )

    def _restore_window_state(self) -> None:
        geometry = self.ui_settings.get("qt_window_geometry")
        if isinstance(geometry, str) and geometry:
            self.restoreGeometry(QByteArray.fromBase64(geometry.encode("ascii")))
        splitter_state = self.ui_settings.get("qt_splitter_state")
        if isinstance(splitter_state, str) and splitter_state:
            self.main_splitter.restoreState(QByteArray.fromBase64(splitter_state.encode("ascii")))
        sizes = self.main_splitter.sizes()
        if sizes and sizes[0] <= self.library_auto_collapse_width:
            self.library_collapsed = True
        self._apply_library_collapsed_state()

    def _setup_shortcuts(self) -> None:
        QShortcut(QKeySequence.StandardKey.Save, self, activated=self._save_current)
        QShortcut(QKeySequence("Ctrl+N"), self, activated=self._create_contextual_item)
        QShortcut(QKeySequence("F11"), self, activated=self._toggle_focus_mode)
        QShortcut(QKeySequence("Ctrl+,"), self, activated=self._open_ai_settings)
        QShortcut(QKeySequence("Ctrl+D"), self, activated=self._detect_current_chapter_ai_probability)
        QShortcut(QKeySequence("Ctrl+Shift+G"), self, activated=self._start_multi_agent_generation)
        QShortcut(QKeySequence("Ctrl+E"), self, activated=lambda: self._export_current("txt"))
        QShortcut(QKeySequence("Ctrl+Shift+E"), self, activated=lambda: self._export_current("docx"))
        QShortcut(QKeySequence("Ctrl+L"), self, activated=self._toggle_library_collapsed)
        QShortcut(QKeySequence("Ctrl+0"), self, activated=self._reset_all_settings)

    def _reset_all_settings(self) -> None:
        if not confirm(self, "重置全部配置", "此操作将清空所有设置（视图、AI、主题、背景图片/视频等），恢复为默认状态。\n\n确定要继续吗？", danger=True):
            return
        self._stop_background_video()
        self.background_image = ""
        self.background_video = ""
        self.background_processed_image = ""
        self.background_filter = "soft_dark"
        self.background_filter_strength = 35
        self.theme_preset = "light_blue"
        self.custom_theme_colors = {}
        self.ai_colors = normalize_ai_color_config({})
        set_ai_probability_meta(self.ai_colors)
        self.library_collapsed = False
        self.library_width = 260
        self.qt_drawer_width = 420
        for path in (self.ui_config_path, self.ai_config_path):
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        self.ui_settings = {}
        self._apply_theme()
        self._apply_background_video()
        self._apply_editor_font_settings({"editor_font_size": 14, "editor_line_height": 145})
        self._refresh_ai_color_dependents()
        self._save_ui_settings()
        self._set_status("全部配置已重置为默认值。")

    def _load_library_tree(
        self,
        preferred_key: tuple[str, int | None, int | None] | None = None,
        *,
        preserve_expansion: bool = True,
    ) -> None:
        expanded_keys = self._collect_expanded_tree_keys() if preserve_expansion else set()
        self.tree_model.removeRows(0, self.tree_model.rowCount())
        root = self.tree_model.invisibleRootItem()
        compact_rows: list[tuple[int, str]] = []
        for book in self.database.list_books():
            book_id = int(book["id"])
            book_title = str(book["title"])
            book_words = self.database.count_book_words(book_id)
            tree_book_title = f"{book_title}  ({book_words}字)" if book_words else book_title
            compact_rows.append((book_id, tree_book_title))
            book_item = self._make_book_tree_item(tree_book_title, book_id)
            root.appendRow(book_item)
            chapters_all = self.database.list_chapters(book_id, ALL_VOLUMES)
            chapters_by_volume: dict[int | None, list[Any]] = {}
            for chapter in chapters_all:
                chapters_by_volume.setdefault(chapter["volume_id"], []).append(chapter)
            for volume in self.database.list_volumes(book_id):
                volume_id = int(volume["id"])
                volume_index = int(volume["sort_order"])
                vol_words = self.database.count_volume_words(volume_id)
                vol_title = f"第{volume_index}卷 · {volume['title']}"
                if vol_words:
                    vol_title += f"  ({vol_words}字)"
                volume_item = self._make_volume_tree_item(vol_title, volume_id)
                book_item.appendRow(volume_item)
                for chapter in chapters_by_volume.get(volume_id, []):
                    volume_item.appendRow(self._make_chapter_tree_item(chapter))
            if chapters_by_volume.get(None):
                group_item = self._make_tree_item("未分卷", "group", None, book_id)
                group_item.setIcon(self._tree_icon("group"))
                book_item.appendRow(group_item)
                for chapter in chapters_by_volume.get(None, []):
                    group_item.appendRow(self._make_chapter_tree_item(chapter))
        self.library_tree.collapseAll()
        self._restore_expanded_tree_keys(expanded_keys)
        self._refresh_library_compact_list(compact_rows)
        target = preferred_key or self.current_tree_key
        if target:
            self._expand_to_key(target)
            self._select_tree_key(target)
        else:
            self._update_library_create_button()

    def _patch_tree_model_drop(self) -> None:
        self._drop_pending = False
        self._drop_applying = False
        def on_rows_removed(parent, first, last):
            if self._drop_applying:
                return
            self._drop_pending = True
        def on_rows_inserted(parent, first, last):
            if self._drop_applying:
                return
            if not self._drop_pending:
                return
            self._drop_pending = False
            QTimer.singleShot(0, self._apply_drop_changes)

        self.tree_model.rowsAboutToBeRemoved.connect(on_rows_removed)
        self.tree_model.rowsInserted.connect(on_rows_inserted)

    def _apply_drop_changes(self) -> None:
        if self._drop_applying:
            return
        self._drop_applying = True
        try:
            self._apply_drop_changes_impl()
        finally:
            self._drop_applying = False

    def _apply_drop_changes_impl(self) -> None:
        updates: list[tuple[int, int | None, int]] = []
        root = self.tree_model.invisibleRootItem()
        for book_row in range(root.rowCount()):
            book_item = root.child(book_row)
            book_id = book_item.data(ID_ROLE)
            vol_order = 0
            for row in range(book_item.rowCount()):
                child = book_item.child(row)
                kind = child.data(KIND_ROLE)
                child_id = child.data(ID_ROLE)
                if kind == "volume" and child_id is not None:
                    vol_order += 1
                    vol_id = int(child_id)
                    chap_order = 0
                    for c_row in range(child.rowCount()):
                        chapter = child.child(c_row)
                        if chapter.data(KIND_ROLE) == "chapter" and chapter.data(ID_ROLE) is not None:
                            chap_order += 1
                            updates.append((int(chapter.data(ID_ROLE)), vol_id, chap_order))
                    self.database.connection.execute(
                        "UPDATE volumes SET sort_order = ? WHERE id = ?", (vol_order, vol_id))
                elif kind == "group":
                    chap_order = 0
                    for c_row in range(child.rowCount()):
                        chapter = child.child(c_row)
                        if chapter.data(KIND_ROLE) == "chapter" and chapter.data(ID_ROLE) is not None:
                            chap_order += 1
                            updates.append((int(chapter.data(ID_ROLE)), None, chap_order))
        self.database.connection.executemany(
            "UPDATE chapters SET volume_id = ?, sort_order = ? WHERE id = ?",
            [(vol_id, order, cid) for cid, vol_id, order in updates],
        )
        self.database.connection.commit()
        self._load_library_tree(self.current_tree_key)
        if self.state.selected_chapter_id:
            self._activate_tree_key(("chapter", self.state.selected_chapter_id, None))

    def _make_tree_item(self, text: str, kind: str, node_id: int | None, group_book_id: int | None = None) -> QStandardItem:
        item = QStandardItem(text)
        item.setEditable(False)
        item.setData(kind, KIND_ROLE)
        item.setData(node_id, ID_ROLE)
        item.setData(group_book_id, GROUP_BOOK_ROLE)
        item.setSizeHint(QSize(0, 24))
        return item

    def _tree_icon(self, kind: str, accent: str = ""):
        key = (
            kind,
            str(accent or ""),
            str(getattr(self.tokens, "primary", "")),
            str(getattr(self.tokens, "surface", "")),
            str(getattr(self.tokens, "mode", "")),
        )
        icon = self._tree_icon_cache.get(key)
        if icon is None:
            icon = make_tree_icon(kind, self.tokens, accent=accent)
            self._tree_icon_cache[key] = icon
        return icon

    def _make_book_tree_item(self, title: str, book_id: int) -> QStandardItem:
        item = self._make_tree_item(title, "book", book_id)
        item.setIcon(self._tree_icon("book"))
        font = QFont()
        font.setBold(True)
        font.setPointSize(11)
        item.setFont(font)
        item.setSizeHint(QSize(0, 28))
        return item

    def _make_volume_tree_item(self, title: str, volume_id: int) -> QStandardItem:
        item = self._make_tree_item(title, "volume", volume_id)
        item.setIcon(self._tree_icon("volume"))
        font = QFont()
        font.setBold(True)
        item.setFont(font)
        return item

    def _make_chapter_tree_item(self, chapter: Any) -> QStandardItem:
        probability, level = normalize_ai_probability_pair(
            str(chapter["ai_probability_level"] or "none"),
            int(chapter["ai_probability"] or 0),
        )
        meta_map = get_ai_probability_meta(self.ai_colors)
        meta = meta_map.get(level, meta_map["none"])
        grade_text = "100%" if level == "certain" else meta["label"].replace("AI ", "")
        data = row_to_dict(chapter)
        is_template = bool(data.get("is_template", False))
        word_count = int(data.get("word_count", 0) or 0)
        title = f"{format_chapter_tree_display_title(chapter['sort_order'], str(chapter['title']), word_count=word_count, is_template=False)}  [{grade_text}]"
        item = self._make_tree_item(title, "chapter", int(chapter["id"]))
        item.setIcon(self._tree_icon("template" if is_template else "chapter", str(meta["tree"])))
        item.setData(level, AI_LEVEL_ROLE)
        item.setForeground(QBrush(QColor(str(meta["tree"]))))
        return item

    def _find_chapter_tree_item(self, chapter_id: int) -> QStandardItem | None:
        def walk(parent: QStandardItem) -> QStandardItem | None:
            for row in range(parent.rowCount()):
                item = parent.child(row)
                if item is None:
                    continue
                if item.data(KIND_ROLE) == "chapter" and int(item.data(ID_ROLE) or 0) == chapter_id:
                    return item
                found = walk(item)
                if found is not None:
                    return found
            return None

        return walk(self.tree_model.invisibleRootItem())

    def _refresh_chapter_tree_item(self, chapter_id: int, content_override: str | None = None) -> None:
        item = self._find_chapter_tree_item(chapter_id)
        if item is None:
            return
        chapter = row_to_dict(self.database.get_chapter(chapter_id))
        if not chapter:
            return
        probability, level = normalize_ai_probability_pair(
            str(chapter.get("ai_probability_level") or "none"),
            int(chapter.get("ai_probability") or 0),
        )
        meta_map = get_ai_probability_meta(self.ai_colors)
        meta = meta_map.get(level, meta_map["none"])
        grade_text = "100%" if level == "certain" else meta["label"].replace("AI ", "")
        content = content_override if content_override is not None else str(chapter.get("content") or "")
        title = format_chapter_tree_display_title(
            chapter.get("sort_order", ""),
            str(chapter.get("title") or ""),
            word_count=count_text_characters(content),
            is_template=False,
        )
        is_template = bool(chapter.get("is_template", False))
        item.setText(f"{title}  [{grade_text}]")
        item.setIcon(self._tree_icon("template" if is_template else "chapter", str(meta["tree"])))
        item.setData(level, AI_LEVEL_ROLE)
        item.setForeground(QBrush(QColor(str(meta["tree"]))))

    def _refresh_library_compact_list(self, books: list[tuple[int, str]]) -> None:
        if not hasattr(self, "library_compact_list"):
            return
        self.library_compact_list.blockSignals(True)
        self.library_compact_list.clear()
        for book_id, title in books:
            label = str(title).strip()[:1] or "书"
            item = QListWidgetItem(label)
            item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            item.setSizeHint(QSize(44, 48))
            item.setToolTip(str(title))
            item.setData(Qt.ItemDataRole.UserRole, book_id)
            self.library_compact_list.addItem(item)
        self.library_compact_list.blockSignals(False)

    def _tree_item_key(self, item: QStandardItem | None) -> tuple[str, int | None, int | None] | None:
        if item is None:
            return None
        kind = item.data(KIND_ROLE)
        if not kind:
            return None
        return (str(kind), item.data(ID_ROLE), item.data(GROUP_BOOK_ROLE))

    def _collect_expanded_tree_keys(self) -> set[tuple[str, int | None, int | None]]:
        expanded_keys: set[tuple[str, int | None, int | None]] = set()
        if not hasattr(self, "library_tree") or not hasattr(self, "tree_model"):
            return expanded_keys

        def walk(parent: QStandardItem) -> None:
            for row in range(parent.rowCount()):
                item = parent.child(row)
                if item is None:
                    continue
                key = self._tree_item_key(item)
                index = self.tree_model.indexFromItem(item)
                if key and self.library_tree.isExpanded(index):
                    expanded_keys.add(key)
                walk(item)

        walk(self.tree_model.invisibleRootItem())
        return expanded_keys

    def _restore_expanded_tree_keys(self, expanded_keys: set[tuple[str, int | None, int | None]]) -> None:
        for key in expanded_keys:
            item = self._find_tree_item(key)
            if item is None:
                continue
            index = self.tree_model.indexFromItem(item)
            self.library_tree.setExpanded(index, True)

    def _tree_key_from_index(self, index) -> tuple[str, int | None, int | None] | None:
        if not index.isValid():
            return None
        return self._tree_item_key(self.tree_model.itemFromIndex(index))

    def _expand_to_key(self, key: tuple[str, int | None, int | None]) -> None:
        item = self._find_tree_item(key)
        if item is None:
            return
        parent = item.parent()
        while parent is not None:
            index = self.tree_model.indexFromItem(parent)
            self.library_tree.expand(index)
            parent = parent.parent()

    def _select_tree_key(self, key: tuple[str, int | None, int | None]) -> None:
        item = self._find_tree_item(key)
        if item is None:
            return
        index = self.tree_model.indexFromItem(item)
        self.suppress_tree_selection = True
        self.library_tree.setCurrentIndex(index)
        self.library_tree.scrollTo(index)
        self.suppress_tree_selection = False

    def _on_compact_book_selected(self, current: QListWidgetItem | None, _previous: QListWidgetItem | None) -> None:
        if not current:
            return
        book_id = int(current.data(Qt.ItemDataRole.UserRole) or 0)
        if not book_id:
            return
        key = ("book", book_id, None)
        if not self._confirm_editor_navigation() or not self._confirm_side_form_navigation():
            if _previous:
                self.library_compact_list.setCurrentItem(_previous)
            return
        self._activate_tree_key(key)
        self._select_tree_key(key)

    def _on_library_double_clicked(self, index) -> None:
        key = self._tree_key_from_index(index)
        if not key or key[0] != "book":
            return
        item = self.tree_model.itemFromIndex(index)
        if item is None:
            return
        expanded = self.library_tree.isExpanded(index)
        self._set_tree_item_expanded_recursive(item, not expanded)

    def _set_tree_item_expanded_recursive(self, item: QStandardItem, expanded: bool) -> None:
        index = self.tree_model.indexFromItem(item)
        self.library_tree.setExpanded(index, expanded)
        for row in range(item.rowCount()):
            child = item.child(row)
            if child:
                self._set_tree_item_expanded_recursive(child, expanded)

    def _find_tree_item(self, key: tuple[str, int | None, int | None]) -> QStandardItem | None:
        def walk(parent: QStandardItem) -> QStandardItem | None:
            for row in range(parent.rowCount()):
                item = parent.child(row)
                if item is None:
                    continue
                if self._tree_item_key(item) == key:
                    return item
                found = walk(item)
                if found:
                    return found
            return None
        return walk(self.tree_model.invisibleRootItem())

    def _on_tree_current_changed(self, current, _previous) -> None:
        if self.suppress_tree_selection:
            return
        key = self._tree_key_from_index(current)
        if not key:
            return
        if not self._confirm_editor_navigation():
            if self.current_tree_key:
                self._select_tree_key(self.current_tree_key)
            return
        if not self._confirm_side_form_navigation():
            if self.current_tree_key:
                self._select_tree_key(self.current_tree_key)
            return
        self._activate_tree_key(key)

    def _activate_tree_key(self, key: tuple[str, int | None, int | None]) -> None:
        kind, node_id, group_book_id = key
        self.current_tree_key = key
        self.state.current_node_kind = kind
        if kind == "book" and node_id is not None:
            self._load_book(node_id)
        elif kind == "volume" and node_id is not None:
            self._load_volume(node_id)
        elif kind == "group" and group_book_id is not None:
            self._load_group(group_book_id)
        elif kind == "chapter" and node_id is not None:
            self._load_chapter(node_id)
        self._refresh_drawer()
        self._update_library_create_button()

    def _update_library_create_button(self) -> None:
        if not hasattr(self, "library_create_button"):
            return
        kind = self.state.current_node_kind
        if kind == "book":
            self.library_create_button.setText("新建卷")
            self.library_create_button.setEnabled(True)
        elif kind in {"volume", "group", "chapter"} and self.state.selected_book_id:
            self.library_create_button.setText("新建章节")
            self.library_create_button.setEnabled(True)
        else:
            self.library_create_button.setText("新建书籍")
            self.library_create_button.setEnabled(True)

    def _confirm_editor_navigation(self) -> bool:
        if not self.state.editor_dirty:
            return True
        choice = ask_unsaved(self)
        if choice == "cancel":
            return False
        if choice == "save":
            return self._save_current()
        self._mark_clean()
        return True

    def _ask_save_discard_cancel(self, title: str, message: str) -> str:
        box = QMessageBox(self)
        box.setWindowTitle(title)
        box.setText(message)
        box.setIcon(QMessageBox.Icon.Question)
        save_button = box.addButton("保存", QMessageBox.ButtonRole.AcceptRole)
        discard_button = box.addButton("不保存", QMessageBox.ButtonRole.DestructiveRole)
        cancel_button = box.addButton("取消", QMessageBox.ButtonRole.RejectRole)
        box.setDefaultButton(save_button)
        box.exec()
        clicked = box.clickedButton()
        if clicked == save_button:
            return "save"
        if clicked == discard_button:
            return "discard"
        if clicked == cancel_button:
            return "cancel"
        return "cancel"

    def _confirm_side_form_navigation(self) -> bool:
        return self._confirm_character_edit_navigation() and self._confirm_world_edit_navigation()

    def _book_outline_volume_specs(self, book_id: int | None) -> list[dict[str, Any]]:
        if not book_id:
            return []
        return [row_to_dict(volume) for volume in self.database.list_volumes(book_id)]

    def _load_book(self, book_id: int) -> None:
        book = self.database.get_book(book_id)
        if not book:
            return
        data = row_to_dict(book)
        self.state.selected_book_id = book_id
        self.state.selected_book_title = str(data["title"])
        self.state.selected_volume_id = None
        self.state.selected_chapter_id = None
        self.state.editor_scope_kind = "book"
        self.state.editor_scope_id = book_id
        self.loaded_chapter_content = ""
        self.loaded_chapter_summary_text = ""
        self._set_editor_title(str(data["title"]), str(data["title"]))
        self.outline_editor.setReadOnly(False)
        self.outline_template_btn.setVisible(True)
        self.outline_markdown_toggle_btn.setVisible(True)
        self.outline_checkbox_btn.setVisible(False)
        self.loading_editor = True
        volume_specs = self._book_outline_volume_specs(book_id)
        outline_data = self.database.get_book_outline_form(book_id)
        self.book_outline_form.set_data(outline_data, self.state.selected_book_title, volume_specs)
        self.outline_editor.setPlainText(book_outline_form_to_markdown(outline_data, self.state.selected_book_title, volume_specs))
        self.loading_editor = False
        self._refresh_outline_markdown_preview()
        self._set_markdown_view_mode("outline", False)
        self._sync_editor_text_lengths()
        self.stack.setCurrentWidget(self.outline_page)
        self._clear_ai_span_rendering()
        self.current_book_cover_path = str(data.get("cover_image_path", ""))
        self._refresh_book_cover_preview()
        self._load_book_side_data(book_id)
        self._mark_clean()
        self._set_status(f"已打开书籍：{data['title']}")

    def _load_volume(self, volume_id: int) -> None:
        volume = self.database.get_volume(volume_id)
        if not volume:
            return
        data = row_to_dict(volume)
        book_id = int(data["book_id"])
        book = self.database.get_book(book_id)
        book_data = row_to_dict(book) if book else {}
        self.state.selected_book_id = book_id
        self.state.selected_book_title = str(book_data.get("title", ""))
        self.state.selected_volume_id = volume_id
        self.state.selected_chapter_id = None
        self.state.editor_scope_kind = "volume"
        self.state.editor_scope_id = volume_id
        self.loaded_chapter_content = ""
        self.loaded_chapter_summary_text = ""
        self._set_editor_title(str(data["title"]), f"{self.state.selected_book_title} / {data['title']}")
        self.outline_editor.setReadOnly(False)
        self.outline_template_btn.setVisible(False)
        self.outline_markdown_toggle_btn.setVisible(False)
        self.outline_checkbox_btn.setVisible(False)
        self._set_markdown_view_mode("outline", False)
        self.loading_editor = True
        self.outline_editor.setPlainText(str(data.get("outline_text", "")))
        self.loading_editor = False
        self.outline_markdown_stack.setCurrentWidget(self.outline_editor)
        self._refresh_outline_markdown_preview()
        self._sync_editor_text_lengths()
        self.stack.setCurrentWidget(self.outline_page)
        self._clear_ai_span_rendering()
        self.current_book_cover_path = str(book_data.get("cover_image_path", ""))
        self._refresh_book_cover_preview()
        self._load_book_side_data(book_id)
        self._mark_clean()
        self._set_status(f"已打开卷：{data['title']}")

    def _load_group(self, book_id: int) -> None:
        self._load_book(book_id)
        self.state.current_node_kind = "group"
        self.state.selected_volume_id = None
        self.state.selected_chapter_id = None
        self.state.editor_scope_kind = None
        self.state.editor_scope_id = None
        self.loaded_chapter_content = ""
        self.loaded_chapter_summary_text = ""
        self._set_editor_title("未分卷章节", f"{self.state.selected_book_title} / 未分卷")
        self.outline_template_btn.setVisible(False)
        self.outline_markdown_toggle_btn.setVisible(False)
        self.outline_checkbox_btn.setVisible(False)
        self._set_markdown_view_mode("outline", False)
        self.loading_editor = True
        self.outline_editor.setPlainText("")
        self.loading_editor = False
        self.outline_markdown_stack.setCurrentWidget(self.outline_editor)
        self._refresh_outline_markdown_preview()
        self._sync_editor_text_lengths()
        self.stack.setCurrentWidget(self.outline_page)
        self.outline_editor.setReadOnly(True)
        self._clear_ai_span_rendering()
        self._mark_clean()

    def _load_chapter(self, chapter_id: int) -> None:
        chapter = self.database.get_chapter(chapter_id)
        if not chapter:
            return
        data = row_to_dict(chapter)
        self._load_book(int(data["book_id"]))
        self.outline_editor.setReadOnly(False)
        self.chapter_outline_editor.setReadOnly(False)
        self.content_editor.setReadOnly(False)
        self.state.current_node_kind = "chapter"
        self.state.selected_chapter_id = chapter_id
        self.state.selected_volume_id = int(data["volume_id"]) if data.get("volume_id") is not None else None
        self.state.editor_scope_kind = "chapter"
        self.state.editor_scope_id = chapter_id
        self._load_book_side_data(int(data["book_id"]))
        self._set_editor_title(str(data["title"]), f"{self.state.selected_book_title} / {data['volume_title']} / {data['title']}")
        self._update_ai_probability_badge(
            str(data.get("ai_probability_level", "none")),
            int(data.get("ai_probability", 0) or 0),
        )
        outline_text = str(data.get("outline", ""))
        content_text = str(data.get("content", ""))
        self.loaded_chapter_content = content_text
        self.loaded_chapter_summary_text = str(data.get("summary_text", ""))
        self.loading_editor = True
        self.chapter_outline_editor.setPlainText(outline_text)
        self.content_editor.setPlainText(content_text)
        self.loading_editor = False
        self._refresh_ai_spans_for_current_chapter()
        self._sync_editor_text_lengths()
        self.stack.setCurrentWidget(self.chapter_page)
        self._mark_clean()
        self._set_status(f"已打开章节：{data['title']}")

    def _load_book_side_data(self, book_id: int) -> None:
        self.global_skills = [dict(item) for item in self.database.list_bound_skills_for_global()]
        self.book_skills = [dict(item) for item in self.database.list_bound_skills_for_book(book_id)]
        self.characters = [dict(item) for item in self.database.list_characters(book_id)]
        self.world_entries = [dict(item) for item in self.database.list_world_entries(book_id)]
        self._load_world_form_for_book(book_id)
        if self.state.selected_chapter_id:
            self.bound_skills = [dict(item) for item in self.database.list_bound_skills_for_chapter(self.state.selected_chapter_id)]
        else:
            self.bound_skills = []

    def _load_world_form_for_book(self, book_id: int | None) -> None:
        if not book_id or not hasattr(self, "world_form"):
            return
        if getattr(self, "_world_form_dirty", False) and getattr(self, "_world_form_book_id", None) == book_id:
            return
        self.world_form.set_data(self.database.get_book_world_form(book_id))
        self._world_form_book_id = book_id
        self._refresh_world_markdown_preview()
        if hasattr(self, "world_markdown_stack"):
            self._set_markdown_view_mode("world", False)
        self._world_form_dirty = False

    def _set_editor_title(self, title: str, path: str) -> None:
        self.editor_title.setText(title)
        self.editor_path.setText(path)
        self._update_ai_probability_badge("none", 0, visible=False)

    def _set_markdown_preview(self, preview: QTextEdit, text: str) -> None:
        try:
            preview.setMarkdown(str(text or ""))
        except Exception:
            preview.setPlainText(str(text or ""))

    def _current_book_outline_markdown(self) -> str:
        if self.state.editor_scope_kind == "book" and hasattr(self, "book_outline_form"):
            return self.book_outline_form.to_markdown(self.state.selected_book_title)
        return self.outline_editor.toPlainText()

    def _current_world_markdown(self) -> str:
        if hasattr(self, "world_form"):
            return self.world_form.to_markdown()
        if self.state.selected_book_id:
            book = self.database.get_book(self.state.selected_book_id)
            if book:
                return str(book["world_text"] or "")
        return ""

    def _on_book_outline_form_changed(self) -> None:
        markdown = self._current_book_outline_markdown()
        self.outline_editor.blockSignals(True)
        try:
            self.outline_editor.setPlainText(markdown)
        finally:
            self.outline_editor.blockSignals(False)
        self._refresh_outline_markdown_preview()
        self._on_editor_changed()

    def _on_world_form_changed(self) -> None:
        self._refresh_world_markdown_preview()
        self._mark_world_form_dirty()

    def _current_character_markdown(self) -> str:
        if hasattr(self, "character_profile_form"):
            return self.character_profile_form.to_markdown(self.character_name.text(), self.character_role.text())
        if self.selected_character_id:
            character = self.database.get_character(self.selected_character_id)
            if character:
                return str(character["profile_text"] or "")
        return ""

    def _on_character_form_changed(self) -> None:
        self._refresh_character_markdown_preview()
        self._mark_character_form_dirty()

    def _on_character_identity_changed(self) -> None:
        if getattr(self, "loading_editor", False):
            return
        if hasattr(self, "character_profile_form"):
            self.character_profile_form.refresh_identity(self.character_name.text(), self.character_role.text())
        self._refresh_character_markdown_preview()
        self._mark_character_form_dirty()

    def _refresh_outline_markdown_preview(self) -> None:
        if hasattr(self, "outline_markdown_preview"):
            self._set_markdown_preview(self.outline_markdown_preview, self._current_book_outline_markdown())

    def _refresh_character_markdown_preview(self) -> None:
        if hasattr(self, "character_profile_preview"):
            self._set_markdown_preview(self.character_profile_preview, self._current_character_markdown())

    def _refresh_world_markdown_preview(self) -> None:
        if hasattr(self, "world_content_preview"):
            self._set_markdown_preview(self.world_content_preview, self._current_world_markdown())

    def _toggle_markdown_view(self, target: str) -> None:
        if target == "outline" and self.state.editor_scope_kind == "book":
            self._set_markdown_view_mode("outline", self.outline_markdown_stack.currentWidget() is self.book_outline_form)
            return
        if target == "character" and hasattr(self, "character_profile_form"):
            self._set_markdown_view_mode("character", self.character_markdown_stack.currentWidget() is self.character_profile_form)
            return
        if target == "world" and hasattr(self, "world_form"):
            self._set_markdown_view_mode("world", self.world_markdown_stack.currentWidget() is self.world_form)
            return
        stack, editor, preview, button = self._markdown_widgets(target)
        if not stack or not editor or not preview or not button:
            return
        self._set_markdown_view_mode(target, stack.currentWidget() is editor)

    def _set_markdown_view_mode(self, target: str, preview_mode: bool) -> None:
        if target == "outline" and hasattr(self, "book_outline_form") and self.state.editor_scope_kind == "book":
            if preview_mode:
                self._set_markdown_preview(self.outline_markdown_preview, self._current_book_outline_markdown())
                self.outline_markdown_stack.setCurrentWidget(self.outline_markdown_preview)
                self.outline_markdown_toggle_btn.setText("表单")
            else:
                self.outline_markdown_stack.setCurrentWidget(self.book_outline_form)
                self.outline_markdown_toggle_btn.setText("Markdown")
            return
        if target == "character" and hasattr(self, "character_profile_form"):
            if preview_mode:
                self._set_markdown_preview(self.character_profile_preview, self._current_character_markdown())
                self.character_markdown_stack.setCurrentWidget(self.character_profile_preview)
                self.character_markdown_toggle_btn.setText("表单")
            else:
                self.character_markdown_stack.setCurrentWidget(self.character_profile_form)
                self.character_markdown_toggle_btn.setText("Markdown")
            return
        if target == "world" and hasattr(self, "world_form"):
            if preview_mode:
                self._set_markdown_preview(self.world_content_preview, self._current_world_markdown())
                self.world_markdown_stack.setCurrentWidget(self.world_content_preview)
                self.world_markdown_toggle_btn.setText("表单")
            else:
                self.world_markdown_stack.setCurrentWidget(self.world_form)
                self.world_markdown_toggle_btn.setText("Markdown")
            return
        stack, editor, preview, button = self._markdown_widgets(target)
        if not stack or not editor or not preview or not button:
            return
        if preview_mode:
            self._set_markdown_preview(preview, editor.toPlainText())
            stack.setCurrentWidget(preview)
            button.setText("Markdown")
        else:
            stack.setCurrentWidget(editor)
            button.setText("预览")

    def _markdown_widgets(self, target: str) -> tuple[QStackedWidget | None, QPlainTextEdit | None, QTextEdit | None, QPushButton | None]:
        if target == "outline" and hasattr(self, "outline_markdown_stack"):
            return self.outline_markdown_stack, self.outline_editor, self.outline_markdown_preview, self.outline_markdown_toggle_btn
        return None, None, None, None

    def _toggle_markdown_checkbox(self, target: str) -> None:
        stack, editor, preview, _button = self._markdown_widgets(target)
        if not stack or not editor:
            return
        if stack.currentWidget() is not editor:
            self._set_markdown_view_mode(target, False)
        if not self._toggle_markdown_checkbox_selection(editor):
            return
        if preview:
            self._set_markdown_preview(preview, editor.toPlainText())
        self._set_status("已切换勾选状态")

    def _toggle_markdown_checkbox_selection(self, editor: QPlainTextEdit) -> bool:
        text = editor.toPlainText()
        cursor = editor.textCursor()
        if not text:
            editor.setPlainText("- [x] ")
            return True

        lines = text.splitlines(keepends=True)
        if not lines:
            return False

        start = min(cursor.selectionStart(), len(text))
        end = min(cursor.selectionEnd(), len(text))
        end_for_line = max(start, end - 1) if end > start else start
        start_line = min(text.count("\n", 0, start), len(lines) - 1)
        end_line = min(text.count("\n", 0, end_for_line), len(lines) - 1)
        changed = False
        for line_index in range(start_line, end_line + 1):
            original = lines[line_index]
            body = original.rstrip("\r\n")
            newline = original[len(body):]
            toggled = self._toggle_markdown_checkbox_line(body)
            if toggled != body:
                lines[line_index] = toggled + newline
                changed = True
        if not changed:
            return False

        old_position = cursor.position()
        old_anchor = cursor.anchor()
        editor.setPlainText("".join(lines))
        new_cursor = editor.textCursor()
        text_length = len(editor.toPlainText())
        new_cursor.setPosition(min(old_anchor, text_length))
        new_cursor.setPosition(min(old_position, text_length), QTextCursor.MoveMode.KeepAnchor)
        editor.setTextCursor(new_cursor)
        return True

    def _toggle_markdown_checkbox_line(self, line: str) -> str:
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or stripped.startswith("```") or stripped in {"---", "***", "___"}:
            return line
        match = re.match(r"^(\s*(?:[-*+]\s+)?)\[( |x|X)\](.*)$", line)
        if match:
            next_box = "[x]" if match.group(2) == " " else "[ ]"
            return f"{match.group(1)}{next_box}{match.group(3)}"
        match = re.match(r"^(\s*)☐(.*)$", line)
        if match:
            return f"{match.group(1)}☑{match.group(2)}"
        match = re.match(r"^(\s*)☑(.*)$", line)
        if match:
            return f"{match.group(1)}☐{match.group(2)}"
        match = re.match(r"^(\s*)([-*+]\s+)(.*)$", line)
        if match:
            return f"{match.group(1)}{match.group(2)}[x] {match.group(3).lstrip()}"
        indent_match = re.match(r"^(\s*)", line)
        indent = indent_match.group(1) if indent_match else ""
        return f"{indent}- [x] {line[len(indent):]}"

    def _fill_book_outline_template(self) -> None:
        if self.state.editor_scope_kind != "book":
            info(self, "书籍模板", "请先选中书籍级大纲。")
            return
        if not confirm(self, "重置书籍模板", "当前书籍级设定会被完整空模板替换，确定继续吗？", danger=True):
            return
        volume_specs = self._book_outline_volume_specs(self.state.selected_book_id)
        self.book_outline_form.set_data(
            book_outline_form_defaults(self.state.selected_book_title, volume_specs),
            self.state.selected_book_title,
            volume_specs,
        )
        self._on_book_outline_form_changed()

    def _fill_character_template(self) -> None:
        if not hasattr(self, "character_profile_form"):
            return
        if getattr(self, "_character_form_dirty", False) and not confirm(
            self,
            "重置人物模板",
            "当前人物卡会被完整空模板替换，确定继续吗？",
            danger=True,
        ):
            return
        self.character_profile_form.set_data(
            character_profile_form_defaults(self.character_name.text(), self.character_role.text()),
            self.character_name.text(),
            self.character_role.text(),
        )
        self._on_character_form_changed()
        self._set_markdown_view_mode("character", False)

    def _fill_world_template(self) -> None:
        if not hasattr(self, "world_form"):
            return
        if getattr(self, "_world_form_dirty", False) and not confirm(
            self,
            "重置世界观模板",
            "当前世界观设定会被完整空模板替换，确定继续吗？",
            danger=True,
        ):
            return
        self.world_form.set_data(world_form_defaults())
        self._on_world_form_changed()
        self._set_markdown_view_mode("world", False)

    def _replace_with_template(self, editor: QPlainTextEdit, template: str, label: str) -> None:
        if editor.toPlainText().strip() and not confirm(
            self,
            f"填入{label}模板",
            f"当前{label}内容会被模板替换，确定继续吗？",
            danger=True,
        ):
            return
        editor.setPlainText(template)
        self._set_status(f"已填入{label} Markdown 模板")

    def _update_ai_probability_badge(self, level: str, probability: int, visible: bool = True) -> None:
        if not hasattr(self, "ai_probability_badge"):
            return
        probability, normalized = normalize_ai_probability_pair(level, probability)
        meta = get_ai_probability_meta(self.ai_colors)[normalized]
        label = meta["label"] if normalized == "certain" else f"{meta['label']} · {probability}%"
        self.ai_probability_badge.setText(label)
        self.ai_probability_badge.setStyleSheet(
            "QLabel#AiProbabilityBadge {"
            f"background: {meta['bg']};"
            f"color: {meta['fg']};"
            "border-radius: 10px;"
            "padding: 4px 10px;"
            "font-weight: 700;"
            "}"
        )
        self.ai_probability_badge.setVisible(visible)
        self.ai_probability_button.setVisible(visible)

    def _detect_current_chapter_ai_probability(self) -> None:
        if self.state.current_node_kind != "chapter" or not self.state.selected_chapter_id:
            return
        ai_service = self._require_ai_service("detector", "AI 概率检测")
        if ai_service is None:
            return
        content = self.content_editor.toPlainText()
        if not content.strip():
            info(self, "AI 概率检测", "当前章节正文为空，无法检测。")
            return
        if len(content.strip()) < 200:
            chapter_id = int(self.state.selected_chapter_id)
            self.database.update_chapter_ai_probability(chapter_id, 0, "none")
            self._update_ai_probability_badge("none", 0)
            self._refresh_chapter_tree_item(chapter_id, content)
            self._set_status("正文过短，已跳过 AI 概率检测。")
            return
        chapter_id = int(self.state.selected_chapter_id)
        self._open_drawer("tasks")

        def work() -> tuple[int, int, str]:
            probability, level = ai_service.detect_ai_probability(content)
            return chapter_id, probability, level

        def done(payload: Any) -> None:
            result_chapter_id, probability, level = payload
            probability, level = normalize_ai_probability_pair(str(level), int(probability))
            self.database.update_chapter_ai_probability(int(result_chapter_id), probability, level)
            if self.state.selected_chapter_id == int(result_chapter_id):
                self._update_ai_probability_badge(level, probability)
            self._refresh_chapter_tree_item(int(result_chapter_id))
            self._switch_drawer_tab("review")

        self._run_background_ai_job(label="AI 概率检测", callback=work, on_done=done)

    def _on_editor_changed(self) -> None:
        if self.loading_editor:
            return
        sender = self.sender()
        if (
            sender is self.content_editor
            and self.state.selected_chapter_id
            and not self.ai_writing
            and not self.suppress_ai_span_clear
            and self.current_ai_spans
        ):
            self.database.delete_chapter_ai_spans(self.state.selected_chapter_id)
            self.current_ai_spans = []
            self.content_editor.setExtraSelections([])
            self._refresh_ai_ratio_status()
        self._record_editor_delta()
        self.state.editor_dirty = True
        self.save_label.setText("未保存")
        self._update_dirty_indicator()
        self.word_timer.start(180)

    def _sync_editor_text_lengths(self) -> None:
        for editor_name in ("outline_editor", "chapter_outline_editor", "content_editor"):
            editor = getattr(self, editor_name, None)
            if editor is not None:
                self.editor_text_lengths[id(editor)] = len(editor.toPlainText())

    def _record_editor_delta(self) -> None:
        sender = self.sender()
        if not isinstance(sender, QPlainTextEdit):
            return
        key = id(sender)
        current_length = len(sender.toPlainText())
        previous_length = self.editor_text_lengths.get(key, current_length)
        self.editor_text_lengths[key] = current_length
        delta = max(0, current_length - previous_length)
        if delta <= 0:
            return
        source = "ai" if self.ai_writing else "user"
        self._record_input_stat(source, delta)

    def _record_input_stat(self, source: str, count: int) -> None:
        if count <= 0:
            return
        now = datetime.now()
        self.input_events.append((now, source, int(count)))
        if source == "ai":
            self.ai_total_chars += int(count)
            self.ai_total_tokens = int(self.ai_total_chars * 1.5)
        self._refresh_input_stats_table()

    def _refresh_input_stats_table(self) -> None:
        if not hasattr(self, "input_stat_value_labels"):
            return
        now = datetime.now()
        cutoff = now.timestamp() - 3600.0
        self.input_events = [(ts, src, cnt) for ts, src, cnt in self.input_events if ts.timestamp() >= cutoff]
        user_count = sum(cnt for ts, src, cnt in self.input_events if src == "user")
        ai_count = sum(cnt for ts, src, cnt in self.input_events if src == "ai")
        total_count = user_count + ai_count
        cells = {
            "user": f"{user_count} 字/小时",
            "ai": f"{ai_count} 字/小时 · {self.ai_total_tokens} tokens",
            "total": f"{total_count} 字/小时",
        }
        for key, value in cells.items():
            self.input_stat_value_labels[key].setText(value)

    def _copy_full_chapter(self) -> None:
        if self.state.editor_scope_kind != "chapter":
            return
        text = self.content_editor.toPlainText()
        if not text.strip():
            return
        QApplication.clipboard().setText(text)
        self._set_status("全章已复制到剪贴板")

    def _toggle_content_visibility(self) -> None:
        if self.state.editor_scope_kind != "chapter":
            return
        if not hasattr(self, "content_editor"):
            return
        if self.content_editor.isVisible():
            self.content_editor.hide()
            if hasattr(self, "content_label"):
                self.content_label.hide()
            self._set_status("正文已隐藏")
        else:
            self.content_editor.show()
            if hasattr(self, "content_label"):
                self.content_label.show()
            self._set_status("正文已显示")

    def _show_search_bar(self) -> None:
        if not self.state.selected_book_id:
            info(self, "查找", "请先选择一本书、一个卷或一个章节。")
            return
        if self.search_dialog is None:
            self.search_dialog = SearchDialog(self)
            self.search_dialog.destroyed.connect(lambda _obj=None: setattr(self, "search_dialog", None))
        options, current_scope = self._search_scope_options()
        self.search_dialog.set_scope_options(options, current_scope)
        self.search_dialog.show()
        self.search_dialog.raise_()
        self.search_dialog.activateWindow()
        self.search_dialog.focus_query()

    def _search_scope_options(self) -> tuple[list[tuple[str, str]], str]:
        options: list[tuple[str, str]] = []
        book_id = self.state.selected_book_id
        if not book_id:
            return options, ""
        default_scope = "book"
        chapter_id = self.state.selected_chapter_id
        volume_id = self.state.selected_volume_id
        if self.state.current_node_kind == "chapter" and chapter_id:
            chapter = self.database.get_chapter(chapter_id)
            chapter_title = str(chapter["title"]) if chapter else "当前章节"
            options.append(("chapter", f"当前章节：{chapter_title}"))
            default_scope = "chapter"
            if chapter and chapter["volume_id"] is not None:
                volume_id = int(chapter["volume_id"])
            elif chapter:
                volume_id = None
        if self.state.current_node_kind in {"volume", "group"}:
            default_scope = "volume"
        if volume_id:
            volume = self.database.get_volume(volume_id)
            volume_title = str(volume["title"]) if volume else "当前卷"
            options.append(("volume", f"当前卷：{volume_title}"))
        elif self.state.current_node_kind == "group" or (chapter_id and not volume_id):
            options.append(("volume", "未分卷章节"))
        book_label = self.state.selected_book_title or "当前书籍"
        options.append(("book", f"当前书籍：{book_label}"))
        seen: set[str] = set()
        unique_options: list[tuple[str, str]] = []
        for scope, label in options:
            if scope in seen:
                continue
            unique_options.append((scope, label))
            seen.add(scope)
        if default_scope not in seen:
            default_scope = unique_options[0][0] if unique_options else ""
        return unique_options, default_scope

    def _run_search_from_dialog(
        self,
        *,
        query: str,
        scope: str,
        case_sensitive: bool,
        whole_word: bool,
    ) -> None:
        query = query.strip()
        if not query:
            self._clear_search_results()
            return
        units = self._search_units_for_scope(scope)
        matches: list[dict[str, Any]] = []
        for unit in units:
            text = str(unit.get("text", ""))
            for start, end in _find_plain_text_matches(
                text,
                query,
                case_sensitive=case_sensitive,
                whole_word=whole_word,
            ):
                matches.append(
                    {
                        **unit,
                        "start": start,
                        "end": end,
                        "snippet": _search_snippet(text, start, end),
                    }
                )
                if len(matches) >= 500:
                    break
            if len(matches) >= 500:
                break
        self._search_query = query
        self._search_scope = scope
        self._search_case_sensitive = case_sensitive
        self._search_whole_word = whole_word
        self._search_matches = matches
        if not self._search_matches:
            self._search_match_index = -1
            if self.search_dialog:
                self.search_dialog.set_results([], -1)
            self._set_status(f"未找到「{query}」")
            return
        self._search_match_index = 0
        if self.search_dialog:
            self.search_dialog.set_results(self._search_matches, self._search_match_index)
        self._activate_search_result(self._search_match_index)

    def _clear_search_results(self) -> None:
        self._search_matches = []
        self._search_match_index = -1
        self._search_query = ""
        self._search_scope = ""
        self._search_case_sensitive = False
        self._search_whole_word = False
        if self.search_dialog:
            self.search_dialog.set_results([], -1)

    def _search_units_for_scope(self, scope: str) -> list[dict[str, Any]]:
        book_id = self.state.selected_book_id
        if not book_id:
            return []
        units: list[dict[str, Any]] = []
        if scope == "book":
            book = self.database.get_book(book_id)
            if book:
                book_outline = self.outline_editor.toPlainText() if self.state.editor_scope_kind == "book" else str(book["outline_text"] or "")
                self._append_search_unit(
                    units,
                    title=f"{book['title']} / 书籍大纲",
                    field="outline",
                    field_label="书籍大纲",
                    text=book_outline,
                    node_key=("book", book_id, None),
                )
            for volume in self.database.list_volumes(book_id):
                volume_data = row_to_dict(volume)
                self._append_search_unit(
                    units,
                    title=f"{volume_data.get('title', '')} / 卷大纲",
                    field="outline",
                    field_label="卷大纲",
                    text=str(volume_data.get("outline_text", "") or ""),
                    node_key=("volume", int(volume_data["id"]), None),
                )
            chapters = self.database.list_chapters(book_id)
        elif scope == "volume":
            volume_filter: int | str
            if self.state.selected_volume_id:
                volume_filter = self.state.selected_volume_id
                volume = self.database.get_volume(self.state.selected_volume_id)
                if volume:
                    volume_outline = self.outline_editor.toPlainText() if self.state.editor_scope_kind == "volume" else str(volume["outline_text"] or "")
                    self._append_search_unit(
                        units,
                        title=f"{volume['title']} / 卷大纲",
                        field="outline",
                        field_label="卷大纲",
                        text=volume_outline,
                        node_key=("volume", self.state.selected_volume_id, None),
                    )
            else:
                volume_filter = UNASSIGNED_VOLUMES
            chapters = self.database.list_chapters(book_id, volume_filter)
        elif scope == "chapter" and self.state.selected_chapter_id:
            chapter = self.database.get_chapter(self.state.selected_chapter_id)
            chapters = [chapter] if chapter else []
        else:
            chapters = []

        for chapter_row in chapters:
            if not chapter_row:
                continue
            chapter_summary = row_to_dict(chapter_row)
            detail = self.database.get_chapter(int(chapter_summary["id"]))
            if not detail:
                continue
            self._append_chapter_search_units(units, row_to_dict(detail))
        return units

    def _append_chapter_search_units(self, units: list[dict[str, Any]], chapter: dict[str, Any]) -> None:
        chapter_id = int(chapter["id"])
        title = str(chapter.get("title", "") or "")
        outline = str(chapter.get("outline", "") or "")
        content = str(chapter.get("content", "") or "")
        summary = str(chapter.get("summary_text", "") or "")
        if self.state.editor_scope_kind == "chapter" and self.state.editor_scope_id == chapter_id:
            outline = self.chapter_outline_editor.toPlainText()
            content = self.content_editor.toPlainText()
            summary = self.loaded_chapter_summary_text
        node_key = ("chapter", chapter_id, None)
        volume_title = str(chapter.get("volume_title", "") or "未分卷")
        base_title = f"{volume_title} / {title}"
        self._append_search_unit(units, title=base_title, field="title", field_label="标题", text=title, node_key=node_key)
        self._append_search_unit(units, title=base_title, field="outline", field_label="章节大纲", text=outline, node_key=node_key)
        self._append_search_unit(units, title=base_title, field="summary", field_label="章节摘要", text=summary, node_key=node_key)
        self._append_search_unit(units, title=base_title, field="content", field_label="正文", text=content, node_key=node_key)

    def _append_search_unit(
        self,
        units: list[dict[str, Any]],
        *,
        title: str,
        field: str,
        field_label: str,
        text: str,
        node_key: tuple[str, int | None, int | None],
    ) -> None:
        if not str(text or "").strip():
            return
        units.append(
            {
                "title": title,
                "field": field,
                "field_label": field_label,
                "text": str(text),
                "node_key": node_key,
            }
        )

    def _highlight_search_match(self) -> None:
        if not self._search_matches or self._search_match_index < 0:
            return
        self._activate_search_result(self._search_match_index)

    def _search_next_match(self) -> None:
        if not self._search_matches:
            return
        self._activate_search_result((self._search_match_index + 1) % len(self._search_matches))

    def _search_previous_match(self) -> None:
        if not self._search_matches:
            return
        self._activate_search_result((self._search_match_index - 1) % len(self._search_matches))

    def _activate_search_result(self, index: int) -> None:
        if not self._search_matches:
            return
        index = index % len(self._search_matches)
        result = self._search_matches[index]
        node_key = result.get("node_key")
        if not isinstance(node_key, tuple) or len(node_key) != 3:
            return
        if self.current_tree_key != node_key:
            if not self._confirm_editor_navigation() or not self._confirm_side_form_navigation():
                return
            self._activate_tree_key(node_key)
            self._select_tree_key(node_key)
        self._search_match_index = index
        field = str(result.get("field", ""))
        start = int(result.get("start", 0) or 0)
        end = int(result.get("end", start) or start)
        editor: QPlainTextEdit | None = None
        if field == "content" and self.state.editor_scope_kind == "chapter":
            editor = self.content_editor
        elif field == "outline":
            if self.state.editor_scope_kind == "chapter":
                editor = self.chapter_outline_editor
            elif self.state.editor_scope_kind in {"book", "volume"}:
                editor = self.outline_editor
        if editor is not None:
            self._select_search_text(editor, start, end)
        if self.search_dialog:
            self.search_dialog.set_results(self._search_matches, self._search_match_index)
        self._set_status(f"匹配 {self._search_match_index + 1}/{len(self._search_matches)} · {result.get('field_label', '')}")

    def _select_search_text(self, editor: QPlainTextEdit, start: int, end: int) -> None:
        text_length = len(editor.toPlainText())
        start = max(0, min(start, text_length))
        end = max(start, min(end, text_length))
        cursor = editor.textCursor()
        cursor.setPosition(start)
        cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
        editor.setTextCursor(cursor)
        editor.setFocus()
        editor.ensureCursorVisible()

    def keyPressEvent(self, event) -> None:
        if event.key() == Qt.Key.Key_F3 and self._search_matches and hasattr(self, "_search_match_index"):
            if event.modifiers() & Qt.KeyboardModifier.ShiftModifier:
                self._search_previous_match()
            else:
                self._search_next_match()
            return
        if event.key() == Qt.Key.Key_F and event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            self._show_search_bar()
            return
        super().keyPressEvent(event)

    def _auto_correct_chapter(self) -> None:
        if self.state.editor_scope_kind != "chapter" or not self.state.selected_chapter_id:
            info(self, "纠错", "请先选择一个章节。")
            return
        content = self.content_editor.toPlainText()
        if not content.strip():
            info(self, "纠错", "当前章节正文为空。")
            return
        ai_service = self._require_ai_service("correction", "AI 纠错")
        if ai_service is None:
            return
        if self._find_active_background_job("AI 纠错"):
            info(self, "任务运行中", "AI 纠错已经在运行。")
            return

        def work() -> list[dict[str, Any]]:
            return ai_service.detect_chinese_typos(content, require_remote=True)

        def done(payload: Any) -> None:
            if self.content_editor.toPlainText() != content:
                self._clear_correction_highlights()
                info(self, "纠错", "章节内容已变化，本次 AI 纠错结果已丢弃。请重新运行纠错。")
                return
            findings = payload if isinstance(payload, list) else []
            self._show_correction_findings(content, findings, f"AI 纠错完成：{len(findings)} 处。")

        self._clear_correction_highlights()
        self._set_status("AI 纠错校对中。")
        self._run_background_ai_job(label="AI 纠错", callback=work, on_done=done)

    def _show_correction_findings(self, content: str, findings: list[dict[str, Any]], status_message: str = "") -> None:
        if not findings:
            info(self, "纠错", "未发现错别字。")
            self._clear_correction_highlights()
            return
        self._apply_correction_highlights(content, findings)
        self._set_status(status_message or f"纠错完成：{len(findings)} 处疑似错别字已高亮")

    def _apply_correction_highlights(self, content: str, findings: list[dict[str, Any]]) -> None:
        if not hasattr(self, "content_editor"):
            return
        doc = self.content_editor.document()
        selections: list[QTextEdit.ExtraSelection] = []
        for f_item in findings:
            start = max(0, min(len(content), int(f_item.get("start", 0))))
            end = max(start, min(len(content), int(f_item.get("end", start))))
            if end <= start:
                continue
            cursor = QTextCursor(doc)
            cursor.setPosition(start)
            cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
            fmt = QTextCharFormat()
            if f_item.get("severity") == "medium":
                fmt.setBackground(QColor(255, 180, 100))
                fmt.setUnderlineStyle(QTextCharFormat.UnderlineStyle.WaveUnderline)
                fmt.setUnderlineColor(QColor(220, 80, 20))
            else:
                fmt.setBackground(QColor(255, 230, 160))
                fmt.setUnderlineStyle(QTextCharFormat.UnderlineStyle.WaveUnderline)
                fmt.setUnderlineColor(QColor(180, 140, 40))
            fmt.setToolTip(f"{f_item.get('wrong', '')} → {f_item.get('suggestion', '')}")
            sel = QTextEdit.ExtraSelection()
            sel.cursor = cursor
            sel.format = fmt
            selections.append(sel)
        self.content_editor.setExtraSelections(selections)

    def _clear_correction_highlights(self) -> None:
        if hasattr(self, "content_editor") and self.content_editor:
            self.content_editor.setExtraSelections([])

    def _auto_format_content(self) -> None:
        if self.state.editor_scope_kind != "chapter" or not hasattr(self, "content_editor"):
            return
        text = self.content_editor.toPlainText()
        if not text.strip():
            return
        text = re.sub(r"([a-zA-Z0-9]+)([\u4e00-\u9fff])", r"\1 \2", text)
        text = re.sub(r"([\u4e00-\u9fff])([a-zA-Z0-9]+)", r"\1 \2", text)
        text = re.sub(r",", "，", text)
        text = re.sub(r"(?<!\d)\.(?!\d)", "。", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        self.content_editor.setPlainText(text)
        self.state.editor_dirty = True
        self.save_label.setText("未保存")
        self._set_status("已排版（中英文空格、全角标点、多余空行清理）")

    def _show_polish_diff_dialog(self, job: dict[str, Any]) -> None:
        original = str(job.get("base_content", ""))
        polished = self.content_editor.toPlainText()
        target_chapter_id = job.get("target_chapter_id")
        dialog = QDialog(self)
        dialog.setWindowTitle("润色结果对照")
        dialog.resize(900, 600)
        layout = QVBoxLayout(dialog)
        splitter = QSplitter(Qt.Orientation.Horizontal)
        left_widget = QWidget()
        left_layout = QVBoxLayout(left_widget)
        left_layout.setContentsMargins(0, 0, 0, 0)
        left_label = QLabel("原文")
        left_label.setObjectName("MetaLabel")
        left_layout.addWidget(left_label)
        left_editor = QPlainTextEdit()
        left_editor.setReadOnly(True)
        left_editor.setPlainText(original)
        left_layout.addWidget(left_editor)
        right_widget = QWidget()
        right_layout = QVBoxLayout(right_widget)
        right_layout.setContentsMargins(0, 0, 0, 0)
        right_label = QLabel("润色后")
        right_label.setObjectName("MetaLabel")
        right_layout.addWidget(right_label)
        right_editor = QPlainTextEdit()
        right_editor.setPlainText(polished)
        right_layout.addWidget(right_editor)
        splitter.addWidget(left_widget)
        splitter.addWidget(right_widget)
        layout.addWidget(splitter)
        buttons = QHBoxLayout()
        accept_btn = QPushButton("接受润色")
        accept_btn.setObjectName("PrimaryButton")
        reject_btn = QPushButton("拒绝，恢复原文")
        reject_btn.setObjectName("DangerButton")
        manual_btn = QPushButton("在编辑器中手动调整")
        buttons.addWidget(accept_btn)
        buttons.addWidget(reject_btn)
        buttons.addWidget(manual_btn)
        buttons.addStretch(1)
        layout.addLayout(buttons)
        result = {"action": "manual"}

        def on_accept() -> None:
            result["action"] = "accept"
            dialog.accept()

        def on_reject() -> None:
            if confirm(dialog, "恢复原文", "确定放弃润色结果，恢复为原文吗？"):
                result["action"] = "reject"
                dialog.accept()

        def on_manual() -> None:
            dialog.accept()

        accept_btn.clicked.connect(on_accept)
        reject_btn.clicked.connect(on_reject)
        manual_btn.clicked.connect(on_manual)
        dialog.exec()
        action = result["action"]
        if action == "accept":
            if target_chapter_id:
                self.content_editor.setPlainText(polished)
                self.state.editor_dirty = True
                self.save_label.setText("未保存")
            self._set_status("已接受润色结果")
        elif action == "reject":
            self.content_editor.setPlainText(original)
            self.state.editor_dirty = True
            self.save_label.setText("未保存")
            self._set_status("已恢复原文")
        else:
            self.state.editor_dirty = True
            self.save_label.setText("未保存")
            self._set_status("润色结果已加载到编辑器，可手动调整")

    def _update_word_count(self) -> None:
        if self.state.editor_scope_kind == "chapter":
            if hasattr(self, "content_editor") and self.content_editor.isVisible():
                text = self.content_editor.toPlainText()
            elif hasattr(self, "outline_editor"):
                text = self.outline_editor.toPlainText()
            else:
                self.word_label.setText("字数 0")
                return
        else:
            text = self.outline_editor.toPlainText()
        self.word_label.setText(f"字数 {len(text.strip())}")
        self._update_selection_count()
        self._refresh_ai_ratio_status()

    def _update_selection_count(self) -> None:
        if not hasattr(self, "selection_label"):
            return
        editor = QApplication.focusWidget()
        if editor not in {getattr(self, "content_editor", None), getattr(self, "chapter_outline_editor", None), getattr(self, "outline_editor", None)}:
            if self.state.editor_scope_kind == "chapter":
                editor = self.content_editor
            else:
                editor = self.outline_editor
        selected = ""
        if hasattr(editor, "textCursor"):
            selected = editor.textCursor().selectedText().replace("\u2029", "\n").strip()
        self.selection_label.setText(f"选中 {len(selected)} 字")

    def _clear_ai_span_rendering(self) -> None:
        self.current_ai_spans = []
        if hasattr(self, "content_editor"):
            self.content_editor.setExtraSelections([])
        self._refresh_ai_ratio_status()

    def _refresh_ai_spans_for_current_chapter(self) -> None:
        if not self.state.selected_chapter_id:
            self._clear_ai_span_rendering()
            return
        self.current_ai_spans = [dict(item) for item in self.database.list_chapter_ai_spans(self.state.selected_chapter_id)]
        self._render_ai_spans()
        self._refresh_ai_ratio_status()

    def _render_ai_spans(self) -> None:
        selections: list[QTextEdit.ExtraSelection] = []
        text_length = len(self.content_editor.toPlainText())
        for span in self.current_ai_spans:
            start = max(0, min(text_length, int(span.get("start_offset", 0))))
            end = max(start, min(text_length, int(span.get("end_offset", 0))))
            if end <= start:
                continue
            cursor = QTextCursor(self.content_editor.document())
            cursor.setPosition(start)
            cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
            selection = QTextEdit.ExtraSelection()
            selection.cursor = cursor
            fmt = QTextCharFormat()
            fmt.setForeground(QBrush(QColor("#C92A2A")))
            fmt.setBackground(QBrush(QColor(216, 58, 52, 32)))
            selection.format = fmt
            selections.append(selection)
        self.content_editor.setExtraSelections(selections)

    def _refresh_ai_ratio_status(self) -> None:
        if not hasattr(self, "ai_ratio_label"):
            return
        if self.state.editor_scope_kind != "chapter":
            self.ai_ratio_label.setText("AI标记 0%")
            return
        total = max(1, len(self.content_editor.toPlainText()))
        marked = 0
        for span in self.current_ai_spans:
            start = max(0, int(span.get("start_offset", 0)))
            end = min(total, int(span.get("end_offset", 0)))
            marked += max(0, end - start)
        ratio = int(round(marked * 100 / total)) if total else 0
        self.ai_ratio_label.setText(f"AI标记 {ratio}%")

    def _clear_current_ai_spans(self) -> None:
        if not self.state.selected_chapter_id:
            return
        self.database.delete_chapter_ai_spans(self.state.selected_chapter_id)
        self._refresh_ai_spans_for_current_chapter()
        self._set_status("已清除本章 AI 标记")

    def _mark_clean(self) -> None:
        self.state.editor_dirty = False
        self.save_label.setText("已保存")
        self._update_dirty_indicator()
        self._update_word_count()

    def _update_dirty_indicator(self) -> None:
        dirty = self.state.editor_dirty
        if hasattr(self, "editor_title"):
            current = self.editor_title.text()
            if dirty and not current.startswith("● "):
                self.editor_title.setText(f"● {current}")
            elif not dirty and current.startswith("● "):
                self.editor_title.setText(current[2:])
        base_title = APP_DISPLAY_NAME
        if hasattr(self, "state") and dirty:
            self.setWindowTitle(f"* {base_title}")
        else:
            self.setWindowTitle(base_title)

    def _save_current(self) -> bool:
        kind = self.state.editor_scope_kind
        scope_id = self.state.editor_scope_id
        if not kind or scope_id is None:
            return True
        try:
            if kind == "book":
                outline_data = self.book_outline_form.to_data(self.state.selected_book_title)
                volume_specs = self._book_outline_volume_specs(int(scope_id))
                outline = book_outline_form_to_markdown(outline_data, self.state.selected_book_title, volume_specs)
                self.database.update_book_outline(scope_id, outline, outline_data=outline_data)
                self.outline_editor.blockSignals(True)
                try:
                    self.outline_editor.setPlainText(outline)
                finally:
                    self.outline_editor.blockSignals(False)
                self._refresh_outline_markdown_preview()
                self._sync_truth_files(int(scope_id))
                self._index_book_design_for_rag(int(scope_id))
            elif kind == "volume":
                self.database.update_volume_outline(scope_id, self.outline_editor.toPlainText())
            elif kind == "chapter":
                content = self.content_editor.toPlainText()
                summary_text = "" if content != self.loaded_chapter_content else self.loaded_chapter_summary_text
                self.database.update_chapter(
                    scope_id,
                    outline=self.chapter_outline_editor.toPlainText(),
                    content=content,
                    summary_text=summary_text,
                )
                self.loaded_chapter_content = content
                self.loaded_chapter_summary_text = summary_text
                self._refresh_chapter_tree_item(int(scope_id), content)
                self._sync_truth_files(self.state.selected_book_id, int(scope_id))
                self._index_chapter_for_rag(self.state.selected_book_id, int(scope_id))
            self._mark_clean()
            self._set_status("已保存")
            return True
        except Exception as exc:  # noqa: BLE001
            error(self, "保存失败", str(exc))
            return False

    def _show_library_context_menu(self, pos) -> None:
        index = self.library_tree.indexAt(pos)
        key = self._tree_key_from_index(index)
        menu = QMenu(self)
        if not key:
            menu.addAction("新建书籍", self._create_book)
            menu.addAction("导入为新书", lambda: self._import_book_file(None))
            menu.exec(self.library_tree.viewport().mapToGlobal(pos))
            return
        kind, node_id, group_book_id = key
        self.library_tree.setCurrentIndex(index)
        if kind == "book" and node_id is not None:
            menu.addAction("新建卷", lambda: self._create_volume(node_id))
            menu.addAction("新建章节", lambda: self._create_chapter(node_id, None))
            menu.addSeparator()
            item = self.tree_model.itemFromIndex(index)
            if item is not None:
                menu.addAction("展开全书", lambda item=item: self._set_tree_item_expanded_recursive(item, True))
                menu.addAction("折叠全书", lambda item=item: self._set_tree_item_expanded_recursive(item, False))
                menu.addSeparator()
            menu.addAction("书籍设置", lambda: self._open_book_settings(node_id))
            menu.addAction("重命名书籍", lambda: self._rename_book(node_id))
            menu.addAction("删除书籍", lambda: self._delete_book(node_id))
            menu.addSeparator()
            menu.addAction("导入到本书", lambda: self._import_book_file(node_id))
            export_menu = menu.addMenu("导出本书")
            export_menu.addAction("TXT", lambda: self._export_book("txt", node_id))
            export_menu.addAction("Markdown", lambda: self._export_book("md", node_id))
            export_menu.addAction("DOCX", lambda: self._export_book("docx", node_id))
            export_menu.addSeparator()
            export_menu.addAction("仅正文 TXT", lambda: self._export_book_part("content", node_id))
            export_menu.addAction("仅大纲 TXT", lambda: self._export_book_part("outline", node_id))
            export_menu.addAction("仅人物卡 TXT", lambda: self._export_book_part("characters", node_id))
            export_menu.addAction("仅世界观 TXT", lambda: self._export_book_part("world", node_id))
            export_menu.addAction("按卷分文件 TXT", lambda: self._export_book_part("per_volume", node_id))
        elif kind == "volume" and node_id is not None:
            volume = self.database.get_volume(node_id)
            if volume:
                book_id = int(volume["book_id"])
                menu.addAction("新建章节", lambda: self._create_chapter(book_id, node_id))
                menu.addSeparator()
                menu.addAction("重命名卷", lambda: self._rename_volume(node_id))
                menu.addAction("删除卷", lambda: self._delete_volume(node_id))
        elif kind == "group" and group_book_id is not None:
            menu.addAction("新建未分卷章节", lambda: self._create_chapter(group_book_id, None))
        elif kind == "chapter" and node_id is not None:
            menu.addAction("查看历史", lambda: self._open_history_dialog(node_id))
            menu.addSeparator()
            menu.addAction("导入到本章", lambda: self._import_chapter_file(node_id))
            export_menu = menu.addMenu("导出本章")
            export_menu.addAction("TXT", lambda: self._export_chapter("txt", node_id))
            export_menu.addAction("Markdown", lambda: self._export_chapter("md", node_id))
            export_menu.addAction("DOCX", lambda: self._export_chapter("docx", node_id))
            menu.addSeparator()
            if self.database.has_chapter_tag(node_id, "template"):
                menu.addAction("取消对照模板", lambda: self._toggle_chapter_template(node_id, False))
            else:
                menu.addAction("设为对照模板", lambda: self._toggle_chapter_template(node_id, True))
            menu.addSeparator()
            menu.addAction("重命名章节", lambda: self._rename_chapter(node_id))
            menu.addAction("移动到卷…", lambda: self._move_chapter_to_volume(node_id))
            menu.addAction("删除章节", lambda: self._delete_chapter(node_id))
        menu.exec(self.library_tree.viewport().mapToGlobal(pos))

    def _toggle_chapter_template(self, chapter_id: int, enabled: bool) -> None:
        chapter = self.database.get_chapter(chapter_id)
        if not chapter:
            return
        book_id = int(chapter["book_id"])
        if enabled and not self.database.has_chapter_tag(chapter_id, "template"):
            templates = self.database.list_template_chapters(book_id)
            if len(templates) >= 3:
                info(self, "对照模板", "同一本书最多保留 3 个对照模板章节，请先取消一个旧模板。")
                return
            self.database.set_chapter_tag(chapter_id, "template")
            self._set_status("已设为对照模板章节")
        else:
            self.database.remove_chapter_tag(chapter_id, "template")
            self._set_status("已取消对照模板章节")
        self._load_library_tree(self.current_tree_key)
        self._refresh_review_tab()

    def _create_contextual_item(self) -> None:
        if self.state.current_node_kind == "book" and self.state.selected_book_id:
            self._create_volume(self.state.selected_book_id)
        elif self.state.current_node_kind == "volume" and self.state.selected_book_id:
            self._create_chapter(self.state.selected_book_id, self.state.selected_volume_id)
        elif self.state.current_node_kind == "group" and self.state.selected_book_id:
            self._create_chapter(self.state.selected_book_id, None)
        elif self.state.current_node_kind == "chapter" and self.state.selected_book_id:
            self._create_chapter(self.state.selected_book_id, self.state.selected_volume_id)
        else:
            self._create_book()

    def _import_for_current_context(self) -> None:
        if self.state.current_node_kind == "chapter" and self.state.selected_chapter_id:
            self._import_chapter_file(self.state.selected_chapter_id)
            return
        if self.state.selected_book_id:
            self._import_book_file(self.state.selected_book_id)
            return
        self._import_book_file(None)

    def _import_long_text(self) -> None:
        self._import_book_file(None)

    def _import_book_file(self, book_id: int | None = None) -> None:
        if not self._confirm_editor_navigation():
            return
        target_book = self.database.get_book(book_id) if book_id else None
        if book_id and not target_book:
            info(self, "导入文件", "目标书籍不存在。")
            return
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "导入到书籍" if target_book else "导入为新书",
            "",
            "All Supported Files (*.txt *.md *.docx);;Text Files (*.txt *.md);;Word Documents (*.docx);;All Files (*.*)",
        )
        if not file_path:
            return
        source_path = Path(file_path)
        try:
            raw_text = self._read_any_file(source_path)
        except Exception as exc:  # noqa: BLE001
            error(self, "导入失败", f"无法读取文本文件：\n{exc}")
            return
        if not raw_text.strip():
            info(self, "导入文件", "这个文件没有可导入的正文内容。")
            return

        if target_book:
            title = str(target_book["title"])
        else:
            default_title = source_path.stem.strip() or "导入作品"
            title = ask_text(self, "导入为新书", "请输入导入后的书名：", default_title)
            if not title:
                return

        try:
            parsed = parse_long_text(title, raw_text)
        except Exception as exc:  # noqa: BLE001
            error(self, "导入失败", f"文本拆分失败：\n{exc}")
            return
        if not parsed.chapters:
            info(self, "导入文件", "没有识别到可导入的章节。")
            return

        volume_count = len({chapter.volume_title for chapter in parsed.chapters if chapter.volume_title})
        total_chars = sum(len(chapter.content.strip()) for chapter in parsed.chapters)
        if target_book:
            message = (
                f"将导入到《{target_book['title']}》。\n"
                f"识别到 {volume_count} 卷 / {len(parsed.chapters)} 章 / {total_chars} 字。\n\n"
                "导入内容会作为新卷/新章节追加到当前书籍中，是否继续？"
            )
        else:
            message = (
                f"识别到 {volume_count} 卷 / {len(parsed.chapters)} 章 / {total_chars} 字。\n\n"
                "将创建一本新书，并把每章正文写入章节正文区、自动提炼内容写入大纲。是否继续？"
            )
        if not confirm(self, "确认导入", message):
            return

        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            if target_book:
                first_chapter_id = self._append_parsed_to_book(int(target_book["id"]), parsed)
                imported_book_id = int(target_book["id"])
            else:
                imported_book_id, first_chapter_id = self._create_book_from_parsed(parsed)
        except Exception as exc:  # noqa: BLE001
            error(self, "导入失败", f"写入作品库失败：\n{exc}")
            return
        finally:
            QApplication.restoreOverrideCursor()

        key = ("chapter", first_chapter_id, None) if first_chapter_id else ("book", imported_book_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)
        self._set_status(f"已导入 {parsed.title}：{volume_count} 卷 / {len(parsed.chapters)} 章")
        info(self, "导入完成", f"已导入《{parsed.title}》。")

    def _read_text_file(self, path: Path) -> str:
        data = path.read_bytes()
        max_bytes = 50 * 1024 * 1024
        if len(data) > max_bytes:
            raise ValueError("文件超过 50MB，请先拆分后再导入。")
        if data.startswith((b"\xff\xfe", b"\xfe\xff")):
            return data.decode("utf-16")
        if data.count(b"\x00") > max(8, len(data) // 100):
            raise ValueError("文件看起来不是纯文本，请确认格式后再导入。")
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError as exc:
                last_error = exc
        raise ValueError(f"无法识别文本编码：{last_error}")

    def _read_docx_text(self, path: Path) -> str:
        doc = _DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)

    def _read_any_file(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".docx":
            if not HAS_DOCX:
                raise RuntimeError("当前环境缺少 python-docx，无法导入 DOCX。")
            return self._read_docx_text(path)
        return self._read_text_file(path)

    def _import_chapter_file(self, chapter_id: int | None = None) -> None:
        target_chapter_id = chapter_id or self.state.selected_chapter_id
        if not target_chapter_id:
            info(self, "导入章节", "请先选择一个章节。")
            return
        chapter = self.database.get_chapter(target_chapter_id)
        if not chapter:
            return
        if self.state.editor_dirty and self.state.selected_chapter_id == target_chapter_id:
            if not confirm(self, "确认导入章节", "当前章节有未保存内容，导入会替换正文。是否继续？"):
                return
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "导入到本章",
            "",
            "All Supported Files (*.txt *.md *.docx);;Text Files (*.txt *.md);;Word Documents (*.docx);;All Files (*.*)",
        )
        if not file_path:
            return
        source_path = Path(file_path)
        try:
            raw_text = self._read_any_file(source_path).strip()
        except Exception as exc:  # noqa: BLE001
            error(self, "导入失败", f"无法读取文件：\n{exc}")
            return
        if not raw_text:
            info(self, "导入章节", "文件内容为空。")
            return
        if not confirm(self, "确认导入章节", f"将用文件内容替换《{chapter['title']}》正文，原正文会保存为历史快照。是否继续？"):
            return

        label = f"导入覆盖前快照 · {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        self.database.create_snapshot(
            target_chapter_id,
            label,
            str(chapter["outline"] or ""),
            str(chapter["content"] or ""),
        )
        self.database.update_chapter(
            target_chapter_id,
            outline=str(chapter["outline"] or ""),
            content=raw_text,
            summary_text=str(chapter["summary_text"] or ""),
        )
        if self.state.selected_chapter_id == target_chapter_id:
            self.loading_editor = True
            try:
                self.chapter_outline_editor.setPlainText(str(chapter["outline"] or ""))
                self.content_editor.setPlainText(raw_text)
            finally:
                self.loading_editor = False
            self.loaded_chapter_content = raw_text
            self.loaded_chapter_summary_text = str(chapter["summary_text"] or "")
            self._mark_clean()
        self._set_status(f"已导入到章节：{chapter['title']}")

    def _create_book_from_parsed(self, parsed: ParsedBook) -> tuple[int, int | None]:
        book_id: int | None = None
        try:
            book_id = self.database.create_book(parsed.title)
            outline_data = merge_book_outline_supplement(
                self.database.get_book_outline_form(book_id),
                parsed.title,
                parsed.outline,
                supplement_title="导入补充",
            )
            self.database.update_book_outline_form(book_id, outline_data)

            volume_ids: dict[str, int] = {}
            volume_chapters: dict[str, list[Any]] = {}
            first_chapter_id: int | None = None
            for chapter in parsed.chapters:
                volume_id: int | None = None
                if chapter.volume_title:
                    volume_id = volume_ids.get(chapter.volume_title)
                    if volume_id is None:
                        volume_id = self.database.create_volume(book_id, chapter.volume_title)
                        volume_ids[chapter.volume_title] = volume_id
                    volume_chapters.setdefault(chapter.volume_title, []).append(chapter)

                chapter_id = self.database.create_chapter(book_id, chapter.title, volume_id)
                self.database.update_chapter(
                    chapter_id,
                    outline=chapter.outline,
                    content=chapter.content,
                    summary_text="",
                )
                if first_chapter_id is None:
                    first_chapter_id = chapter_id

            for volume_title, volume_id in volume_ids.items():
                outline = self._build_imported_volume_outline(volume_title, volume_chapters.get(volume_title, []))
                self.database.update_volume_outline(volume_id, outline)
            self._sync_truth_files(book_id)
            self._index_book_design_for_rag(book_id)
            return book_id, first_chapter_id
        except Exception:
            if book_id is not None:
                self.database.delete_book(book_id)
            raise

    def _append_parsed_to_book(self, book_id: int, parsed: ParsedBook) -> int | None:
        book = self.database.get_book(book_id)
        if not book:
            raise ValueError("Book not found.")

        existing_volumes = {str(volume["title"]): int(volume["id"]) for volume in self.database.list_volumes(book_id)}
        volume_chapters: dict[str, list[Any]] = {}
        touched_volume_ids: dict[str, int] = {}
        first_chapter_id: int | None = None

        for chapter in parsed.chapters:
            volume_id: int | None = None
            if chapter.volume_title:
                volume_id = existing_volumes.get(chapter.volume_title)
                if volume_id is None:
                    volume_id = self.database.create_volume(book_id, chapter.volume_title)
                    existing_volumes[chapter.volume_title] = volume_id
                touched_volume_ids[chapter.volume_title] = volume_id
                volume_chapters.setdefault(chapter.volume_title, []).append(chapter)

            chapter_id = self.database.create_chapter(book_id, chapter.title, volume_id)
            self.database.update_chapter(
                chapter_id,
                outline=chapter.outline,
                content=chapter.content,
                summary_text="",
            )
            if first_chapter_id is None:
                first_chapter_id = chapter_id

        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        outline_data = merge_book_outline_supplement(
            self.database.get_book_outline_form(book_id),
            str(book["title"] or ""),
            parsed.outline,
            supplement_title=f"导入补充 · {stamp}",
        )
        self.database.update_book_outline_form(book_id, outline_data)
        for volume_title, volume_id in touched_volume_ids.items():
            volume = self.database.get_volume(volume_id)
            if not volume:
                continue
            imported_outline = self._build_imported_volume_outline(volume_title, volume_chapters.get(volume_title, []))
            merged_volume_outline = self._merge_imported_outline(str(volume["outline_text"] or ""), imported_outline)
            self.database.update_volume_outline(volume_id, merged_volume_outline)
        self._sync_truth_files(book_id)
        self._index_book_design_for_rag(book_id)
        return first_chapter_id

    def _merge_imported_outline(self, existing: str, imported: str) -> str:
        imported = imported.strip()
        if not imported:
            return existing.strip()
        existing = existing.strip()
        if not existing:
            return imported
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        return f"{existing}\n\n【导入补充 · {stamp}】\n{imported}"

    def _build_imported_volume_outline(self, volume_title: str, chapters: list[Any]) -> str:
        total_chars = sum(len(getattr(chapter, "content", "").strip()) for chapter in chapters)
        lines = [
            f"自动导入卷大纲：{volume_title}",
            f"规模：{len(chapters)} 章 / {total_chars} 字",
            "",
            "章节结构：",
        ]
        for chapter in chapters[:40]:
            first_outline_line = getattr(chapter, "outline", "").splitlines()
            detail = first_outline_line[0] if first_outline_line else getattr(chapter, "title", "")
            lines.append(f"- {getattr(chapter, 'title', '')}：{detail}")
        if len(chapters) > 40:
            lines.append(f"- ... 另有 {len(chapters) - 40} 章")
        return "\n".join(lines).strip()

    def _smart_import(self) -> None:
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "智能导入文档",
            "",
            "All Supported Files (*.txt *.md *.docx);;Text Files (*.txt *.md);;Word Documents (*.docx);;All Files (*.*)",
        )
        if not file_path:
            return
        source_path = Path(file_path)
        try:
            raw_text = self._read_any_file(source_path)
        except Exception as exc:
            error(self, "导入失败", f"无法读取文件：\n{exc}")
            return
        if not raw_text.strip():
            info(self, "智能导入", "文件内容为空。")
            return

        ai_service = self._ai_service_for_purpose("import")
        if ai_service is None:
            info(self, "智能导入", "智能导入需要远程模型识别文档类型和章节结构。请先在 AI 设置中配置“智能导入分类”。")
            return
        self._open_drawer("tasks")

        def work() -> dict[str, Any]:
            doc_type = ai_service.analyze_document(raw_text[:3000], require_remote=True)
            normalized_type = str(doc_type).strip().lower()
            chapter_structure: list[dict[str, Any]] = []
            if normalized_type == "body":
                chapter_structure = ai_service.detect_chapter_structure(raw_text, require_remote=True)
            return {"doc_type": normalized_type, "chapter_structure": chapter_structure}

        def done(result: Any) -> None:
            if isinstance(result, dict):
                doc_type = str(result.get("doc_type", "body")).strip().lower()
                chapter_structure = result.get("chapter_structure", [])
            else:
                doc_type = str(result).strip().lower()
                chapter_structure = []
            self._dispatch_smart_import(doc_type, raw_text, source_path.stem, chapter_structure)

        def failed(message: str) -> None:
            self._append_task_log(f"智能导入失败：{message}")
            self._set_status("智能导入失败，请检查远程模型配置")
            error(self, "智能导入失败", f"远程模型未能完成文档分类或章节结构识别：\n{message}")

        self._run_background_ai_job(label="智能导入分类", callback=work, on_done=done, on_error=failed)

    def _dispatch_smart_import(
        self,
        doc_type: str,
        raw_text: str,
        default_title: str,
        chapter_structure: list[dict[str, Any]] | None = None,
    ) -> None:
        doc_type = doc_type if doc_type in {"body", "outline", "character", "world"} else "body"
        if doc_type == "outline":
            self._smart_import_outline(raw_text, default_title)
        elif doc_type == "character":
            self._smart_import_character(raw_text, default_title)
        elif doc_type == "world":
            self._smart_import_world(raw_text, default_title)
        else:
            self._smart_import_body(raw_text, default_title, chapter_structure)

    def _smart_import_body(
        self,
        raw_text: str,
        default_title: str,
        chapter_structure: list[dict[str, Any]] | None = None,
    ) -> None:
        if not self._confirm_editor_navigation():
            return
        title = ask_text(self, "智能导入 — 正文", "请输入导入后的书名：", default_title)
        if not title:
            return
        try:
            parsed = parse_long_text(
                title,
                raw_text,
                chapter_structure=chapter_structure,
                use_local_headings=False,
            )
        except Exception as exc:
            error(self, "智能导入失败", f"文本拆分失败：\n{exc}")
            return
        if not parsed.chapters:
            info(self, "智能导入", "没有识别到可导入的章节。")
            return
        volume_count = len({c.volume_title for c in parsed.chapters if c.volume_title})
        total_chars = sum(len(c.content.strip()) for c in parsed.chapters)
        message = f"识别为正文文档。\n识别到 {volume_count} 卷 / {len(parsed.chapters)} 章 / {total_chars} 字。\n\n是否导入为新书？"
        if not confirm(self, "确认导入", message):
            return
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            book_id, first_chapter_id = self._create_book_from_parsed(parsed)
        except Exception as exc:
            error(self, "智能导入失败", f"写入作品库失败：\n{exc}")
            return
        finally:
            QApplication.restoreOverrideCursor()
        key = ("chapter", first_chapter_id, None) if first_chapter_id else ("book", book_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)
        self._set_status(f"已智能导入正文：《{parsed.title}》")
        info(self, "导入完成", f"已导入《{parsed.title}》：{volume_count} 卷 / {len(parsed.chapters)} 章")

    def _smart_import_outline(self, raw_text: str, default_title: str) -> None:
        if not self.state.selected_book_id:
            info(self, "智能导入", "请先选择一本书，大纲将写入当前选中范围。")
            return
        message = f"识别为大纲文档（共 {len(raw_text)} 字符）。\n\n写入当前编辑区的大纲栏？"
        if not confirm(self, "确认导入大纲", message):
            return
        if self.state.editor_scope_kind == "chapter":
            self.chapter_outline_editor.setPlainText(raw_text.strip())
        elif self.state.editor_scope_kind == "book":
            volume_specs = self._book_outline_volume_specs(self.state.selected_book_id)
            data = merge_book_outline_supplement(
                self.book_outline_form.to_data(self.state.selected_book_title),
                self.state.selected_book_title,
                raw_text,
                supplement_title=f"导入补充 · {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            )
            self.book_outline_form.set_data(data, self.state.selected_book_title, volume_specs)
            self._on_book_outline_form_changed()
        else:
            self.outline_editor.setPlainText(raw_text.strip())
        self.state.editor_dirty = True
        self.save_label.setText("未保存")
        self._set_status(f"大纲已导入，请核对后保存。")

    def _smart_import_character(self, raw_text: str, default_name: str) -> None:
        if not self.state.selected_book_id:
            info(self, "智能导入", "请先选择一本书以创建人物。")
            return
        if not self._confirm_character_edit_navigation():
            return
        name = ask_text(self, "智能导入 — 人物", "请输入人物名称：", default_name)
        if not name:
            return
        char_id = self.database.create_character(self.state.selected_book_id, name)
        profile_data = merge_character_profile_supplement(
            self.database.get_character_profile_form(char_id),
            name,
            "",
            raw_text.strip(),
            supplement_title=f"智能导入补充 · {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        )
        self.database.update_character_form(char_id, name, "", profile_data)
        self.selected_character_id = char_id
        self._sync_truth_files(self.state.selected_book_id)
        self._index_book_design_for_rag(self.state.selected_book_id)
        self._refresh_drawer()
        self._open_drawer("details")
        self._set_drawer_tab("characters")
        self._set_status(f"人物「{name}」已创建并导入设定。")
        info(self, "导入完成", f"人物「{name}」已创建。请在详情抽屉中核对设定内容。")

    def _smart_import_world(self, raw_text: str, default_name: str) -> None:
        if not self.state.selected_book_id:
            info(self, "智能导入", "请先选择一本书以导入世界观。")
            return
        base_data = (
            self.world_form.to_data()
            if hasattr(self, "world_form") and getattr(self, "_world_form_book_id", None) == self.state.selected_book_id
            else self.database.get_book_world_form(self.state.selected_book_id)
        )
        world_data = merge_world_supplement(
            base_data,
            raw_text,
            supplement_title=f"智能导入补充 · {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        )
        self.database.update_book_world_form(self.state.selected_book_id, world_data)
        self._sync_truth_files(self.state.selected_book_id)
        self._index_book_design_for_rag(self.state.selected_book_id)
        if hasattr(self, "world_form"):
            self.world_form.set_data(world_data)
            self._world_form_book_id = self.state.selected_book_id
            self._refresh_world_markdown_preview()
            self._world_form_dirty = False
        else:
            self._load_world_form_for_book(self.state.selected_book_id)
        self._refresh_drawer()
        self._open_drawer("details")
        self._set_drawer_tab("world")
        self._set_status("世界观内容已导入统一模板。")
        info(self, "导入完成", "世界观内容已加入统一模板补充区。请在详情抽屉中核对内容。")

    def _create_book(self) -> None:
        title = ask_text(self, "新建书籍", "请输入书名：")
        if not title:
            return
        book_id = self.database.create_book(title)
        self._sync_truth_files(book_id)
        self._index_book_design_for_rag(book_id)
        key = ("book", book_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _rename_book(self, book_id: int | None = None) -> None:
        book_id = book_id or self.state.selected_book_id
        if not book_id:
            return
        book = self.database.get_book(book_id)
        if not book:
            return
        title = ask_text(self, "重命名书籍", "请输入新书名：", str(book["title"]))
        if not title:
            return
        self.database.rename_book(book_id, title)
        key = ("book", book_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _delete_book(self, book_id: int | None = None) -> None:
        book_id = book_id or self.state.selected_book_id
        if not book_id:
            return
        book = self.database.get_book(book_id)
        if not book:
            return
        if not confirm(self, "删除书籍", f"确定删除《{book['title']}》及其所有内容吗？", danger=True):
            return
        self.database.delete_book(book_id)
        self.state.clear_selection()
        self.current_tree_key = None
        self._load_library_tree()
        self.stack.setCurrentWidget(self.empty_page)
        self._refresh_drawer()
        self._set_status("已删除书籍")

    def _create_volume(self, book_id: int | None = None) -> None:
        book_id = book_id or self.state.selected_book_id
        if not book_id:
            info(self, "提示", "请先选择一本书。")
            return
        title = ask_text(self, "新建卷", "请输入卷名：")
        if not title:
            return
        volume_id = self.database.create_volume(book_id, title)
        self._sync_truth_files(book_id)
        self._index_book_design_for_rag(book_id)
        key = ("volume", volume_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _rename_volume(self, volume_id: int) -> None:
        volume = self.database.get_volume(volume_id)
        if not volume:
            return
        title = ask_text(self, "重命名卷", "请输入新卷名：", str(volume["title"]))
        if not title:
            return
        book_id = int(volume["book_id"])
        self.database.rename_volume(volume_id, title)
        self._sync_truth_files(book_id)
        self._index_book_design_for_rag(book_id)
        key = ("volume", volume_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _delete_volume(self, volume_id: int) -> None:
        volume = self.database.get_volume(volume_id)
        if not volume:
            return
        if not confirm(self, "删除卷", f"确定删除卷《{volume['title']}》吗？章节会移入未分卷。", danger=True):
            return
        book_id = int(volume["book_id"])
        self.database.delete_volume(volume_id)
        self._sync_truth_files(book_id)
        self._index_book_design_for_rag(book_id)
        key = ("book", book_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _move_volume(self, volume_id: int, direction: int) -> None:
        volume = self.database.get_volume(volume_id)
        book_id = int(volume["book_id"]) if volume else None
        if self.database.move_volume(volume_id, direction):
            if book_id:
                self._sync_truth_files(book_id)
                self._index_book_design_for_rag(book_id)
            key = ("volume", volume_id, None)
            self._load_library_tree(key)
            self._activate_tree_key(key)

    def _create_chapter(self, book_id: int | None = None, volume_id: int | None = None) -> None:
        book_id = book_id or self.state.selected_book_id
        if not book_id:
            info(self, "提示", "请先选择一本书。")
            return
        title = ask_text(self, "新建章节", "请输入章节名：")
        if not title:
            return
        chapter_id = self.database.create_chapter(book_id, title, volume_id)
        key = ("chapter", chapter_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _rename_chapter(self, chapter_id: int) -> None:
        chapter = self.database.get_chapter(chapter_id)
        if not chapter:
            return
        title = ask_text(self, "重命名章节", "请输入新章节名：", str(chapter["title"]))
        if not title:
            return
        self.database.rename_chapter(chapter_id, title)
        key = ("chapter", chapter_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _delete_chapter(self, chapter_id: int) -> None:
        chapter = self.database.get_chapter(chapter_id)
        if not chapter:
            return
        if not confirm(self, "删除章节", f"确定删除章节《{chapter['title']}》吗？", danger=True):
            return
        book_id = int(chapter["book_id"])
        volume_id = chapter["volume_id"]
        key = ("group", None, book_id) if volume_id is None else ("volume", int(volume_id), None)
        self.database.delete_chapter(chapter_id)
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _move_chapter(self, chapter_id: int, direction: int) -> None:
        if self.database.move_chapter(chapter_id, direction):
            key = ("chapter", chapter_id, None)
            self._load_library_tree(key)
            self._activate_tree_key(key)

    def _move_chapter_to_volume(self, chapter_id: int) -> None:
        chapter = self.database.get_chapter(chapter_id)
        if not chapter:
            return
        book_id = int(chapter["book_id"])
        volumes = self.database.list_volumes(book_id)
        current_volume_id = chapter["volume_id"]
        volume_labels = []
        volume_ids = []
        for idx, vol in enumerate(volumes, start=1):
            label = f"第{idx}卷 · {vol['title']}"
            volume_labels.append(label)
            volume_ids.append(int(vol["id"]))
        if not volume_labels:
            info(self, "提示", "当前书籍没有卷，请先创建卷。")
            return
        dialog = QInputDialog(self)
        dialog.setWindowTitle("移动到卷")
        dialog.setLabelText("选择目标卷：")
        dialog.setComboBoxItems(volume_labels)
        dialog.setComboBoxEditable(False)
        if not dialog.exec():
            return
        selected_label = dialog.textValue()
        if not selected_label:
            return
        try:
            idx = volume_labels.index(selected_label)
            target_volume_id = volume_ids[idx]
        except ValueError:
            return
        if target_volume_id == current_volume_id:
            return
        self.database.update_chapter_volume(chapter_id, target_volume_id)
        key = ("chapter", chapter_id, None)
        self._load_library_tree(key)
        self._activate_tree_key(key)
        self._set_status("章节已移动到目标卷")

    def _refresh_drawer(self) -> None:
        return self.drawer_controller.refresh()

    def _on_drawer_tab_changed(self) -> None:
        return self.drawer_controller.on_tab_changed()

    def _refresh_visible_drawer_tab(self) -> None:
        return self.drawer_controller.refresh_visible_tab()

    def _is_book_level_selected(self) -> bool:
        return self.drawer_controller.is_book_level_selected()

    def _sync_overview_tab_visibility(self) -> None:
        return self.drawer_controller.sync_overview_tab_visibility()

    def _sync_shortcuts_tab_visibility(self) -> None:
        return self.drawer_controller.sync_shortcuts_tab_visibility()

    def _sync_outline_tab_for_level(self) -> None:
        return self.drawer_controller.sync_outline_tab_for_level()

    def _set_tab_visible(self, index: int, visible: bool, *, fallback_key: str = "outline") -> None:
        return self.drawer_controller.set_tab_visible(index, visible, fallback_key=fallback_key)

    @staticmethod
    def _is_named_character_record(character: dict[str, Any]) -> bool:
        name = str(character.get("name", "") or "").strip()
        if not name:
            return False
        compact = re.sub(r"\s+", "", name).casefold()
        generic_names = {
            "未命名",
            "未命名人物",
            "新人物",
            "人物",
            "角色",
            "主角",
            "配角",
            "反派",
            "无名",
            "匿名",
            "unknown",
        }
        if compact in generic_names:
            return False
        return re.fullmatch(r"(人物|角色|character|char)\d*", compact) is None

    def _named_book_characters(self, book_id: int) -> list[dict[str, Any]]:
        characters: list[dict[str, Any]] = []
        for row in self.database.list_characters(book_id):
            record = dict(row)
            detail = self.database.get_character(int(record["id"]))
            if detail:
                record.update(dict(detail))
            if self._is_named_character_record(record):
                characters.append(record)
        return characters

    def _refresh_generation_outline_context(self) -> None:
        if not hasattr(self, "generation_book_outline_editor"):
            return
        book_outline = ""
        if self.state.selected_book_id:
            book = self.database.get_book(self.state.selected_book_id)
            if book:
                book_outline = str(book["outline_text"] or "")
        self.generation_book_outline_editor.blockSignals(True)
        self.generation_book_outline_editor.setPlainText(book_outline)
        self.generation_book_outline_editor.blockSignals(False)

    def _save_generation_book_outline(self) -> None:
        if not self.state.selected_book_id or not hasattr(self, "generation_book_outline_editor"):
            info(self, "提示", "请先选择一本书。")
            return
        outline = coerce_book_outline_markdown(
            self.generation_book_outline_editor.toPlainText(),
            self.state.selected_book_title,
        )
        self.database.update_book_outline(self.state.selected_book_id, outline)
        book = self.database.get_book(self.state.selected_book_id)
        if book:
            outline = str(book["outline_text"] or outline)
        self._sync_truth_files(self.state.selected_book_id)
        self._index_book_design_for_rag(self.state.selected_book_id)
        if self.state.editor_scope_kind == "book":
            volume_specs = self._book_outline_volume_specs(self.state.selected_book_id)
            outline_data = self.database.get_book_outline_form(self.state.selected_book_id)
            self.loading_editor = True
            try:
                self.book_outline_form.set_data(outline_data, self.state.selected_book_title, volume_specs)
                self.outline_editor.setPlainText(outline)
            finally:
                self.loading_editor = False
            self._refresh_outline_markdown_preview()
            self._sync_editor_text_lengths()
            self._mark_clean()
        self.generation_book_outline_editor.setPlainText(outline)
        self._refresh_overview()
        self._set_status("书籍全局描述已保存。")

    def _on_outline_action_btn(self) -> None:
        kind = self.state.current_node_kind
        if kind == "volume":
            self._save_volume_outline()
        else:
            self._start_multi_agent_generation()

    def _save_volume_outline(self) -> None:
        if not self.state.selected_volume_id:
            info(self, "提示", "请先选择一个卷。")
            return
        if not hasattr(self, "volume_outline_editor"):
            return
        self.database.update_volume_outline(
            self.state.selected_volume_id,
            self.volume_outline_editor.toPlainText(),
        )
        self._set_status("卷大纲已保存。")

    def _start_ai_book_summary(self) -> None:
        book_id = self.state.selected_book_id
        if not book_id:
            info(self, "提示", "请先选择一本书。")
            return
        ai_service = self._require_ai_service("book_analysis", "AI 总结全书")
        if ai_service is None:
            return
        book = self.database.get_book(book_id)
        if not book:
            return
        chapters = [dict(item) for item in self.database.get_book_export_data(book_id).get("chapters", [])]
        if not chapters:
            info(self, "AI 总结全书", "当前书籍还没有章节。")
            return
        book_title = str(book["title"])
        book_outline = str(book["outline_text"] or "")
        characters = [dict(item) for item in self.database.list_characters(book_id)]
        world_entries = [dict(item) for item in self.database.list_world_entries(book_id)]

        if len(characters) > 20:
            dialog = QDialog(self)
            dialog.setWindowTitle("选择分析人物")
            dialog.resize(400, 500)
            layout = QVBoxLayout(dialog)
            layout.addWidget(QLabel(f"全书分析共有 {len(characters)} 个人物，请选择要分析的人物："))
            list_widget = QListWidget()
            list_widget.setSelectionMode(QAbstractItemView.SelectionMode.MultiSelection)
            all_items = []
            for ch in characters:
                item = QListWidgetItem(f"{ch.get('name', '')}  [{ch.get('role', '未设定')}]")
                item.setData(Qt.ItemDataRole.UserRole, ch)
                item.setSelected(True)
                list_widget.addItem(item)
                all_items.append(item)
            layout.addWidget(list_widget)
            btn_layout = QHBoxLayout()
            select_all = QPushButton("全选")
            deselect_all = QPushButton("反选")
            ok_btn = QPushButton("确定")
            cancel_btn = QPushButton("取消")
            select_all.clicked.connect(lambda: [it.setSelected(True) for it in all_items])
            deselect_all.clicked.connect(lambda: [it.setSelected(not it.isSelected()) for it in all_items])
            btn_layout.addWidget(select_all)
            btn_layout.addWidget(deselect_all)
            btn_layout.addStretch(1)
            btn_layout.addWidget(ok_btn)
            btn_layout.addWidget(cancel_btn)
            layout.addLayout(btn_layout)
            ok_btn.clicked.connect(dialog.accept)
            cancel_btn.clicked.connect(dialog.reject)
            if not dialog.exec():
                return
            characters = [it.data(Qt.ItemDataRole.UserRole) for it in all_items if it.isSelected()]
            if not characters:
                return
        self._open_drawer("tasks")
        self._append_task_log(f"AI 总结全书开始：《{book_title}》，{len(chapters)} 个章节。")

        def work(log: Callable[[str], None], cancel_event: Event | None) -> dict[str, Any]:
            log("正在分析全书结构、人物、世界观与情节脉络...")
            return self._analyze_book_in_batches(
                ai_service=ai_service,
                book_title=book_title,
                book_outline=book_outline,
                chapters=chapters,
                characters=characters,
                world_entries=world_entries,
                log=log,
                cancel_event=cancel_event,
            )

        def done(payload: Any) -> None:
            if not isinstance(payload, dict):
                return
            self._apply_book_analysis_result(book_id, payload)

        self._run_background_ai_job(
            label="AI 总结全书",
            callback=work,
            on_done=done,
            callback_uses_log=True,
            callback_uses_cancel=True,
        )

    def _refresh_shortcuts_list(self) -> None:
        if not hasattr(self, "shortcuts_list"):
            return
        self.shortcuts_list.clear()
        if not self.state.selected_book_id:
            return
        rows = self.database.list_name_shortcuts(self.state.selected_book_id)
        for row in rows:
            item = QListWidgetItem(f"{row['key_seq']}  →  「{row['character_name']}」")
            item.setData(Qt.ItemDataRole.UserRole, int(row["id"]))
            self.shortcuts_list.addItem(item)
        self._register_name_shortcuts()

    def _register_name_shortcuts(self) -> None:
        self._unregister_name_shortcuts()
        if not self.state.selected_book_id:
            return
        rows = self.database.list_name_shortcuts(self.state.selected_book_id)
        for row in rows:
            key_seq = str(row["key_seq"]).strip()
            character_name = str(row["character_name"]).strip()
            if not key_seq or not character_name:
                continue
            try:
                shortcut = QShortcut(QKeySequence(key_seq), self)
                shortcut.activated.connect(self._make_insert_name_handler(character_name))
                self._cs_name_shortcut_objects.append(shortcut)
            except Exception:
                pass

    def _unregister_name_shortcuts(self) -> None:
        if not hasattr(self, "_cs_name_shortcut_objects"):
            self._cs_name_shortcut_objects: list[QShortcut] = []
            return
        for sc in self._cs_name_shortcut_objects:
            sc.setEnabled(False)
            sc.deleteLater()
        self._cs_name_shortcut_objects.clear()

    def _make_insert_name_handler(self, name: str):
        def handler() -> None:
            editor = getattr(self, "content_editor", None)
            if editor is None or not editor.isVisible():
                return
            cursor = editor.textCursor()
            cursor.insertText(name)
            editor.setTextCursor(cursor)
            editor.setFocus()
        return handler

    def _add_name_shortcut(self) -> None:
        if not self.state.selected_book_id:
            info(self, "提示", "请先选择一本书。")
            return
        key_seq = self.shortcut_key_input.text().strip()
        character_name = self.shortcut_name_input.text().strip()
        if not key_seq or not character_name:
            info(self, "提示", "请输入快捷键和人名。")
            return
        self.database.add_name_shortcut(self.state.selected_book_id, key_seq, character_name)
        self.shortcut_key_input.clear()
        self.shortcut_name_input.clear()
        self._refresh_shortcuts_list()

    def _delete_name_shortcut(self) -> None:
        current = self.shortcuts_list.currentItem()
        if not current:
            info(self, "提示", "请先选择一条快捷键。")
            return
        shortcut_id = current.data(Qt.ItemDataRole.UserRole)
        if shortcut_id is None:
            return
        self.database.delete_name_shortcut(int(shortcut_id))
        self._refresh_shortcuts_list()

    def _pick_characters_for_shortcuts(self) -> None:
        if not self.state.selected_book_id:
            info(self, "提示", "请先选择一本书。")
            return
        characters = [dict(item) for item in self.database.list_characters(self.state.selected_book_id)]
        if not characters:
            info(self, "提示", "当前书籍还没有人物，请先在人物面板创建。")
            return
        names = [str(ch.get("name", "")).strip() for ch in characters if str(ch.get("name", "")).strip()]
        if not names:
            return
        name, ok = QInputDialog.getItem(self, "选取人物", "选择要添加快捷键的人物：", names, 0, False)
        if not ok or not name:
            return
        key_seq, ok2 = QInputDialog.getText(self, "设置快捷键", f"为「{name}」设置快捷键（如 Ctrl+D）：", text="")
        if not ok2 or not key_seq.strip():
            return
        self.database.add_name_shortcut(self.state.selected_book_id, key_seq.strip(), name)
        self._refresh_shortcuts_list()

    def _refresh_overview(self) -> None:
        return drawer_pages.refresh_overview(self)

    def _refresh_characters(self) -> None:
        return drawer_pages.refresh_characters(self)

    def _refresh_character_list(self) -> None:
        return drawer_pages.refresh_character_list(self)

    def _select_character_list_row(self, character_id: int | None) -> None:
        return drawer_pages.select_character_list_row(self, character_id)

    def _open_character_from_list_item(self, item: QListWidgetItem) -> None:
        return drawer_pages.open_character_from_list_item(self, item)

    def _open_character_detail_dialog(self, character_id: int) -> None:
        return drawer_pages.open_character_detail_dialog(self, character_id)

    def _refresh_character_picker_label(self) -> None:
        return drawer_pages.refresh_character_picker_label(self)

    def _show_character_picker_menu(self) -> None:
        return drawer_pages.show_character_picker_menu(self)

    def _select_character_from_picker(self, character_id: int) -> None:
        return drawer_pages.select_character_from_picker(self, character_id)

    def _refresh_star_graph(self) -> None:
        if not hasattr(self, "star_graph") or not self.state.selected_book_id:
            return
        if not self._is_book_level_selected():
            self.star_graph.load_book([], [], lambda stored: self.database.resolve_media_path(stored))
            return
        if not hasattr(self, "star_graph_style_manager"):
            from novel_app.qt.star_graph_style import StyleManager

            self.star_graph_style_manager = StyleManager(self.database, self.state.selected_book_id)
            self.star_graph_style_manager.initialize_defaults()
            self.star_graph.set_style_manager(self.star_graph_style_manager)
        characters = self._named_book_characters(self.state.selected_book_id)
        character_ids = {int(item["id"]) for item in characters}
        relationships = [
            dict(item)
            for item in self.database.list_relationships(self.state.selected_book_id)
            if int(item["source_character_id"]) in character_ids and int(item["target_character_id"]) in character_ids
        ]
        self.star_graph.load_book(
            characters,
            relationships,
            lambda stored: self.database.resolve_media_path(stored),
        )
        self._apply_styles_to_graph()

    def _apply_styles_to_graph(self) -> None:
        if not hasattr(self, "star_graph_style_manager") or not hasattr(self.star_graph, "view"):
            return
        for character_id, node in self.star_graph.view.nodes.items():
            style_id = self.database.get_character_node_style(character_id)
            if style_id:
                style = self.star_graph_style_manager.get_node_style(style_id)
                if style:
                    node.apply_style(style.id, style.background, style.border_color, style.text_color, style.border_width)
            else:
                default_style = self.star_graph_style_manager.get_default_node_style()
                node.apply_style(default_style.id, default_style.background, default_style.border_color, default_style.text_color, default_style.border_width)

    def _create_character_from_graph(self, graph_x: float, graph_y: float) -> None:
        if not self.state.selected_book_id:
            return
        name = ask_text(self, "新建人物", "请输入人物名称：")
        if not name:
            return
        character_id = self.database.create_character(self.state.selected_book_id, name)
        self.database.update_character_position(character_id, graph_x, graph_y)
        self._refresh_drawer()

    def _move_character_in_graph(self, character_id: int, graph_x: float, graph_y: float) -> None:
        self.database.update_character_position(character_id, graph_x, graph_y)

    def _create_relationship_from_graph(self, source_character_id: int, target_character_id: int) -> None:
        if not self.state.selected_book_id or source_character_id == target_character_id:
            return
        relationship_type = ask_relationship_type(self, "关系")
        if not relationship_type:
            return
        self.database.create_relationship(
            self.state.selected_book_id,
            source_character_id,
            target_character_id,
            relationship_type,
        )
        self._refresh_star_graph()

    def _edit_character_from_graph(self, character_id: int) -> None:
        self._open_drawer("characters")
        self._set_drawer_tab("characters")
        self._select_character_from_picker(character_id)

    def _delete_character_from_graph(self, character_id: int) -> None:
        character = self.database.get_character(character_id)
        if not character:
            return
        if not confirm(self, "删除人物", f"确定删除人物「{character['name']}」吗？相关关系也会删除。", danger=True):
            return
        self.database.delete_character(character_id)
        if self.selected_character_id == character_id:
            self.selected_character_id = None
        self._refresh_drawer()

    def _edit_relationship_from_graph(self, relationship_id: int) -> None:
        relationship = self.database.get_relationship(relationship_id)
        if not relationship:
            return
        value = ask_relationship_type(self, str(relationship["relationship_type"] or "关系"))
        if not value:
            return
        self.database.update_relationship(relationship_id, value, str(relationship["description"] or ""))
        self._refresh_star_graph()

    def _delete_relationship_from_graph(self, relationship_id: int) -> None:
        relationship = self.database.get_relationship(relationship_id)
        if not relationship:
            return
        if not confirm(self, "删除关系", f"确定删除「{relationship['source_name']} → {relationship['target_name']}」吗？", danger=True):
            return
        self.database.delete_relationship(relationship_id)
        self._refresh_star_graph()

    def _change_character_style_from_graph(self, character_id: int) -> None:
        from novel_app.qt.star_graph_editors import NodeStyleSelectorDialog

        if not hasattr(self, "star_graph_style_manager"):
            return
        dialog = NodeStyleSelectorDialog(self, self.star_graph_style_manager)
        if dialog.exec() == dialog.Accepted and dialog.selected_style_id is not None:
            style = self.star_graph_style_manager.get_node_style(dialog.selected_style_id)
            if style:
                self.database.set_character_node_style(character_id, style.id)
                node = self.star_graph.view.nodes.get(character_id)
                if node:
                    node.apply_style(style.id, style.background, style.border_color, style.text_color, style.border_width)
                self._refresh_star_graph()

    def _change_relationship_type_from_graph(self, relationship_id: int) -> None:
        from novel_app.qt.star_graph_editors import RelationshipTypeSelectorDialog

        if not hasattr(self, "star_graph_style_manager"):
            return
        relationship = self.database.get_relationship(relationship_id)
        current_type_id = None
        relationship_data = dict(relationship) if relationship else {}
        if relationship_data.get("relationship_type_id"):
            current_type_id = int(relationship_data["relationship_type_id"])
        dialog = RelationshipTypeSelectorDialog(self, self.star_graph_style_manager, current_type_id)
        if dialog.exec() == dialog.Accepted and dialog.selected_type_id is not None:
            self.database.update_relationship_type_info(relationship_id, dialog.selected_type_id, dialog.description)
            self._refresh_star_graph()

    def _jump_to_chapter_from_graph(self, chapter_id: int) -> None:
        key = ("chapter", chapter_id, None)
        if not self._confirm_editor_navigation() or not self._confirm_side_form_navigation():
            return
        self._load_library_tree(key)
        self._activate_tree_key(key)

    def _refresh_world(self) -> None:
        return drawer_pages.refresh_world(self)

    def _refresh_skills(self) -> None:
        return drawer_pages.refresh_skills(self)

    def _on_enabled_skill_selected(self, current: QListWidgetItem | None, _previous: QListWidgetItem | None) -> None:
        return drawer_pages.on_enabled_skill_selected(self, current, _previous)

    def _refresh_review_tab(self) -> None:
        return drawer_pages.refresh_review_tab(self)

    def _on_review_run_selected(self, current: QListWidgetItem | None, _previous: QListWidgetItem | None) -> None:
        return drawer_pages.on_review_run_selected(self, current, _previous)

    def _render_review_run_detail(self, run_id: int) -> None:
        return drawer_pages.render_review_run_detail(self, run_id)

    def _selected_review_run(self) -> dict[str, Any] | None:
        return drawer_pages.selected_review_run(self)

    def _view_selected_review_snapshot(self) -> None:
        return drawer_pages.view_selected_review_snapshot(self)

    def _restore_selected_review_snapshot(self) -> None:
        return drawer_pages.restore_selected_review_snapshot(self)

    def _copy_selected_review_truth_file(self) -> None:
        return drawer_pages.copy_selected_review_truth_file(self)

    def _load_character_form(self, character_id: int) -> None:
        self._ensure_character_detail_dialog()
        detail = self.database.get_character(character_id)
        if not detail:
            return
        name = str(detail["name"])
        role = str(detail["role"] or "")
        profile_data = self.database.get_character_profile_form(character_id)
        self.selected_character_id = character_id
        self.loading_editor = True
        self.character_name.setText(name)
        self.character_role.setText(role)
        self.character_profile_form.set_data(profile_data, name, role)
        self.loading_editor = False
        self.character_image_path = str(detail["image_path"] or "")
        self._refresh_character_image_preview()
        self._refresh_character_markdown_preview()
        self._refresh_character_picker_label()
        self._character_form_dirty = False

    def _clear_character_form(self) -> None:
        self._ensure_character_detail_dialog()
        self.selected_character_id = None
        self.character_name.clear()
        self.character_role.clear()
        if hasattr(self, "character_profile_form"):
            self.character_profile_form.set_data(character_profile_form_defaults("", ""), "", "")
        self.character_image_path = ""
        self._refresh_character_image_preview()
        self._refresh_character_markdown_preview()
        self._refresh_character_picker_label()
        self._character_form_dirty = False

    def _confirm_character_edit_navigation(self) -> bool:
        if not getattr(self, "_character_form_dirty", False):
            return True
        if not self.selected_character_id:
            return True
        choice = self._ask_save_discard_cancel("未保存的人物修改", "当前人物有未保存的修改，切换前要保存吗？")
        if choice == "cancel":
            return False
        if choice == "save":
            return self._save_character(refresh=False)
        self._character_form_dirty = False
        return True

    def _mark_character_form_dirty(self) -> None:
        if getattr(self, "loading_editor", False):
            return
        self._character_form_dirty = True

    def _refresh_book_cover_preview(self) -> None:
        if not hasattr(self, "book_cover_preview"):
            return
        path = self.database.resolve_media_path(self.current_book_cover_path)
        set_adaptive_image(
            self.book_cover_preview,
            path,
            placeholder="封面",
            max_width=170,
            max_height=240,
            crop=False,
        )

    def _refresh_character_image_preview(self) -> None:
        if not hasattr(self, "character_image_preview"):
            return
        path = self.database.resolve_media_path(self.character_image_path)
        set_adaptive_image(
            self.character_image_preview,
            path,
            placeholder="立绘",
            max_width=170,
            max_height=240,
            crop=False,
        )

    def _choose_book_cover(self) -> None:
        if not self.state.selected_book_id:
            return
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择书籍封面",
            "",
            "Image Files (*.png *.jpg *.jpeg *.bmp *.webp);;All Files (*.*)",
        )
        if not path:
            return
        source = Path(path)
        dest = self.database.book_covers_dir / f"book_{self.state.selected_book_id}_{uuid.uuid4().hex[:8]}{source.suffix.lower()}"
        shutil.copy2(source, dest)
        self.current_book_cover_path = self.database.normalize_media_path(dest)
        self.database.update_book_cover(self.state.selected_book_id, self.current_book_cover_path)
        self._refresh_book_cover_preview()
        self._set_status("已更新书籍封面")

    def _clear_book_cover(self) -> None:
        if not self.state.selected_book_id:
            return
        self.current_book_cover_path = ""
        self.database.update_book_cover(self.state.selected_book_id, "")
        self._refresh_book_cover_preview()
        self._set_status("已清除书籍封面")

    def _choose_character_image(self) -> None:
        if not self.selected_character_id:
            return
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择人物立绘",
            "",
            "Image Files (*.png *.jpg *.jpeg *.bmp *.webp);;All Files (*.*)",
        )
        if not path:
            return
        source = Path(path)
        dest = self.database.character_images_dir / f"character_{self.selected_character_id}_{uuid.uuid4().hex[:8]}{source.suffix.lower()}"
        shutil.copy2(source, dest)
        self.character_image_path = self.database.normalize_media_path(dest)
        self.database.update_character_image(self.selected_character_id, self.character_image_path)
        self._refresh_character_image_preview()

    def _clear_character_image(self) -> None:
        if not self.selected_character_id:
            return
        self.character_image_path = ""
        self.database.update_character_image(self.selected_character_id, "")
        self._refresh_character_image_preview()

    def _create_character(self) -> None:
        if not self.state.selected_book_id:
            info(self, "提示", "请先选择一本书。")
            return
        name = ask_text(self, "新建人物", "请输入人物名称：")
        if not name:
            return
        self.selected_character_id = self.database.create_character(self.state.selected_book_id, name)
        self._character_form_dirty = False
        self._refresh_drawer()
        self._open_character_detail_dialog(self.selected_character_id)

    def _save_character(self, _checked: bool = False, *, refresh: bool = True) -> bool:
        if not self.selected_character_id:
            info(self, "提示", "请先选择或创建一个人物。")
            return False
        name = self.character_name.text().strip() or "未命名人物"
        role = self.character_role.text().strip()
        profile_data = self.character_profile_form.to_data(name, role)
        self.database.update_character_form(
            self.selected_character_id,
            name,
            role,
            profile_data,
            image_path=self.database.normalize_media_path(self.character_image_path),
        )
        self.character_name.setText(name)
        self.character_role.setText(role)
        self.character_profile_form.set_data(profile_data, name, role)
        self._refresh_character_markdown_preview()
        self._character_form_dirty = False
        if self.state.selected_book_id:
            self._sync_truth_files(self.state.selected_book_id)
            self._index_book_design_for_rag(self.state.selected_book_id)
        self._set_status("人物设定已保存")
        if refresh:
            self._refresh_drawer()
        else:
            self._refresh_character_list()
        return True

    def _import_character_file(self) -> None:
        if not self.selected_character_id:
            info(self, "提示", "请先选择或创建一个人物。")
            return
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "导入人物设定文件",
            "",
            "All Supported Files (*.txt *.md *.docx);;Text Files (*.txt *.md);;Word Documents (*.docx);;All Files (*.*)",
        )
        if not file_path:
            return
        try:
            text = self._read_any_file(Path(file_path))
        except Exception as exc:
            error(self, "导入失败", f"无法读取文件：\n{exc}")
            return
        if not text.strip():
            info(self, "导入文件", "文件内容为空。")
            return
        data = merge_character_profile_supplement(
            self.character_profile_form.to_data(self.character_name.text(), self.character_role.text()),
            self.character_name.text(),
            self.character_role.text(),
            text.strip(),
            supplement_title=f"导入补充 · {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        )
        self.character_profile_form.set_data(data, self.character_name.text(), self.character_role.text())
        self._refresh_character_markdown_preview()
        self._mark_character_form_dirty()
        self._set_status("人物设定文件已载入补充区，请核对后保存。")

    def _delete_character(self) -> None:
        if not self.selected_character_id:
            return
        if confirm(self, "删除人物", "确定删除当前人物吗？", danger=True):
            self.database.delete_character(self.selected_character_id)
            self.selected_character_id = None
            if self.character_detail_dialog:
                self.character_detail_dialog.close()
            self._refresh_drawer()

    def _confirm_world_edit_navigation(self) -> bool:
        if not getattr(self, "_world_form_dirty", False):
            return True
        if not self.state.selected_book_id:
            return True
        choice = self._ask_save_discard_cancel("未保存的世界观修改", "当前世界观设定有未保存的修改，切换前要保存吗？")
        if choice == "cancel":
            return False
        if choice == "save":
            return self._save_world_entry(refresh=False)
        self._world_form_dirty = False
        return True

    def _mark_world_form_dirty(self) -> None:
        self._world_form_dirty = True

    def _create_world_entry(self) -> None:
        self._load_world_form_for_book(self.state.selected_book_id)

    def _save_world_entry(self, _checked: bool = False, *, refresh: bool = True) -> bool:
        if not self.state.selected_book_id:
            info(self, "提示", "请先选择一本书。")
            return False
        if not hasattr(self, "world_form"):
            return False
        world_data = self.world_form.to_data()
        self.database.update_book_world_form(self.state.selected_book_id, world_data)
        self._refresh_world_markdown_preview()
        self._world_form_dirty = False
        self._world_form_book_id = self.state.selected_book_id
        self._sync_truth_files(self.state.selected_book_id)
        self._index_book_design_for_rag(self.state.selected_book_id)
        self._set_status("世界观设定已保存")
        if refresh:
            self._refresh_drawer()
        return True

    def _import_world_file(self) -> None:
        if not self.state.selected_book_id:
            info(self, "提示", "请先选择一本书。")
            return
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "导入世界观设定文件",
            "",
            "All Supported Files (*.txt *.md *.docx);;Text Files (*.txt *.md);;Word Documents (*.docx);;All Files (*.*)",
        )
        if not file_path:
            return
        try:
            text = self._read_any_file(Path(file_path))
        except Exception as exc:
            error(self, "导入失败", f"无法读取文件：\n{exc}")
            return
        if not text.strip():
            info(self, "导入文件", "文件内容为空。")
            return
        data = merge_world_supplement(
            self.world_form.to_data() if hasattr(self, "world_form") else self.database.get_book_world_form(self.state.selected_book_id),
            text,
            supplement_title=f"导入补充 · {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        )
        self.world_form.set_data(data)
        self._on_world_form_changed()
        self._set_status("世界观设定文件已载入补充区，请核对后保存。")

    def _delete_world_entry(self) -> None:
        self._fill_world_template()

    def _get_bound_skill_payload(self) -> list[dict[str, Any]]:
        payload: list[dict[str, Any]] = []
        seen: set[int] = set()
        for item in [*self.global_skills, *self.book_skills, *self.bound_skills]:
            skill_id = int(item["id"])
            if skill_id in seen:
                continue
            payload.append(dict(item))
            seen.add(skill_id)
        return payload

    def _build_sourcebook_context(self) -> str:
        if not self.state.selected_book_id:
            return ""
        try:
            export_data = self.database.get_book_export_data(self.state.selected_book_id)
        except Exception:  # noqa: BLE001
            return ""
        lines = [f"书籍：{self.state.selected_book_title}"]
        book_outline = str(export_data.get("book", {}).get("outline_text", "")).strip()
        if book_outline:
            lines.append(f"书籍级结构化设定：{book_outline[:2000]}")
        if self.state.selected_volume_id:
            for volume in export_data.get("volumes", []):
                if int(volume.get("id", -1)) == self.state.selected_volume_id:
                    lines.append(f"当前卷：{volume.get('title', '')}")
                    outline = str(volume.get("outline_text", "")).strip()
                    if outline:
                        lines.append(f"卷大纲：{outline[:500]}")
        characters = export_data.get("characters", [])[:8]
        if characters:
            lines.append("人物设定：")
            for item in characters:
                label = f"{item.get('name', '')}（{item.get('role', '')}）" if item.get("role") else str(item.get("name", ""))
                lines.append(f"- {label}：{str(item.get('profile_text', '')).strip()[:600] or '暂无简介'}")
        world_entries = export_data.get("world_entries", [])[:8]
        if world_entries:
            lines.append("世界观规则：")
            for item in world_entries:
                label = f"{item.get('name', '')} [{item.get('category', '')}]" if item.get("category") else str(item.get("name", ""))
                lines.append(f"- {label}：{str(item.get('content_text', '')).strip()[:600] or '暂无内容'}")
        current_chapter_id = self.state.selected_chapter_id
        recent: list[str] = []
        for item in export_data.get("chapters", []):
            if current_chapter_id is not None and int(item.get("id", -1)) == current_chapter_id:
                continue
            summary = str(item.get("summary_text", "")).strip()
            if summary:
                recent.append(f"- {item.get('title', '未命名章节')}：{summary[:160]}")
            if len(recent) >= 4:
                break
        if recent:
            lines.append("可参考的前后章节摘要：")
            lines.extend(recent)
        project_skills = [skill for skill in self._get_bound_skill_payload() if str(skill.get("source_type", "")) == "github_project"]
        if project_skills:
            lines.append("蓝本提炼 Skills：")
            for skill in project_skills[:6]:
                lines.append(f"- {skill.get('name', '未命名 Skill')}：{str(skill.get('instruction_text', '')).strip()[:180]}")
        lines.append("安全边界：蓝本项目只提供抽象工作流和写作约束，不复制代码、README 原文或具体表达。")
        return "\n".join(line for line in lines if str(line).strip())

    def _index_chapter_for_rag(self, book_id: int, chapter_id: int) -> None:
        try:
            chapter = self.database.get_chapter(chapter_id)
            if not chapter:
                return
            self.rag_retriever.index_chapter(
                book_id=book_id,
                chapter_id=chapter_id,
                title=str(chapter["title"] or ""),
                summary=str(chapter["summary_text"] or ""),
                outline=str(chapter["outline"] or ""),
                content_snippet=str(chapter["content"] or "")[:800],
            )
        except Exception:
            pass

    def _index_book_design_for_rag(self, book_id: int) -> None:
        try:
            book = self.database.get_book(book_id)
            if not book:
                return
            character_blocks: list[str] = []
            for character in self.database.list_characters(book_id):
                name = str(character["name"] or "").strip()
                role = str(character["role"] or "").strip()
                profile = str(character["profile_text"] or "").strip()
                if profile:
                    heading = f"### {name}" + (f"（{role}）" if role else "")
                    character_blocks.append(f"{heading}\n\n{profile}")
            character_markdown = "\n\n".join(character_blocks)
            self.rag_retriever.index_book_design(
                book_id=book_id,
                title=str(book["title"] or ""),
                markdown=(
                    f"{str(book['outline_text'] or '').strip()}"
                    f"\n\n## 世界观设定\n{str(book['world_text'] or '').strip()}"
                    f"\n\n## 人物卡设定\n{character_markdown}"
                ).strip(),
            )
        except Exception:
            pass

    def _build_rag_context(self, book_id: int, query: str = "", top_k: int = 5) -> str:
        try:
            if not query.strip():
                return ""
            results = self.rag_retriever.search(book_id, query, top_k=top_k, use_hybrid=True)
            return self.rag_retriever.build_context_from_results(results)
        except Exception:
            return ""

    def _merge_sourcebook_with_rag(self, target_payload: dict[str, Any]) -> str:
        sourcebook = str(target_payload.get("sourcebook_context", ""))
        rag = str(target_payload.get("rag_context", ""))
        if rag:
            separator = "\n\n【智能检索相关章节】\n"
            return sourcebook + separator + rag
        return sourcebook

    def _build_character_immersion_block(
        self,
        book_id: int,
        characters: list[dict[str, Any]] | None = None,
    ) -> str:
        characters = characters if characters is not None else self._named_book_characters(book_id)
        if not characters:
            return ""
        lines = ["【人物 AI 代入与 OOC 约束】"]
        for item in characters[:24]:
            name = str(item.get("name", "") or "").strip()
            if not name:
                continue
            role = str(item.get("role", "") or "").strip()
            profile = str(item.get("profile_text", "") or "").strip()
            label = f"{name}（{role}）" if role else name
            lines.append(f"- {label}：{profile[:320] or '暂无人物设定，请保持行为动机前后一致。'}")
        relationships = [dict(item) for item in self.database.list_relationships(book_id)]
        if relationships:
            lines.append("人物关系约束：")
            for item in relationships[:60]:
                source = str(item.get("source_name", "") or "").strip()
                target = str(item.get("target_name", "") or "").strip()
                relation = str(item.get("relationship_type", "") or "关联").strip()
                description = str(item.get("description", "") or "").strip()
                if source and target:
                    lines.append(f"- {source} -> {target}：{relation}；{description[:120]}")
        lines.append("生成时必须先代入相关人物的立场、动机、口吻和关系，再输出正文；最终校验必须检查 OOC 并在 findings 中记录。")
        return "\n".join(lines)

    def _set_drawer_tab(self, key: str) -> None:
        return self.drawer_controller.set_tab(key)

    def _update_drawer_title(self) -> None:
        return self.drawer_controller.update_title()

    def _open_drawer(self, kind: str = "details") -> None:
        return self.drawer_controller.open(kind)

    def _hide_drawer(self) -> None:
        if hasattr(self, "drawer") and self.drawer.isVisible():
            self.qt_drawer_width = max(320, self.drawer.width())
        self.drawer.setVisible(False)

    def _preserve_workspace_width_for_drawer(self) -> None:
        if not hasattr(self, "main_splitter"):
            return
        sizes = self.main_splitter.sizes()
        if len(sizes) < 2:
            return
        min_workspace = 640
        if sizes[1] >= min_workspace:
            self._apply_responsive_chrome()
            return
        total = max(1, sum(sizes))
        left_width = max(76, total - min_workspace)
        if left_width <= self.library_auto_collapse_width:
            self.library_collapsed = True
            self._apply_library_collapsed_state()
        else:
            self.library_width = max(180, min(left_width, self.library_width))
            self.main_splitter.setSizes([left_width, max(min_workspace, total - left_width)])
        self._apply_responsive_chrome()

    def _toggle_library_collapsed(self) -> None:
        self.focus_mode = False
        self.library_collapsed = not self.library_collapsed
        self._apply_library_collapsed_state()
        self._save_ui_settings()

    def _apply_library_collapsed_state(self) -> None:
        if not hasattr(self, "main_splitter") or not hasattr(self, "library_tree"):
            return
        if self._applying_library_layout:
            return
        self._applying_library_layout = True
        sizes = self.main_splitter.sizes()
        can_resize = self.main_splitter.count() >= 2 and len(sizes) >= 2
        try:
            if self.library_collapsed:
                if can_resize:
                    if sizes[0] > self.library_auto_collapse_width:
                        self.library_width = max(180, sizes[0])
                    self.main_splitter.setSizes([76, max(800, sizes[1])])
                if hasattr(self, "library_layout"):
                    self.library_layout.setContentsMargins(8, 10, 8, 10)
                    self.library_layout.setSpacing(6)
                self.library_title.setVisible(False)
                self.library_create_button.setVisible(False)
                if hasattr(self, "library_generate_button"):
                    self.library_generate_button.setVisible(False)
                if hasattr(self, "library_skills_button"):
                    self.library_skills_button.setVisible(False)
                self.collapse_button.setText(">")
                self.collapse_button.setToolTip("展开导航")
                self.collapse_button.setFixedWidth(42)
                self.library_tree.setVisible(False)
                self.library_compact_list.setVisible(True)
                self.library_compact_list.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
            else:
                if can_resize:
                    self.main_splitter.setSizes([self.library_width, max(800, sizes[1])])
                if hasattr(self, "library_layout"):
                    self.library_layout.setContentsMargins(12, 12, 12, 12)
                    self.library_layout.setSpacing(10)
                self.library_title.setVisible(True)
                self.library_title.setText("作品轨道")
                self.library_create_button.setVisible(True)
                if hasattr(self, "library_generate_button"):
                    self.library_generate_button.setVisible(True)
                if hasattr(self, "library_skills_button"):
                    self.library_skills_button.setVisible(True)
                self.collapse_button.setText("收起")
                self.collapse_button.setToolTip("")
                self.collapse_button.setMinimumWidth(0)
                self.collapse_button.setMaximumWidth(16777215)
                self.library_tree.setVisible(True)
                self.library_compact_list.setVisible(False)
        finally:
            self._applying_library_layout = False

    def _on_splitter_moved(self, _pos: int, _index: int) -> None:
        if self._applying_library_layout:
            return
        sizes = self.main_splitter.sizes()
        if not sizes:
            return
        left_width = sizes[0]
        if left_width <= self.library_auto_collapse_width:
            if not self.library_collapsed:
                self.library_collapsed = True
                self.focus_mode = False
                self._apply_library_collapsed_state()
            self._save_ui_settings()
            return
        self.library_width = max(180, left_width)
        if self.library_collapsed:
            self.library_collapsed = False
            self.focus_mode = False
            self._apply_library_collapsed_state()
        elif self.focus_mode:
            self.focus_mode = False
        self._save_ui_settings()

    def _apply_responsive_chrome(self) -> None:
        workspace_width = self.workspace_panel.width() if hasattr(self, "workspace_panel") else self.width()
        if hasattr(self, "title_label"):
            self.title_label.setVisible(False)
        if hasattr(self, "input_stats_card"):
            self.input_stats_card.setVisible(workspace_width >= 820)
        if hasattr(self, "ai_quick_bar"):
            self.ai_quick_bar.setVisible(True)
        if hasattr(self, "ai_quick_buttons"):
            for button in self.ai_quick_buttons:
                button.setVisible(True)

    def _toggle_focus_mode(self) -> None:
        if self.state.current_node_kind != "chapter":
            info(self, "提示", "专注写作只在章节中可用。")
            return
        sizes = self.main_splitter.sizes()
        workspace_width = sizes[1] if len(sizes) > 1 else 900
        if self.focus_mode:
            self.main_splitter.setSizes([self.library_width, max(700, workspace_width)])
            self.focus_mode = False
            self._set_status("已退出专注写作。")
            return
        self._hide_drawer()
        self.main_splitter.setSizes([72, max(900, workspace_width)])
        self.focus_mode = True
        self._set_status("已进入专注写作；拖动分隔条或点击视图菜单可恢复。")

    def _set_status(self, message: str) -> None:
        if not message:
            return
        text = str(message)
        self.status_label.setToolTip(text)
        if len(text) > 120:
            text = text[:117] + "..."
        self.status_label.setText(text)
        if hasattr(self, "_status_fade_timer"):
            self._status_fade_timer.start(5000)

    def _toggle_auto_save(self, checked: bool) -> None:
        self.auto_save_enabled = bool(checked)
        self._sync_auto_save_timer()
        self._save_ui_settings()
        self._set_status("定时保存已启用" if self.auto_save_enabled else "定时保存已关闭")

    def _auto_save_interval_action_text(self) -> str:
        return f"定时保存间隔：{int(self.auto_save_interval_seconds)} 秒"

    def _configure_auto_save_interval(self) -> None:
        seconds, accepted = QInputDialog.getInt(
            self,
            "定时保存",
            "每几秒自动保存当前章节正文：",
            int(self.auto_save_interval_seconds),
            5,
            3600,
            5,
        )
        if not accepted:
            return
        self.auto_save_interval_seconds = int(seconds)
        self.auto_save_enabled = True
        if hasattr(self, "auto_save_action"):
            self.auto_save_action.setChecked(True)
        if hasattr(self, "auto_save_interval_action"):
            self.auto_save_interval_action.setText(self._auto_save_interval_action_text())
        self._sync_auto_save_timer()
        self._save_ui_settings()
        self._set_status(f"定时保存已设置为每 {self.auto_save_interval_seconds} 秒")

    def _sync_auto_save_timer(self) -> None:
        if not hasattr(self, "auto_save_timer"):
            return
        self.auto_save_timer.stop()
        if self.auto_save_enabled:
            self.auto_save_timer.start(max(5, int(self.auto_save_interval_seconds)) * 1000)
        if hasattr(self, "auto_save_action"):
            self.auto_save_action.setChecked(self.auto_save_enabled)
        if hasattr(self, "auto_save_interval_action"):
            self.auto_save_interval_action.setText(self._auto_save_interval_action_text())

    def _auto_save_current(self) -> None:
        if not self.auto_save_enabled or not self.state.editor_dirty:
            return
        if self.state.editor_scope_kind != "chapter" or not self.state.selected_chapter_id:
            return
        if self._has_active_write_job(int(self.state.selected_chapter_id), {"content", "outline"}):
            return
        if self._save_current():
            self._set_status("已自动保存")

    def _open_ai_chat(self) -> None:
        ai_service = self._require_ai_service("chat", "AI 对话")
        if ai_service is None:
            return
        if self.chat_dialog is None:
            self.chat_dialog = ChatDialog(
                self,
                ai_service,
                self._build_chat_context_prompt,
                self._chat_character_provider,
            )
            self.chat_dialog.destroyed.connect(lambda _obj=None: setattr(self, "chat_dialog", None))
        else:
            self.chat_dialog.update_ai_service(ai_service)
            self.chat_dialog.refresh_characters()
        self.chat_dialog.show()
        self.chat_dialog.raise_()
        self.chat_dialog.activateWindow()

    def _chat_character_provider(self) -> list[dict[str, Any]]:
        if not self.state.selected_book_id:
            return []
        return self._named_book_characters(self.state.selected_book_id)

    def _build_chat_context_prompt(
        self,
        ai_character: dict[str, Any] | None = None,
        self_character: dict[str, Any] | None = None,
        latest_user_message: str = "",
    ) -> str:
        lines = [
            self._prompt_text("chat.system", BUILTIN_CHAT_SYSTEM_PROMPT),
        ]
        latest_user_message = str(latest_user_message or "").strip()
        if latest_user_message:
            lines.append(f"【当前用户请求】\n{latest_user_message[:1200]}")
        if self.state.selected_book_title:
            lines.append(f"作品名称：{self.state.selected_book_title}")
        if self.state.selected_volume_id:
            volume = self.database.get_volume(self.state.selected_volume_id)
            if volume:
                lines.append(f"当前卷：{volume['title']}")
        if self.state.selected_chapter_id:
            chapter = self.database.get_chapter(self.state.selected_chapter_id)
            if chapter:
                lines.append(f"当前章节：{chapter['title']}")
                outline = self.chapter_outline_editor.toPlainText().strip() or str(chapter["outline"] or "").strip()
                content = self.content_editor.toPlainText().strip() or str(chapter["content"] or "").strip()
                if outline:
                    lines.append(f"章节大纲：{outline[:1200]}")
                if content:
                    lines.append(f"当前正文节选：{content[:1800]}")
                runs = self.database.list_review_runs(int(chapter["book_id"]), int(chapter["id"]), limit=1)
                if runs:
                    lines.append(f"最近审查摘要：{str(runs[0]['summary'] or '')[:500]}")
        if self.state.selected_book_id:
            book_id = int(self.state.selected_book_id)
            chapter_id = int(self.state.selected_chapter_id) if self.state.selected_chapter_id else None
            truth_context = self._build_truth_context(book_id, chapter_id, max_chars=9000)
            if truth_context:
                lines.append(f"【Truth File：必须优先遵守】\n{truth_context}")
            rag_query = self._build_chat_rag_query(latest_user_message, ai_character, self_character)
            if rag_query:
                self._index_book_design_for_rag(book_id)
                rag_context = self._build_rag_context(book_id, rag_query, top_k=6)
                if rag_context:
                    lines.append(f"【RAG 检索结果：与当前问题相关】\n{rag_context}")
            lines.append(self._build_book_chat_context(book_id, max_chars=7000))
        if ai_character:
            lines.append(self._build_character_chat_prompt(ai_character, self_character))
        skills = self._get_bound_skill_payload() if self.state.selected_book_id else []
        if skills:
            lines.append("启用 Skills：")
            for skill in skills[:10]:
                lines.append(f"- {skill.get('name', '')}：{str(skill.get('instruction_text', '')).strip()[:180]}")
        return "\n".join(line for line in lines if str(line).strip())

    def _build_chat_rag_query(
        self,
        latest_user_message: str,
        ai_character: dict[str, Any] | None = None,
        self_character: dict[str, Any] | None = None,
    ) -> str:
        query_parts = [str(latest_user_message or "").strip()]
        if self.state.selected_book_title:
            query_parts.append(str(self.state.selected_book_title))
        if self.state.selected_chapter_id:
            chapter = self.database.get_chapter(self.state.selected_chapter_id)
            if chapter:
                query_parts.append(str(chapter["title"] or ""))
                query_parts.append(self.chapter_outline_editor.toPlainText().strip() or str(chapter["outline"] or ""))
                query_parts.append(self.content_editor.toPlainText().strip()[:800])
        for character in (ai_character, self_character):
            if character:
                query_parts.append(str(character.get("name", "") or ""))
                query_parts.append(str(character.get("role", "") or ""))
                query_parts.append(str(character.get("profile_text", "") or "")[:600])
        return "\n".join(part for part in query_parts if part.strip())[:3000]

    def _build_book_chat_context(self, book_id: int, max_chars: int = 16000) -> str:
        book = self.database.get_book(book_id)
        chapters = [dict(item) for item in self.database.get_book_export_data(book_id).get("chapters", [])]
        pieces: list[str] = ["【全书上下文】"]
        if book:
            outline = str(book["outline_text"] or "").strip()
            if outline:
                pieces.append(f"书籍全局描述：{outline[:3000]}")
        character_lines: list[str] = []
        for item in self._named_book_characters(book_id)[:40]:
            detail = self.database.get_character(int(item["id"]))
            record = dict(detail) if detail else item
            character_lines.append(
                f"- {record.get('name', '')}（{record.get('role', '') or '未设定'}）：{str(record.get('profile_text', '') or '')[:500]}"
            )
        if character_lines:
            pieces.append("人物设定：\n" + "\n".join(character_lines))
        chapter_blocks: list[str] = []
        used = 0
        for chapter in chapters:
            title = str(chapter.get("title", "") or "")
            outline = str(chapter.get("outline", "") or "").strip()
            content = str(chapter.get("content", "") or "").strip()
            block = f"--- {title} ---\n大纲：{outline[:600]}\n正文：{content[:1800]}"
            if used + len(block) > max_chars:
                break
            chapter_blocks.append(block)
            used += len(block)
        if chapter_blocks:
            pieces.append("全书正文与章节摘要节选：\n" + "\n\n".join(chapter_blocks))
        return "\n\n".join(pieces)

    def _build_character_chat_prompt(self, ai_character: dict[str, Any], self_character: dict[str, Any] | None = None) -> str:
        name = str(ai_character.get("name", "") or "").strip()
        detail = self.database.get_character(int(ai_character.get("id", 0) or 0)) if ai_character.get("id") else None
        record = dict(detail) if detail else ai_character
        role = str(record.get("role", "") or "").strip()
        profile = str(record.get("profile_text", "") or "").strip()
        self_name = str(self_character.get("name", "") or "").strip() if self_character else ""
        self_line = (
            f"正在与你对话的人是「{self_name}」，请以「{name}」的视角和口吻回应。"
            if self_name
            else ""
        )
        return render_prompt(
            self._prompt_text("chat.character_template", CHARACTER_CHAT_PROMPT_TEMPLATE),
            {
                "name": name,
                "role": role or "未设定",
                "profile": profile or "暂无明确设定，请从全书上下文推断但不要编造关键事实。",
                "self_name": self_name,
                "self_line": self_line,
            },
        )
        result_lines = [
            "【角色对话模式】",
            f"现在你直接代入书中人物「{name}」。",
            f"角色定位：{role or '未设定'}",
            f"人物设定：{profile or '暂无明确设定，请从全书上下文推断但不要编造关键事实。'}",
        ]
        if self_character:
            self_name = str(self_character.get("name", "") or "").strip()
            result_lines.append(f"正在与你对话的人是「{self_name}」，请以「{name}」的视角和口吻回应。")
        result_lines.extend([
            "回答规则：",
            "- 以该角色的第一人称口吻和认知边界回答作者。",
            "- 你已阅读上方全书上下文，应结合剧情、关系和人物经历说话。",
            "- 不要跳出角色解释自己是 AI；除非作者要求写作建议，否则保持角色身份。",
            "- 不确定的设定可以用角色视角表达犹疑，不要硬编书中没有的事实。",
        ])
        return "\n".join(result_lines)

    def _get_previous_chapters_context(
        self, book_id: int, current_chapter_id: int, limit: int = 3
    ) -> list[dict[str, Any]]:
        chapters = self.database.list_chapters(book_id)
        previous: list[dict[str, Any]] = []
        for cid in previous_chapter_ids_in_order(chapters, current_chapter_id, limit):
            chapter = self.database.get_chapter(cid)
            if chapter:
                previous.append({
                    "chapter_id": cid,
                    "title": str(chapter["title"]),
                    "outline": str(chapter["outline"] or ""),
                    "content": str(chapter["content"] or ""),
                    "sort_order": int(chapter["sort_order"] or 0),
                })
        return previous

    def _start_multi_agent_generation(self) -> None:
        if not self.state.selected_chapter_id:
            info(self, "提示", "请先选择一个章节。")
            return
        if self.state.editor_dirty and not self._save_current():
            return
        anchor_chapter = self.database.get_chapter(self.state.selected_chapter_id)
        if not anchor_chapter:
            return
        anchor = dict(anchor_chapter)
        anchor_outline = str(anchor["outline"] or self.chapter_outline_editor.toPlainText()).strip()
        anchor_content = str(anchor["content"] or self.content_editor.toPlainText()).strip()
        has_existing_content = bool(anchor_content)
        if not has_existing_content and not anchor_outline:
            info(self, "提示", "请先为当前章节写入大纲，再启动长篇生成。")
            return
        if self._find_active_background_job("长篇生成"):
            info(self, "任务运行中", "长篇生成已经在运行。")
            return
        ai_service = self._require_ai_service("review", "长篇生成")
        if ai_service is None:
            return
        if self._has_active_write_job(int(anchor["id"])):
            info(self, "任务运行中", "当前章节已有写入型 AI 任务在运行，请等待它完成或先停止任务。")
            return

        if has_existing_content:
            new_outline = ask_multiline(
                self,
                "新章节大纲",
                "当前章节已有正文。请先写入新章节大纲，长篇生成会在当前章节下方插入新章节并按这份大纲生成正文。",
                "",
            )
            if not new_outline:
                return
            title = self._next_chapter_title_for_book(int(anchor["book_id"]))
            new_id = self._create_chapter_after_current(
                book_id=int(anchor["book_id"]),
                after_chapter_id=int(anchor["id"]),
                title=title,
                content="",
                outline=new_outline,
            )
            if new_id is None:
                info(self, "长篇生成", "无法在当前章节后创建新章节。")
                return
            target_chapter_id = new_id
            target_chapter = self.database.get_chapter(target_chapter_id)
            if not target_chapter:
                return
            chapter = dict(target_chapter)
            current_content = ""
            target_title = str(chapter["title"])
            outline = new_outline
        else:
            target_chapter_id = int(anchor["id"])
            chapter = anchor
            current_content = anchor_content
            target_title = str(anchor["title"])
            outline = anchor_outline

        if self._has_active_write_job(target_chapter_id):
            info(self, "任务运行中", "目标章节已有写入型 AI 任务在运行，请等待它完成或先停止任务。")
            return
        snapshot_id = self.database.create_snapshot(
            target_chapter_id,
            f"长篇生成前 · {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            outline,
            current_content,
        )
        target_payload = {
            "chapter_id": target_chapter_id,
            "book_id": int(chapter["book_id"]),
            "title": target_title,
            "outline": outline,
            "content": current_content,
            "snapshot_id": snapshot_id,
            "truth_file": self._build_truth_context(int(chapter["book_id"]), target_chapter_id),
            "template_context": self._build_template_context(int(chapter["book_id"]), target_chapter_id),
            "skills": self._get_review_skill_payload(target_chapter_id),
            "previous_chapters": self._get_previous_chapters_context(int(chapter["book_id"]), target_chapter_id, limit=3),
            "sourcebook_context": self._build_sourcebook_context(),
            "rag_context": self._build_rag_context(int(chapter["book_id"]), outline, top_k=5),
            "anchor_chapter_id": int(anchor["id"]),
            "anchor_title": str(anchor["title"]),
        }
        named_characters = self._named_book_characters(int(chapter["book_id"]))
        target_payload["character_names"] = [
            str(item.get("name", "")).strip()
            for item in named_characters
            if str(item.get("name", "")).strip()
        ]
        target_payload["character_immersion"] = self._build_character_immersion_block(
            int(chapter["book_id"]),
            named_characters,
        )
        self._open_drawer("tasks")
        self._append_task_log(
            f"长篇生成启动：根据目标大纲生成《{target_payload['title']}》正文，"
            f"参考前 {len(target_payload['previous_chapters'])} 个章节。"
        )

        def work(log: Callable[[str], None], cancel_event: Event | None) -> dict[str, Any]:
            log("长篇生成：大纲拆解、人物代入、正文写作和 OOC 校验开始。")
            style_voices_block = ""
            prev_chs = list(target_payload["previous_chapters"])
            character_names = list(target_payload.get("character_names", []))
            voice_sample_chapters = prev_chs or [
                {
                    "title": str(target_payload["title"]),
                    "content": str(target_payload["content"]),
                }
            ]
            if voice_sample_chapters:
                log("提取写作风格特征、角色声音卡片和人物代入约束...")
                style_result = ai_service.extract_style_profile(
                    sample_texts=[str(ch.get("content", "")) for ch in voice_sample_chapters if str(ch.get("content", "")).strip()],
                )
                voice_result = ai_service.extract_character_voices(chapters=voice_sample_chapters, character_names=character_names)
                style_voices_block = ai_service.format_style_and_voices_block(
                    style_profile=str(style_result.get("style_profile", "")),
                    character_voices=list(voice_result.get("character_voices", [])),
                )
                character_immersion = str(target_payload.get("character_immersion", ""))
                if character_immersion:
                    style_voices_block = "\n\n".join(part for part in [style_voices_block, character_immersion] if part.strip())
                if style_voices_block:
                    log("人物代入信息已注入生成上下文。")
            review = ai_service.multi_agent_generate_from_outline(
                book_title=self.state.selected_book_title,
                chapter_title=str(target_payload["title"]),
                outline=outline,
                current_content=current_content,
                truth_file=str(target_payload["truth_file"]),
                selected_skills=list(target_payload["skills"]),
                sourcebook_context=self._merge_sourcebook_with_rag(target_payload),
                require_remote=True,
                cancel_event=cancel_event,
                template_content=str(target_payload["template_context"]),
                previous_chapters=prev_chs,
                style_voices_block=style_voices_block,
            )
            if cancel_event and cancel_event.is_set():
                log("长篇生成已停止，结果不会覆盖正文。")
                return {
                    "scope_type": "generation",
                    "source_label": "长篇生成",
                    "results": [],
                    "cancelled": True,
                    "cancelled_targets": [target_payload],
                    "total": 1,
                }
            log("长篇生成：最终校验完成，准备写回。")
            return {
                "scope_type": "generation",
                "source_label": "长篇生成",
                "results": [{**target_payload, "review": review}],
                "cancelled": False,
                "cancelled_targets": [],
                "total": 1,
                "activate_chapter_id": target_chapter_id,
            }

        self._run_background_ai_job(
            label="长篇生成",
            callback=work,
            on_done=self._apply_multi_agent_review_result,
            callback_uses_log=True,
            callback_uses_cancel=True,
            deliver_result_on_cancel=True,
        )

    def _next_chapter_title_for_book(self, book_id: int) -> str:
        chapters = self.database.list_chapters(book_id)
        return f"第{len(chapters) + 1}章"

    def _start_generation(self, mode: str) -> None:
        if mode == "draft":
            self._start_multi_agent_generation()
            return
        purpose = AI_MODE_PURPOSE.get(mode, "writing")
        if not self.state.selected_chapter_id:
            info(self, "提示", "请先选择一个章节。")
            return
        chapter = self.database.get_chapter(self.state.selected_chapter_id)
        if not chapter:
            return
        ai_service = self._require_ai_service(purpose, "AI 写作功能")
        if ai_service is None:
            return
        target_chapter_id = int(chapter["id"])
        target = "content" if mode in {"draft", "continue", "polish"} else "log"
        if target == "content" and self._has_active_write_job(target_chapter_id):
            info(self, "任务运行中", "当前章节已有写入型 AI 任务在运行，请等待它完成或先停止任务。")
            return
        outline = self.chapter_outline_editor.toPlainText().strip()
        current_content = self.content_editor.toPlainText().strip()
        summary_text = ""
        if mode in {"continue", "polish"} and not current_content:
            info(self, "提示", "当前正文为空，请先运行长篇生成或手动写入正文。")
            return
        if mode in {"draft", "continue", "polish"}:
            self._create_auto_snapshot("AI改写前快照")
            buffer_prefix = "\n\n" if mode == "continue" and current_content else ""
            if mode != "continue":
                self.suppress_ai_span_clear = True
                self.content_editor.clear()
                self.suppress_ai_span_clear = False
        else:
            buffer_prefix = ""
            self.task_log.clear()
            self.task_log_streaming = False
        self._open_drawer("tasks")
        self._run_ai_worker(
            ai_service=ai_service,
            mode=mode,
            target_chapter_id=target_chapter_id,
            target=target,
            stream_mode="append" if mode == "continue" else "replace",
            chapter_title=str(chapter["title"]),
            outline=outline,
            current_content=current_content,
            summary_text=summary_text,
            base_content=current_content,
            base_outline=outline,
            buffer_prefix=buffer_prefix,
        )

    def _refresh_summary(self) -> None:
        if not self.state.selected_chapter_id:
            info(self, "提示", "请先选择一个章节。")
            return
        if self._has_active_write_job(self.state.selected_chapter_id, targets={"outline"}):
            info(self, "任务运行中", "当前章节已有大纲同步任务在运行。")
            return
        ai_service = self._require_ai_service("outline", "同步章节大纲")
        if ai_service is None:
            return
        content = self.content_editor.toPlainText().strip()
        if not content:
            info(self, "提示", "当前正文为空，无法生成摘要。")
            return
        chapter = self.database.get_chapter(self.state.selected_chapter_id)
        if not chapter:
            return
        self._open_drawer("tasks")
        self._run_ai_worker(
            ai_service=ai_service,
            mode="summary",
            target_chapter_id=int(chapter["id"]),
            target="outline",
            stream_mode="replace",
            chapter_title=str(chapter["title"]),
            outline=self.chapter_outline_editor.toPlainText().strip(),
            current_content=content,
            summary_text="",
            base_content=content,
            base_outline=self.chapter_outline_editor.toPlainText().strip(),
            buffer_prefix="",
        )

    def _analyze_book_full(self) -> None:
        if not self.state.selected_book_id:
            info(self, "提示", "请先选择一本书。")
            return
        ai_service = self._require_ai_service("book_analysis", "全书分析")
        if ai_service is None:
            return
        book = self.database.get_book(self.state.selected_book_id)
        if not book:
            return
        chapters = [dict(item) for item in self.database.get_book_export_data(self.state.selected_book_id).get("chapters", [])]
        if not chapters:
            info(self, "全书分析", "当前书籍还没有章节。")
            return
        if not confirm(self, "AI 分析全书", f"将以三级金字塔分析《{book['title']}》的 {len(chapters)} 个章节，并覆盖写入人物、世界观、大纲和事件。是否继续？"):
            return
        self._open_drawer("tasks")
        book_id = int(self.state.selected_book_id)
        book_title = str(book["title"])
        book_outline = str(book["outline_text"] or "")
        characters = [dict(item) for item in self.database.list_characters(book_id)]
        world_entries = [dict(item) for item in self.database.list_world_entries(book_id)]

        def work(log: Callable[[str], None], cancel_event: Event | None) -> dict[str, Any]:
            return self._analyze_book_in_batches(
                ai_service=ai_service,
                book_title=book_title,
                book_outline=book_outline,
                chapters=chapters,
                characters=characters,
                world_entries=world_entries,
                log=log,
                cancel_event=cancel_event,
            )

        def done(payload: Any) -> None:
            if not isinstance(payload, dict):
                return
            self._apply_book_analysis_result(book_id, payload)

        self._run_background_ai_job(
            label="全书分析",
            callback=work,
            on_done=done,
            callback_uses_log=True,
            callback_uses_cancel=True,
        )

    def _analyze_book_in_batches(
        self,
        *,
        ai_service: SimpleAIService,
        book_title: str,
        book_outline: str,
        chapters: list[dict[str, Any]],
        characters: list[dict[str, Any]],
        world_entries: list[dict[str, Any]],
        log: Callable[[str], None],
        cancel_event: Event | None = None,
    ) -> dict[str, Any]:
        max_retries = self._safe_int(
            getattr(self, "ai_request_retries", DEFAULT_AI_REQUEST_RETRIES),
            DEFAULT_AI_REQUEST_RETRIES,
            0,
            5,
        )
        concurrency = self._safe_int(
            getattr(self, "book_analysis_concurrency", DEFAULT_BOOK_ANALYSIS_CONCURRENCY),
            DEFAULT_BOOK_ANALYSIS_CONCURRENCY,
            1,
            8,
        )
        fast_extra_body = self._parse_json_object_text(
            getattr(self, "book_analysis_fast_extra_body", DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON),
            "一级/二级额外请求参数 JSON",
        )
        final_extra_body = self._parse_json_object_text(
            getattr(self, "book_analysis_final_extra_body", DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON),
            "三级额外请求参数 JSON",
        )
        source_chunks = self._build_book_analysis_source_chunks(chapters, max_chars=20000)
        if not source_chunks:
            raise RuntimeError("全书分析没有可读取的章节内容。")
        log(f"全书分析：已按 20000 字切为 {len(source_chunks)} 个一级片段，并发数 {concurrency}。")

        def level1_task(index: int, chunk: dict[str, Any]) -> dict[str, Any]:
            service = ai_service.clone()
            payload = service.analyze_book_level1(
                book_title=book_title,
                book_outline=book_outline,
                chunk_text=str(chunk.get("text", "")),
                chunk_index=index + 1,
                total_chunks=len(source_chunks),
                chunk_id=str(chunk.get("chunk_id", "")),
                chapter_ids=[int(chapter_id) for chapter_id in chunk.get("chapter_ids", []) if str(chapter_id).strip()],
                extra_body=fast_extra_body,
                require_remote=True,
            )
            payload.setdefault("chunk_id", chunk.get("chunk_id"))
            payload.setdefault("chapter_ids", chunk.get("chapter_ids", []))
            return payload

        level1_results = self._run_book_analysis_parallel_stage(
            stage_label="一级解读",
            items=source_chunks,
            worker=level1_task,
            max_workers=concurrency,
            max_retries=max_retries,
            log=log,
            cancel_event=cancel_event,
        )
        if cancel_event and cancel_event.is_set():
            raise RuntimeError("全书分析已取消。")

        level1_results_for_level2 = self._compact_book_analysis_level1_for_level2(level1_results)
        raw_level1_chars = len(json.dumps(level1_results, ensure_ascii=False))
        compact_level1_chars = len(json.dumps(level1_results_for_level2, ensure_ascii=False))
        log(
            "全书分析：已压缩一级输入，"
            f"{raw_level1_chars} 字符 -> {compact_level1_chars} 字符，发送给二级。"
        )

        level2_groups = self._group_book_analysis_level1_results(level1_results_for_level2)
        log(f"全书分析：一级结果整理为 {len(level2_groups)} 个二级分组。")

        def level2_task(index: int, group: list[dict[str, Any]]) -> dict[str, Any]:
            service = ai_service.clone()
            return service.analyze_book_level2(
                book_title=book_title,
                level1_results=group,
                group_index=index + 1,
                total_groups=len(level2_groups),
                extra_body=fast_extra_body,
                require_remote=True,
            )

        level2_results = self._run_book_analysis_parallel_stage(
            stage_label="二级整理",
            items=level2_groups,
            worker=level2_task,
            max_workers=concurrency,
            max_retries=max_retries,
            log=log,
            cancel_event=cancel_event,
        )
        if cancel_event and cancel_event.is_set():
            raise RuntimeError("全书分析已取消。")

        log("全书分析：三级草案正在整合全部二级结果并检查证据缺口。")

        source_index = self._book_analysis_source_index_summary(source_chunks)
        level2_results_for_level3 = self._compact_book_analysis_level2_for_level3(level2_results)
        raw_level2_chars = len(json.dumps(level2_results, ensure_ascii=False))
        compact_level2_chars = len(json.dumps(level2_results_for_level3, ensure_ascii=False))
        log(
            "全书分析：已压缩二级输入，"
            f"{raw_level2_chars} 字符 -> {compact_level2_chars} 字符，发送给三级。"
        )

        def level3_draft_call() -> dict[str, Any]:
            service = ai_service.clone()
            return service.analyze_book_level3_draft(
                book_title=book_title,
                book_outline=book_outline,
                level2_results=level2_results_for_level3,
                chapters=chapters,
                characters=characters,
                world_entries=world_entries,
                source_index=source_index,
                extra_body=final_extra_body,
                require_remote=True,
            )

        draft_payload = self._run_book_analysis_with_retries(
            label="三级草案",
            call=level3_draft_call,
            max_retries=max_retries,
            log=log,
            cancel_event=cancel_event,
        )
        if not isinstance(draft_payload, dict):
            raise RuntimeError("三级草案没有返回有效 JSON 对象。")
        source_requests = self._analysis_source_requests(draft_payload)
        source_context: dict[str, Any] = {"requests": [], "snippets": [], "unresolved_requests": []}
        payload = draft_payload
        if source_requests:
            source_context = self._build_book_analysis_source_context(source_requests, source_chunks)
            log(
                "全书分析：三级请求原文回查 "
                f"{len(source_requests)} 项，命中 {len(source_context.get('snippets', []))} 段。"
            )

            if source_context.get("snippets"):
                def level3_finalize_call() -> dict[str, Any]:
                    service = ai_service.clone()
                    return service.analyze_book_level3_finalize(
                        book_title=book_title,
                        book_outline=book_outline,
                        level2_results=level2_results_for_level3,
                        draft_payload=draft_payload,
                        source_context=source_context,
                        chapters=chapters,
                        characters=characters,
                        world_entries=world_entries,
                        extra_body=final_extra_body,
                        require_remote=True,
                    )

                payload = self._run_book_analysis_with_retries(
                    label="三级原文复核",
                    call=level3_finalize_call,
                    max_retries=max_retries,
                    log=log,
                    cancel_event=cancel_event,
                )
                local_unresolved = source_context.get("unresolved_requests", [])
                if isinstance(payload, dict) and isinstance(local_unresolved, list) and local_unresolved:
                    unresolved = payload.setdefault("analysis_unresolved", [])
                    if isinstance(unresolved, list):
                        unresolved.extend(local_unresolved)
            else:
                unresolved = draft_payload.setdefault("analysis_unresolved", [])
                if isinstance(unresolved, list):
                    unresolved.extend(source_context.get("unresolved_requests", []))
        if not isinstance(payload, dict):
            raise RuntimeError("三级终审没有返回有效 JSON 对象。")
        payload["_requires_evidence"] = True
        payload = self._filter_untrusted_book_analysis_payload(payload)
        unresolved_count = len(payload.get("analysis_unresolved", [])) if isinstance(payload.get("analysis_unresolved"), list) else 0
        log(
            "全书分析：三级证据门禁完成，"
            f"回查 {len(source_requests)} 项，待确认 {unresolved_count} 项。"
        )
        payload["_meta"] = {
            **(payload.get("_meta", {}) if isinstance(payload.get("_meta"), dict) else {}),
            "processed_chapters": len(chapters),
            "level1_count": len(level1_results),
            "level1_input_chars": raw_level1_chars,
            "level1_compacted_input_chars": compact_level1_chars,
            "level2_count": len(level2_results),
            "level2_input_chars": raw_level2_chars,
            "level2_compacted_input_chars": compact_level2_chars,
            "api_retry_limit": max_retries,
            "concurrency": concurrency,
            "analysis_pipeline": "three_level_pyramid_v2_source_review",
            "source_request_count": len(source_requests),
            "source_snippet_count": len(source_context.get("snippets", [])),
            "analysis_unresolved_count": unresolved_count,
        }
        return payload

    def _build_book_analysis_text_chunks(
        self,
        chapters: list[dict[str, Any]],
        *,
        max_chars: int = 20000,
    ) -> list[str]:
        return [
            str(chunk.get("text", ""))
            for chunk in self._build_book_analysis_source_chunks(chapters, max_chars=max_chars)
        ]

    def _build_book_analysis_source_chunks(
        self,
        chapters: list[dict[str, Any]],
        *,
        max_chars: int = 20000,
    ) -> list[dict[str, Any]]:
        parts: list[str] = []
        chapter_ranges: list[dict[str, Any]] = []
        cursor = 0
        for item in chapters:
            chapter_id = item.get("id", "")
            title = str(item.get("title", "") or "")
            outline = str(item.get("outline", "") or "")
            content = str(item.get("content", "") or "")
            block = (
                "\n\n"
                f"[chapter_id={chapter_id}]\n"
                f"标题：{title}\n"
                f"大纲：{outline}\n"
                f"正文：\n{content}"
            )
            start = cursor
            end = start + len(block)
            parts.append(block)
            chapter_ranges.append(
                {
                    "id": chapter_id,
                    "title": title,
                    "char_start": start,
                    "char_end": end,
                }
            )
            cursor = end
        full_text = "".join(parts).strip()
        if not full_text:
            return []
        limit = max(1, int(max_chars))
        chunks: list[dict[str, Any]] = []
        for chunk_index, start in enumerate(range(0, len(full_text), limit), start=1):
            end = min(len(full_text), start + limit)
            text = full_text[start:end]
            covered = [
                item
                for item in chapter_ranges
                if int(item["char_end"]) > start and int(item["char_start"]) < end
            ]
            chapter_ids: list[int] = []
            chapter_refs: list[dict[str, Any]] = []
            for item in covered:
                try:
                    chapter_id_int = int(item.get("id"))
                except (TypeError, ValueError):
                    continue
                if chapter_id_int not in chapter_ids:
                    chapter_ids.append(chapter_id_int)
                    chapter_refs.append(
                        {
                            "id": chapter_id_int,
                            "title": str(item.get("title", "")),
                        }
                    )
            chunks.append(
                {
                    "chunk_id": f"chunk-{chunk_index:04d}",
                    "chunk_index": chunk_index,
                    "text": text,
                    "chapter_ids": chapter_ids,
                    "chapter_refs": chapter_refs,
                    "char_start": start,
                    "char_end": end,
                }
            )
        return chunks

    @staticmethod
    def _book_analysis_source_index_summary(source_chunks: list[dict[str, Any]]) -> list[dict[str, Any]]:
        summary: list[dict[str, Any]] = []
        for chunk in source_chunks:
            text = str(chunk.get("text", ""))
            summary.append(
                {
                    "chunk_id": chunk.get("chunk_id"),
                    "chunk_index": chunk.get("chunk_index"),
                    "chapter_ids": chunk.get("chapter_ids", []),
                    "chapters": chunk.get("chapter_refs", []),
                    "char_count": len(text),
                    "preview": text[:220],
                }
            )
        return summary

    @staticmethod
    def _analysis_source_requests(payload: dict[str, Any]) -> list[dict[str, Any]]:
        requests_value: Any = None
        for key in ("source_requests", "source_request", "verification_requests", "原文回查请求"):
            value = payload.get(key)
            if value:
                requests_value = value
                break
        if isinstance(requests_value, dict):
            requests_value = [requests_value]
        if not isinstance(requests_value, list):
            return []
        requests: list[dict[str, Any]] = []
        for index, item in enumerate(requests_value, start=1):
            if isinstance(item, dict):
                request_item = dict(item)
            else:
                text = str(item or "").strip()
                if not text:
                    continue
                request_item = {"reason": text, "keywords": [text[:40]]}
            request_item.setdefault("request_id", f"source-{index:02d}")
            requests.append(request_item)
        return requests

    def _build_book_analysis_source_context(
        self,
        source_requests: list[dict[str, Any]],
        source_chunks: list[dict[str, Any]],
        *,
        max_requests: int = BOOK_ANALYSIS_SOURCE_REQUEST_LIMIT,
        snippets_per_request: int = BOOK_ANALYSIS_SOURCE_SNIPPETS_PER_REQUEST,
        snippet_chars: int = BOOK_ANALYSIS_SOURCE_SNIPPET_CHARS,
        max_total_chars: int = BOOK_ANALYSIS_SOURCE_CONTEXT_CHARS,
    ) -> dict[str, Any]:
        requests = [dict(item) for item in source_requests[: max(1, int(max_requests))] if isinstance(item, dict)]
        snippets: list[dict[str, Any]] = []
        unresolved: list[dict[str, Any]] = []
        total_chars = 0
        max_total = max(1, int(max_total_chars))
        per_snippet = max(80, int(snippet_chars))
        per_request = max(1, int(snippets_per_request))

        def keyword_list(item: dict[str, Any]) -> list[str]:
            raw_keywords = item.get("keywords") or item.get("keyword") or item.get("terms")
            values: list[str] = []
            if isinstance(raw_keywords, str):
                values.extend(part.strip() for part in re.split(r"[,，、\s]+", raw_keywords) if part.strip())
            elif isinstance(raw_keywords, list):
                values.extend(str(part).strip() for part in raw_keywords if str(part).strip())
            for key in ("target_field", "subject", "question", "reason"):
                text = str(item.get(key, "") or "").strip()
                if text:
                    values.extend(part.strip() for part in re.split(r"[,，、\s]+", text) if len(part.strip()) >= 2)
            seen: set[str] = set()
            ordered: list[str] = []
            for value in values:
                if value and value not in seen:
                    seen.add(value)
                    ordered.append(value)
            return ordered[:8]

        def candidate_chunks(item: dict[str, Any], keywords: list[str]) -> list[dict[str, Any]]:
            chunk_id = str(item.get("chunk_id", "") or "").strip()
            if chunk_id:
                matched = [chunk for chunk in source_chunks if str(chunk.get("chunk_id")) == chunk_id]
                if matched:
                    return matched
            chapter_id_value = item.get("chapter_id")
            chapter_ids: set[int] = set()
            if isinstance(chapter_id_value, list):
                for value in chapter_id_value:
                    try:
                        chapter_ids.add(int(value))
                    except (TypeError, ValueError):
                        continue
            else:
                try:
                    if chapter_id_value not in (None, ""):
                        chapter_ids.add(int(chapter_id_value))
                except (TypeError, ValueError):
                    chapter_ids = set()
            if chapter_ids:
                matched = []
                for chunk in source_chunks:
                    chunk_chapter_ids: set[int] = set()
                    for value in chunk.get("chapter_ids", []):
                        try:
                            chunk_chapter_ids.add(int(value))
                        except (TypeError, ValueError):
                            continue
                    if chapter_ids.intersection(chunk_chapter_ids):
                        matched.append(chunk)
                if matched:
                    return matched
            if keywords:
                matched = [
                    chunk
                    for chunk in source_chunks
                    if any(keyword and keyword in str(chunk.get("text", "")) for keyword in keywords)
                ]
                if matched:
                    return matched
            return list(source_chunks)

        for request_item in requests:
            if total_chars >= max_total:
                unresolved.append({**request_item, "reason": "回查总字数已达上限"})
                continue
            keywords = keyword_list(request_item)
            candidates = candidate_chunks(request_item, keywords)
            if not candidates:
                unresolved.append({**request_item, "reason": "未找到匹配原文切片"})
                continue
            added = 0
            for chunk in candidates:
                if added >= per_request or total_chars >= max_total:
                    break
                text = str(chunk.get("text", ""))
                if not text:
                    continue
                found_keyword = ""
                found_at = -1
                for keyword in keywords:
                    found_at = text.find(keyword)
                    if found_at >= 0:
                        found_keyword = keyword
                        break
                if found_at < 0:
                    found_at = 0
                half = per_snippet // 2
                start = max(0, found_at - half)
                end = min(len(text), start + per_snippet)
                start = max(0, end - per_snippet)
                snippet_text = text[start:end].strip()
                remaining = max_total - total_chars
                if remaining <= 0:
                    break
                if len(snippet_text) > remaining:
                    snippet_text = snippet_text[:remaining].rstrip()
                if not snippet_text:
                    continue
                snippets.append(
                    {
                        "request_id": request_item.get("request_id"),
                        "target_field": request_item.get("target_field", ""),
                        "reason": request_item.get("reason", ""),
                        "chunk_id": chunk.get("chunk_id"),
                        "chapter_ids": chunk.get("chapter_ids", []),
                        "keyword": found_keyword,
                        "text": snippet_text,
                    }
                )
                total_chars += len(snippet_text)
                added += 1
            if added == 0:
                unresolved.append({**request_item, "reason": "匹配到切片但无法生成非空片段"})
        return {
            "requests": requests,
            "snippets": snippets,
            "unresolved_requests": unresolved,
            "total_chars": total_chars,
        }

    def _group_book_analysis_level1_results(
        self,
        level1_results: list[dict[str, Any]],
        *,
        group_size: int = 3,
        max_group_chars: int = 30000,
    ) -> list[list[dict[str, Any]]]:
        if not level1_results:
            return []
        size = max(1, int(group_size))
        char_budget = max(6000, int(max_group_chars))
        groups: list[list[dict[str, Any]]] = []
        current: list[dict[str, Any]] = []
        current_chars = 2
        for item in level1_results:
            record = item if isinstance(item, dict) else {"value": str(item)}
            record_chars = len(json.dumps([record], ensure_ascii=False))
            if record_chars > char_budget:
                compacted = self._compact_analysis_value_for_level3(record, text_limit=120, item_limit=6)
                if isinstance(compacted, dict):
                    record = compacted
                    record_chars = len(json.dumps([record], ensure_ascii=False))
            if current and (len(current) >= size or current_chars + record_chars > char_budget):
                groups.append(current)
                current = []
                current_chars = 2
            current.append(record)
            current_chars = len(json.dumps(current, ensure_ascii=False))
        if current:
            groups.append(current)
        return groups

    @classmethod
    def _compact_book_analysis_level1_record_for_level2(cls, item: dict[str, Any]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for key in ("level", "chunk_index", "total_chunks", "chunk_id", "chapter_ids"):
            if key in item:
                result[key] = item.get(key)
        digest = cls._compact_text_for_level3(item.get("faithful_digest") or item.get("summary"), 900)
        if digest:
            result["faithful_digest"] = digest
        list_limits = {
            "key_events": (18, 180),
            "character_clues": (24, 220),
            "world_clues": (20, 220),
            "chapter_updates": (36, 220),
            "source_facts": (28, 240),
            "conflicts": (10, 220),
            "coverage_gaps": (10, 220),
            "uncertainties": (10, 220),
        }
        for key, (item_limit, text_limit) in list_limits.items():
            value = item.get(key)
            if not isinstance(value, list) or not value:
                continue
            compacted = cls._compact_analysis_value_for_level3(
                value,
                text_limit=text_limit,
                item_limit=item_limit,
            )
            if cls._analysis_value_has_content(compacted):
                result[key] = compacted
        return result

    @classmethod
    def _compact_book_analysis_level1_for_level2(
        cls,
        level1_results: list[dict[str, Any]],
        *,
        max_chars: int = 90000,
    ) -> list[dict[str, Any]]:
        compacted: list[dict[str, Any]] = []
        max_chars = max(12000, int(max_chars))
        for index, item in enumerate(level1_results):
            if not isinstance(item, dict):
                continue
            record = cls._compact_book_analysis_level1_record_for_level2(item)
            if not record:
                continue
            candidate = compacted + [record]
            if len(json.dumps(candidate, ensure_ascii=False)) > max_chars:
                compacted.append(
                    {
                        "compaction_note": "部分一级结果因二级上下文上限被截断，后续三级可通过 source_requests 回查原文。",
                        "omitted_level1_chunks": len(level1_results) - index,
                    }
                )
                break
            compacted.append(record)
        if not compacted:
            compacted.append(
                {
                    "compaction_note": "一级结果为空或无法压缩，二级应只输出 coverage_gaps 供三级回查。",
                    "original_level1_chunks": len(level1_results),
                }
            )
        return compacted

    @staticmethod
    def _compact_text_for_level3(value: Any, limit: int = 420) -> str:
        text = str(value or "").strip()
        if len(text) <= limit:
            return text
        return text[: max(0, limit - 1)].rstrip() + "…"

    @classmethod
    def _compact_analysis_value_for_level3(
        cls,
        value: Any,
        *,
        text_limit: int = 420,
        item_limit: int = 8,
        depth: int = 0,
    ) -> Any:
        if value is None or isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return value
        if isinstance(value, str):
            return cls._compact_text_for_level3(value, text_limit)
        if isinstance(value, (list, tuple)):
            return [
                cls._compact_analysis_value_for_level3(
                    item,
                    text_limit=max(120, text_limit // 2),
                    item_limit=max(3, item_limit // 2),
                    depth=depth + 1,
                )
                for item in list(value)[:item_limit]
            ]
        if isinstance(value, dict):
            if depth >= 3:
                return cls._compact_text_for_level3(json.dumps(value, ensure_ascii=False), text_limit)
            priority = (
                "chapter_id",
                "chapter_ids",
                "chunk_id",
                "chunk_ids",
                "fact_id",
                "fact_ids",
                "source_fact_ids",
                "name",
                "role",
                "subject",
                "category",
                "claim",
                "rule",
                "evidence",
                "evidence_ref",
                "evidence_refs",
                "evidence_quote",
                "evidence_quotes",
                "source_quote",
                "source_quotes",
                "confidence",
                "uncertainty",
                "outline",
                "events",
                "ai_probability",
                "ai_probability_level",
                "goal",
                "conflict",
                "change",
                "final_state",
                "target_field",
                "reason",
            )
            result: dict[str, Any] = {}
            for key in priority:
                if key in value:
                    result[key] = cls._compact_analysis_value_for_level3(
                        value.get(key),
                        text_limit=text_limit,
                        item_limit=item_limit,
                        depth=depth + 1,
                    )
            for key, item in value.items():
                if key in result or key in {"raw_text", "content", "full_text", "faithful_digest"}:
                    continue
                if len(result) >= 14:
                    break
                if not cls._analysis_value_has_content(item):
                    continue
                result[str(key)] = cls._compact_analysis_value_for_level3(
                    item,
                    text_limit=max(120, text_limit // 2),
                    item_limit=max(3, item_limit // 2),
                    depth=depth + 1,
                )
            return result
        return cls._compact_text_for_level3(value, text_limit)

    @classmethod
    def _compact_book_analysis_level2_record_for_level3(cls, item: dict[str, Any]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        scalar_limits = {
            "story_summary": 1200,
            "summary": 900,
            "group_summary": 900,
        }
        for key in ("level", "group_index", "total_groups"):
            if key in item:
                result[key] = item.get(key)
        for key, limit in scalar_limits.items():
            text = cls._compact_text_for_level3(item.get(key), limit)
            if text:
                result[key] = text
        list_limits = {
            "logic_chain": (18, 220),
            "timeline": (22, 220),
            "output_hints": (28, 220),
            "character_arcs": (36, 260),
            "world_model": (36, 260),
            "chapter_updates": (80, 240),
            "consolidated_facts": (70, 260),
            "source_facts": (50, 240),
            "facts": (50, 240),
            "conflicts": (24, 220),
            "coverage_gaps": (24, 220),
        }
        for key, (item_limit, text_limit) in list_limits.items():
            value = item.get(key)
            if not isinstance(value, list) or not value:
                continue
            compacted = cls._compact_analysis_value_for_level3(
                value,
                text_limit=text_limit,
                item_limit=item_limit,
            )
            if cls._analysis_value_has_content(compacted):
                result[key] = compacted
        return result

    @classmethod
    def _compact_book_analysis_level2_for_level3(
        cls,
        level2_results: list[dict[str, Any]],
        *,
        max_chars: int = 45000,
    ) -> list[dict[str, Any]]:
        if not level2_results:
            return []
        compacted: list[dict[str, Any]] = []
        max_chars = max(8000, int(max_chars))
        for index, item in enumerate(level2_results):
            if not isinstance(item, dict):
                continue
            record = cls._compact_book_analysis_level2_record_for_level3(item)
            if not record:
                continue
            candidate = compacted + [record]
            if len(json.dumps(candidate, ensure_ascii=False)) > max_chars:
                remaining = len(level2_results) - index
                compacted.append(
                    {
                        "compaction_note": "部分二级结果因三级上下文上限被截断，请必要时通过 source_requests 回查原文。",
                        "omitted_level2_groups": remaining,
                    }
                )
                break
            compacted.append(record)
        if not compacted:
            compacted.append(
                {
                    "compaction_note": "二级结果为空或无法压缩，三级应主要依赖 source_index 与原文回查。",
                    "original_level2_groups": len(level2_results),
                }
            )
        return compacted

    def _run_book_analysis_with_retries(
        self,
        *,
        label: str,
        call: Callable[[], dict[str, Any]],
        max_retries: int,
        log: Callable[[str], None],
        cancel_event: Event | None,
    ) -> dict[str, Any]:
        for attempt in range(max_retries + 1):
            if cancel_event and cancel_event.is_set():
                raise RuntimeError("全书分析已取消。")
            try:
                return call()
            except Exception as exc:  # noqa: BLE001
                friendly = self._friendly_ai_error(str(exc))
                if attempt < max_retries:
                    log(f"{label}远程分析失败，正在重试 {attempt + 1}/{max_retries}：{friendly}")
                    continue
                raise RuntimeError(f"{label}远程分析失败，已重试 {max_retries} 次：{friendly}") from exc
        raise RuntimeError(f"{label}远程分析失败。")

    def _run_book_analysis_parallel_stage(
        self,
        *,
        stage_label: str,
        items: list[Any],
        worker: Callable[[int, Any], dict[str, Any]],
        max_workers: int,
        max_retries: int,
        log: Callable[[str], None],
        cancel_event: Event | None,
    ) -> list[dict[str, Any]]:
        if not items:
            return []
        results: list[dict[str, Any] | None] = [None] * len(items)
        worker_count = max(1, min(int(max_workers), len(items)))

        def run_one(index: int, item: Any) -> dict[str, Any]:
            return self._run_book_analysis_with_retries(
                label=f"{stage_label} {index + 1}/{len(items)}",
                call=lambda: worker(index, item),
                max_retries=max_retries,
                log=log,
                cancel_event=cancel_event,
            )

        with ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix="book-analysis") as executor:
            pending: dict[Any, int] = {}
            next_index = 0

            def submit_ready() -> None:
                nonlocal next_index
                while next_index < len(items) and len(pending) < worker_count:
                    if cancel_event and cancel_event.is_set():
                        return
                    log(f"全书分析：{stage_label} {next_index + 1}/{len(items)} 开始。")
                    future = executor.submit(run_one, next_index, items[next_index])
                    pending[future] = next_index
                    next_index += 1

            submit_ready()
            while pending:
                if cancel_event and cancel_event.is_set():
                    for future in pending:
                        future.cancel()
                    raise RuntimeError("全书分析已取消。")
                done, _not_done = wait(pending.keys(), timeout=0.2, return_when=FIRST_COMPLETED)
                if not done:
                    continue
                for future in done:
                    index = pending.pop(future)
                    results[index] = future.result()
                    log(f"全书分析：{stage_label} {index + 1}/{len(items)} 完成。")
                submit_ready()
        return [item for item in results if isinstance(item, dict)]

    @staticmethod
    def _analysis_world_text(payload: dict[str, Any]) -> str:
        for key in (
            "world_text",
            "world_markdown",
            "worldview_markdown",
            "world_setting_markdown",
            "world_settings_markdown",
            "world_form",
            "world",
            "worldview",
            "world_setting",
            "world_settings",
            "world_summary",
            "世界观",
        ):
            value = payload.get(key)
            if isinstance(value, dict):
                text = NovelQtMainWindow._analysis_world_dict_text(value)
                if text:
                    return text
                continue
            if isinstance(value, (list, tuple)):
                continue
            text = str(value or "").strip()
            if text.startswith("{"):
                try:
                    parsed = json.loads(text)
                except (TypeError, ValueError):
                    parsed = None
                if isinstance(parsed, dict):
                    parsed_text = NovelQtMainWindow._analysis_world_dict_text(parsed)
                    if parsed_text:
                        return parsed_text
                    continue
            if text:
                return text
        return ""

    @staticmethod
    def _analysis_world_schema_field_ids() -> set[str]:
        field_ids: set[str] = set()
        for section in WORLD_FORM_SCHEMA:
            for item in section.get("items", []):
                if isinstance(item, dict) and item.get("id"):
                    field_ids.add(str(item["id"]))
        return field_ids

    @staticmethod
    def _analysis_world_dict_text(value: dict[str, Any]) -> str:
        field_ids = NovelQtMainWindow._analysis_world_schema_field_ids()
        if isinstance(value.get("fields"), dict) or any(str(key) in field_ids for key in value.keys()):
            return ""
        lines: list[str] = []
        ignored = {"id", "book_id", "created_at", "updated_at"}
        for key, item in value.items():
            if key in ignored or item is None:
                continue
            if isinstance(item, (dict, list, tuple)):
                item_text = json.dumps(item, ensure_ascii=False)
            else:
                item_text = str(item or "").strip()
            if item_text:
                lines.append(f"{key}: {item_text}")
        return "\n".join(lines).strip()

    def _analysis_world_form(self, payload: dict[str, Any]) -> dict[str, Any] | None:
        field_ids = self._analysis_world_schema_field_ids()
        for key in (
            "world_form",
            "world_data",
            "world_data_json",
            "world_fields",
            "worldview_form",
            "world_setting_form",
            "world_settings_form",
            "world",
            "worldview",
            "world_setting",
            "world_settings",
            "世界观",
        ):
            if key not in payload:
                continue
            parsed = self._analysis_json_dict(payload.get(key))
            if parsed is not None and (
                isinstance(parsed.get("fields"), dict)
                or isinstance(parsed.get("supplements"), list)
                or any(str(field_key) in field_ids for field_key in parsed.keys())
            ):
                return parsed
        if any(str(key) in field_ids for key in payload.keys()):
            return payload
        return None

    @staticmethod
    def _analysis_first_text(item: dict[str, Any], keys: tuple[str, ...]) -> str:
        for key in keys:
            value = item.get(key)
            if isinstance(value, (dict, list, tuple)):
                continue
            text = str(value or "").strip()
            if text:
                return text
        return ""

    @staticmethod
    def _analysis_json_dict(value: Any) -> dict[str, Any] | None:
        if isinstance(value, dict):
            return value
        if not isinstance(value, str):
            return None
        text = value.strip()
        if not text or not text.startswith("{"):
            return None
        try:
            parsed = json.loads(text)
        except (TypeError, ValueError):
            return None
        return parsed if isinstance(parsed, dict) else None

    @staticmethod
    def _analysis_character_items(payload: dict[str, Any]) -> list[dict[str, Any]]:
        for key in (
            "characters",
            "character_cards",
            "character_profiles",
            "character_profile",
            "character_data",
            "人物卡",
            "角色卡",
            "人物",
        ):
            value = payload.get(key)
            if isinstance(value, list):
                return [item for item in value if isinstance(item, dict)]
            if isinstance(value, dict):
                if any(
                    marker in value
                    for marker in (
                        "name",
                        "character_name",
                        "profile_form",
                        "profile_data",
                        "profile_data_json",
                        "profile_text",
                        "profile_markdown",
                    )
                ):
                    return [value]
                records: list[dict[str, Any]] = []
                for name, record in value.items():
                    if isinstance(record, dict):
                        item = dict(record)
                        item.setdefault("name", str(name))
                        records.append(item)
                    else:
                        text = str(record or "").strip()
                        if text:
                            records.append({"name": str(name), "profile_text": text})
                if records:
                    return records
        return []

    @staticmethod
    def _analysis_character_schema_field_ids(role: str, schema_key: str = "") -> set[str]:
        field_ids: set[str] = set()
        for section in character_profile_schema_for_role(role, schema_key):
            for item in section.get("items", []):
                if isinstance(item, dict) and item.get("id"):
                    field_ids.add(str(item["id"]))
        return field_ids

    def _analysis_character_profile_form(self, item: dict[str, Any], role: str) -> dict[str, Any] | None:
        for key in (
            "character_form",
            "character_profile_form",
            "profile_form",
            "profile_data",
            "profile_data_json",
            "character_profile",
            "profile_fields",
            "fields",
            "form",
            "profile",
        ):
            if key not in item:
                continue
            parsed = self._analysis_json_dict(item.get(key))
            if parsed is not None:
                return parsed
        schema_key = str(item.get("schema", "") or item.get("profile_schema", "") or "")
        field_ids = self._analysis_character_schema_field_ids(role, schema_key)
        if any(str(key) in field_ids for key in item.keys()):
            return item
        return None

    @staticmethod
    def _analysis_character_profile_text(item: dict[str, Any]) -> str:
        for key in (
            "profile_text",
            "profile_markdown",
            "markdown",
            "content",
            "description",
            "summary",
            "profile",
        ):
            value = item.get(key)
            if isinstance(value, (dict, list, tuple)):
                continue
            text = str(value or "").strip()
            if text:
                return text
        return ""

    @staticmethod
    def _book_analysis_requires_evidence(payload: dict[str, Any]) -> bool:
        if bool(payload.get("_requires_evidence")):
            return True
        meta = payload.get("_meta", {})
        return isinstance(meta, dict) and str(meta.get("analysis_pipeline", "")).startswith("three_level_pyramid_v2")

    @staticmethod
    def _analysis_value_has_content(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return True
        if isinstance(value, str):
            return bool(value.strip())
        if isinstance(value, dict):
            return any(NovelQtMainWindow._analysis_value_has_content(item) for item in value.values())
        if isinstance(value, (list, tuple, set)):
            return any(NovelQtMainWindow._analysis_value_has_content(item) for item in value)
        return value is not None

    @staticmethod
    def _analysis_has_evidence(value: Any) -> bool:
        if not isinstance(value, dict):
            if isinstance(value, (list, tuple)):
                return any(NovelQtMainWindow._analysis_has_evidence(item) for item in value)
            return False
        evidence_keys = {
            "evidence_refs",
            "evidence_ref",
            "source_quotes",
            "source_quote",
            "evidence_quotes",
            "evidence_quote",
            "evidence",
            "fact_id",
            "fact_ids",
            "chunk_id",
            "chunk_ids",
            "source_fact_ids",
            "source",
            "sources",
        }
        for key in evidence_keys:
            if key in value and NovelQtMainWindow._analysis_value_has_content(value.get(key)):
                return True
        return any(NovelQtMainWindow._analysis_has_evidence(item) for item in value.values())

    @staticmethod
    def _analysis_has_direct_evidence(value: Any) -> bool:
        if not isinstance(value, dict):
            return False
        direct_keys = {
            "evidence_refs",
            "evidence_ref",
            "source_quotes",
            "source_quote",
            "evidence_quotes",
            "evidence_quote",
            "evidence",
            "fact_id",
            "fact_ids",
            "chunk_id",
            "chunk_ids",
            "source_fact_ids",
            "source",
            "sources",
        }
        return any(
            key in value and NovelQtMainWindow._analysis_value_has_content(value.get(key))
            for key in direct_keys
        )

    @staticmethod
    def _analysis_field_has_written_content(field: dict[str, Any]) -> bool:
        if bool(field.get("checked", False)):
            return True
        if str(field.get("value", "") or "").strip():
            return True
        return False

    @staticmethod
    def _analysis_form_has_content(form: dict[str, Any] | None) -> bool:
        if not isinstance(form, dict):
            return False
        fields = form.get("fields", {})
        if isinstance(fields, dict):
            for field in fields.values():
                if isinstance(field, dict) and NovelQtMainWindow._analysis_field_has_written_content(field):
                    return True
        for plan in form.get("volume_plans", []) if isinstance(form.get("volume_plans"), list) else []:
            if not isinstance(plan, dict):
                continue
            plan_fields = plan.get("fields", {})
            if isinstance(plan_fields, dict):
                for field in plan_fields.values():
                    if isinstance(field, dict) and str(field.get("value", "") or "").strip():
                        return True
        supplements = form.get("supplements", [])
        if isinstance(supplements, list):
            for supplement in supplements:
                if isinstance(supplement, dict) and str(supplement.get("content", "") or "").strip():
                    return True
        return False

    def _prune_untrusted_analysis_form(self, form: dict[str, Any]) -> tuple[dict[str, Any], int]:
        pruned = deepcopy(form)
        dropped = 0
        fields = pruned.get("fields", {})
        if isinstance(fields, dict):
            for field_id in list(fields.keys()):
                field = fields.get(field_id)
                if not isinstance(field, dict):
                    fields.pop(field_id, None)
                    continue
                if not self._analysis_field_has_written_content(field):
                    fields.pop(field_id, None)
                    continue
                if not self._analysis_has_evidence(field):
                    fields.pop(field_id, None)
                    dropped += 1
        volume_plans = pruned.get("volume_plans", [])
        if isinstance(volume_plans, list):
            kept_plans: list[dict[str, Any]] = []
            for plan in volume_plans:
                if not isinstance(plan, dict):
                    continue
                plan_copy = deepcopy(plan)
                plan_fields = plan_copy.get("fields", {})
                if isinstance(plan_fields, dict):
                    for field_id in list(plan_fields.keys()):
                        field = plan_fields.get(field_id)
                        if not isinstance(field, dict) or not str(field.get("value", "") or "").strip():
                            plan_fields.pop(field_id, None)
                            continue
                        if not self._analysis_has_evidence(field):
                            plan_fields.pop(field_id, None)
                            dropped += 1
                if self._analysis_form_has_content({"fields": {}, "volume_plans": [plan_copy], "supplements": []}):
                    kept_plans.append(plan_copy)
            pruned["volume_plans"] = kept_plans
        supplements = pruned.get("supplements", [])
        if isinstance(supplements, list):
            kept_supplements: list[dict[str, Any]] = []
            for supplement in supplements:
                if not isinstance(supplement, dict):
                    continue
                if not str(supplement.get("content", "") or "").strip():
                    continue
                if self._analysis_has_evidence(supplement):
                    kept_supplements.append(supplement)
                else:
                    dropped += 1
            pruned["supplements"] = kept_supplements
        return pruned, dropped

    def _overlay_sparse_analysis_form(self, base: dict[str, Any] | None, incoming: dict[str, Any]) -> dict[str, Any]:
        result = deepcopy(base) if isinstance(base, dict) else {}
        for key in ("version", "schema", "name", "role"):
            if key in incoming and str(incoming.get(key, "") or "").strip():
                result[key] = incoming[key]
        incoming_fields = incoming.get("fields", {})
        if isinstance(incoming_fields, dict):
            result_fields = result.setdefault("fields", {})
            if not isinstance(result_fields, dict):
                result["fields"] = result_fields = {}
            for field_id, field in incoming_fields.items():
                if not isinstance(field, dict):
                    continue
                target = result_fields.setdefault(str(field_id), {})
                if not isinstance(target, dict):
                    target = {}
                    result_fields[str(field_id)] = target
                if "checked" in field:
                    target["checked"] = bool(field.get("checked", False))
                if "value" in field:
                    target["value"] = str(field.get("value", "") or "")
        incoming_plans = incoming.get("volume_plans", [])
        if isinstance(incoming_plans, list):
            result_plans = result.setdefault("volume_plans", [])
            if not isinstance(result_plans, list):
                result["volume_plans"] = result_plans = []
            by_id = {str(plan.get("volume_id", "")): plan for plan in result_plans if isinstance(plan, dict)}
            for incoming_plan in incoming_plans:
                if not isinstance(incoming_plan, dict):
                    continue
                key = str(incoming_plan.get("volume_id", ""))
                target_plan = by_id.get(key)
                if target_plan is None:
                    target_plan = {
                        "volume_id": incoming_plan.get("volume_id", len(result_plans) + 1),
                        "volume_title": incoming_plan.get("volume_title", ""),
                        "sort_order": incoming_plan.get("sort_order", len(result_plans) + 1),
                        "fields": {},
                    }
                    result_plans.append(target_plan)
                    by_id[key] = target_plan
                target_fields = target_plan.setdefault("fields", {})
                if not isinstance(target_fields, dict):
                    target_plan["fields"] = target_fields = {}
                incoming_fields = incoming_plan.get("fields", {})
                if not isinstance(incoming_fields, dict):
                    continue
                for field_id, field in incoming_fields.items():
                    if isinstance(field, dict) and "value" in field:
                        target_fields[str(field_id)] = {"value": str(field.get("value", "") or "")}
        incoming_supplements = incoming.get("supplements", [])
        if isinstance(incoming_supplements, list):
            result_supplements = result.setdefault("supplements", [])
            if not isinstance(result_supplements, list):
                result["supplements"] = result_supplements = []
            existing = {
                (str(item.get("title", "") or ""), str(item.get("content", "") or ""))
                for item in result_supplements
                if isinstance(item, dict)
            }
            for supplement in incoming_supplements:
                if not isinstance(supplement, dict):
                    continue
                marker = (str(supplement.get("title", "") or ""), str(supplement.get("content", "") or ""))
                if marker not in existing:
                    result_supplements.append(supplement)
                    existing.add(marker)
        return result

    def _filter_untrusted_book_analysis_payload(self, payload: dict[str, Any]) -> dict[str, Any]:
        if not self._book_analysis_requires_evidence(payload):
            return payload
        filtered = deepcopy(payload)
        dropped = 0
        unresolved = filtered.get("analysis_unresolved", [])
        if not isinstance(unresolved, list):
            unresolved = []
        for key in ("book_outline_form", "world_form"):
            value = filtered.get(key)
            if isinstance(value, dict):
                pruned, count = self._prune_untrusted_analysis_form(value)
                dropped += count
                if self._analysis_form_has_content(pruned):
                    filtered[key] = pruned
                else:
                    filtered.pop(key, None)
        if filtered.get("book_outline") and not self._analysis_has_direct_evidence(filtered):
            filtered.pop("book_outline", None)
            dropped += 1
        if filtered.get("world_text") and not self._analysis_has_direct_evidence(filtered):
            filtered.pop("world_text", None)
            dropped += 1

        characters = self._analysis_character_items(filtered)
        if characters:
            kept_characters: list[dict[str, Any]] = []
            for item in characters:
                role = self._analysis_first_text(item, ("role", "character_role", "position", "定位", "角色定位"))
                profile_form = self._analysis_character_profile_form(item, role)
                item_copy = deepcopy(item)
                has_trusted_content = False
                if isinstance(profile_form, dict):
                    pruned_form, count = self._prune_untrusted_analysis_form(profile_form)
                    dropped += count
                    for key in ("character_form", "character_profile_form", "profile_form", "profile_data", "profile_data_json", "character_profile", "profile_fields", "fields", "form", "profile"):
                        if key in item_copy:
                            item_copy.pop(key, None)
                    if self._analysis_form_has_content(pruned_form):
                        item_copy["profile_form"] = pruned_form
                        has_trusted_content = True
                profile_text = self._analysis_character_profile_text(item)
                if profile_text:
                    if self._analysis_has_evidence(item):
                        item_copy["profile_text"] = profile_text
                        has_trusted_content = True
                    else:
                        item_copy.pop("profile_text", None)
                        item_copy.pop("profile_markdown", None)
                        dropped += 1
                if has_trusted_content:
                    kept_characters.append(item_copy)
                else:
                    dropped += 1
            filtered["characters"] = kept_characters
            for key in ("character_cards", "character_profiles", "character_profile", "character_data"):
                filtered.pop(key, None)

        chapter_items = filtered.get("chapter_updates", [])
        if isinstance(chapter_items, list):
            kept_updates: list[dict[str, Any]] = []
            for item in chapter_items:
                if isinstance(item, dict) and self._analysis_has_evidence(item):
                    kept_updates.append(item)
                elif isinstance(item, dict):
                    dropped += 1
            filtered["chapter_updates"] = kept_updates
        if dropped:
            unresolved.append({"title": "证据门禁剔除", "content": f"{dropped} 个字段或条目缺少证据，未写入。"})
        filtered["analysis_unresolved"] = unresolved
        meta = filtered.get("_meta", {})
        if not isinstance(meta, dict):
            meta = {}
        meta["evidence_gate_dropped"] = dropped
        filtered["_meta"] = meta
        return filtered

    def _merge_book_analysis_payload(self, merged: dict[str, Any], payload: dict[str, Any]) -> None:
        outline = str(payload.get("book_outline", "") or "").strip()
        if outline:
            existing_outline = str(merged.get("book_outline", "") or "").strip()
            merged["book_outline"] = f"{existing_outline}\n\n{outline}".strip() if existing_outline else outline
        outline_form = payload.get("book_outline_form")
        if isinstance(outline_form, dict):
            existing_form = merged.get("book_outline_form")
            merged["book_outline_form"] = merge_book_outline_form(
                existing_form if isinstance(existing_form, dict) else None,
                outline_form,
                "",
            )
        world_form = self._analysis_world_form(payload)
        world_text = self._analysis_world_text(payload)
        world_entries = payload.get("world_entries", [])
        if isinstance(world_form, dict) or world_text or isinstance(world_entries, list):
            existing_world_form = merged.get("world_form")
            merged["world_form"] = merge_world_analysis_payload(
                existing_world_form if isinstance(existing_world_form, dict) else None,
                world_form=world_form,
                world_entries=world_entries,
                world_text=world_text,
            )

        def merge_named_list(target_key: str, content_key: str) -> None:
            incoming = self._analysis_character_items(payload) if target_key == "characters" else payload.get(target_key, [])
            if not isinstance(incoming, list):
                return
            existing: dict[str, dict[str, Any]] = {
                str(item.get("name", "")).strip(): item
                for item in merged.get(target_key, [])
                if isinstance(item, dict) and str(item.get("name", "")).strip()
            }
            for item in incoming:
                if not isinstance(item, dict):
                    continue
                name = str(item.get("name", "")).strip()
                if not name:
                    continue
                current = existing.get(name)
                if current is None:
                    current = {"name": name}
                    existing[name] = current
                role = self._analysis_first_text(item, ("role", "character_role", "position", "定位", "角色定位"))
                if role and not str(current.get("role", "") or "").strip():
                    current["role"] = role
                profile_form = self._analysis_character_profile_form(item, role or str(current.get("role", "") or ""))
                if isinstance(profile_form, dict):
                    current["profile_form"] = merge_character_profile_form(
                        current.get("profile_form") if isinstance(current.get("profile_form"), dict) else None,
                        profile_form,
                        name,
                        role or str(current.get("role", "") or ""),
                    )
                profile_text = self._analysis_character_profile_text(item)
                if profile_text:
                    old = str(current.get("profile_text", "") or "").strip()
                    if old and profile_text not in old:
                        current["profile_text"] = f"{old}\n{profile_text}"
                    elif not old or len(profile_text) > len(old):
                        current["profile_text"] = profile_text
                for key, value in item.items():
                    if key == "name":
                        continue
                    if key in {"character_form", "profile_form", "profile_data", "profile_data_json"} and isinstance(value, dict):
                        old_form = current.get(key)
                        current[key] = merge_character_profile_form(
                            old_form if isinstance(old_form, dict) else None,
                            value,
                            name,
                            str(current.get("role", item.get("role", "")) or ""),
                        )
                        continue
                    text = str(value or "").strip()
                    if not text:
                        continue
                    old = str(current.get(key, "") or "").strip()
                    if key == content_key and old and text not in old:
                        current[key] = f"{old}\n{text}"
                    elif not old or len(text) > len(old):
                        current[key] = text
            merged[target_key] = list(existing.values())

        merge_named_list("characters", "profile_text")

        incoming_chapters = payload.get("chapter_updates", [])
        if not isinstance(incoming_chapters, list):
            return
        existing_updates: dict[int, dict[str, Any]] = {}
        for item in merged.get("chapter_updates", []):
            if not isinstance(item, dict):
                continue
            try:
                existing_updates[int(item.get("chapter_id"))] = item
            except (TypeError, ValueError):
                continue
        for item in incoming_chapters:
            if not isinstance(item, dict):
                continue
            try:
                chapter_id = int(item.get("chapter_id"))
            except (TypeError, ValueError):
                continue
            current = existing_updates.setdefault(chapter_id, {"chapter_id": chapter_id, "events": []})
            outline = str(item.get("outline", "") or "").strip()
            if outline and (not current.get("outline") or len(outline) > len(str(current.get("outline", "")))):
                current["outline"] = outline
            for key in ("ai_probability", "ai_probability_level"):
                if key in item:
                    current[key] = item[key]
            events = item.get("events", [])
            if isinstance(events, list):
                known_events = {str(event).strip() for event in current.get("events", [])}
                for event in events:
                    label = str(event).strip()
                    if label and label not in known_events:
                        current.setdefault("events", []).append(label)
                        known_events.add(label)
        merged["chapter_updates"] = list(existing_updates.values())

    def _apply_book_analysis_result(self, book_id: int, payload: dict[str, Any]) -> None:
        requires_evidence = self._book_analysis_requires_evidence(payload)
        payload = self._filter_untrusted_book_analysis_payload(payload)
        book = self.database.get_book(book_id)
        book_title = str(book["title"] or self.state.selected_book_title) if book else self.state.selected_book_title
        book_outline = str(payload.get("book_outline", "") or "").strip()
        if book:
            volume_specs = self._book_outline_volume_specs(book_id)
            outline_form = payload.get("book_outline_form")
            if isinstance(outline_form, dict):
                if requires_evidence:
                    outline_data = self._overlay_sparse_analysis_form(
                        self.database.get_book_outline_form(book_id),
                        outline_form,
                    )
                else:
                    outline_data = outline_form
                outline_data = normalize_book_outline_form(outline_data, book_title, volume_specs)
            elif book_outline:
                if "# 书籍大纲模板" in book_outline or "## 1." in book_outline:
                    outline_data = book_outline_markdown_to_form(book_outline, book_title, volume_specs)
                else:
                    outline_data = merge_book_outline_supplement(
                        book_outline_form_defaults(book_title, volume_specs),
                        book_title,
                        book_outline,
                        supplement_title=f"AI分析覆盖 · {datetime.now().strftime('%Y-%m-%d %H:%M')}",
                    )
            else:
                outline_data = self.database.get_book_outline_form(book_id)
            self.database.update_book_outline_form(book_id, normalize_book_outline_form(outline_data, book_title, volume_specs))

        existing_characters = {
            str(item["name"]).strip(): int(item["id"])
            for item in self.database.list_characters(book_id)
            if str(item["name"]).strip()
        }
        character_items = self._analysis_character_items(payload)
        for item in character_items:
            if not isinstance(item, dict):
                continue
            name = self._analysis_first_text(item, ("name", "character_name", "姓名", "角色名"))
            if not name:
                continue
            role = self._analysis_first_text(item, ("role", "character_role", "position", "定位", "角色定位"))
            profile = self._analysis_character_profile_text(item)
            character_id = existing_characters.get(name)
            if character_id is None:
                character_id = self.database.create_character(book_id, name)
                existing_characters[name] = character_id
            else:
                existing = self.database.get_character(character_id)
                if existing:
                    role = role or str(existing["role"] or "")
            profile_form = self._analysis_character_profile_form(item, role)
            if isinstance(profile_form, dict):
                if requires_evidence:
                    profile_data = self._overlay_sparse_analysis_form(
                        self.database.get_character_profile_form(character_id),
                        profile_form,
                    )
                else:
                    profile_data = profile_form
                profile_data = normalize_character_profile_form(profile_data, name, role)
            elif profile:
                profile_data = character_profile_markdown_to_form(profile, name, role)
            else:
                profile_data = character_profile_form_defaults(name, role)
            self.database.update_character_form(character_id, name, role, profile_data, image_path=None)

        world_form = self._analysis_world_form(payload)
        world_text = self._analysis_world_text(payload)
        world_entries = payload.get("world_entries", [])
        if isinstance(world_form, dict):
            if requires_evidence:
                world_data = self._overlay_sparse_analysis_form(self.database.get_book_world_form(book_id), world_form)
            else:
                world_data = world_form
            world_data = normalize_world_form(world_data)
        elif world_text or (isinstance(world_entries, list) and world_entries):
            combined_world_parts: list[str] = []
            if world_text:
                combined_world_parts.append(world_text)
            if isinstance(world_entries, list):
                for entry in world_entries:
                    if not isinstance(entry, dict):
                        continue
                    title = str(entry.get("name") or entry.get("title") or entry.get("category") or "世界观").strip()
                    content = str(entry.get("content") or entry.get("description") or "").strip()
                    if content:
                        combined_world_parts.append(f"## {title}\n{content}")
            world_data = world_markdown_to_form("\n\n".join(combined_world_parts))
        else:
            world_data = self.database.get_book_world_form(book_id)
        self.database.update_book_world_form(book_id, world_data)

        chapter_items = payload.get("chapter_updates", [])
        if not isinstance(chapter_items, list):
            chapter_items = []
        existing_events: dict[int, list[int]] = {}
        for event in self.database.list_chapter_events(book_id):
            existing_events.setdefault(int(event["chapter_id"]), []).append(int(event["id"]))
        for item in chapter_items:
            if not isinstance(item, dict):
                continue
            try:
                chapter_id = int(item.get("chapter_id"))
            except (TypeError, ValueError):
                continue
            chapter = self.database.get_chapter(chapter_id)
            if not chapter or int(chapter["book_id"]) != book_id:
                continue
            outline = str(item.get("outline", "")).strip()
            if outline:
                self.database.update_chapter(chapter_id, outline=outline, content=str(chapter["content"] or ""), summary_text="")
            probability = item.get("ai_probability")
            level = str(item.get("ai_probability_level", "none")).strip().lower()
            if probability is not None:
                try:
                    self.database.update_chapter_ai_probability(chapter_id, int(probability), level)
                except (TypeError, ValueError):
                    pass
            events = item.get("events", [])
            if isinstance(events, list):
                for event_id in existing_events.get(chapter_id, []):
                    self.database.delete_chapter_event(event_id)
                written: set[str] = set()
                for event_label in events[:6]:
                    label = str(event_label).strip()
                    if label and label not in written:
                        self.database.create_chapter_event(book_id, chapter_id, label)
                        written.add(label)

        self._sync_truth_files(book_id)
        self._index_book_design_for_rag(book_id)
        self._load_library_tree(self.current_tree_key)
        self._refresh_drawer()
        self._refresh_structured_markdown_views_after_analysis(book_id)
        if self.state.selected_chapter_id:
            self._load_chapter(self.state.selected_chapter_id)
        meta = payload.get("_meta", {}) if isinstance(payload.get("_meta", {}), dict) else {}
        processed = int(meta.get("processed_chapters", 0) or 0)
        unresolved_items = payload.get("analysis_unresolved", [])
        if isinstance(unresolved_items, list) and unresolved_items:
            self._append_task_log(f"全书分析待确认：{len(unresolved_items)} 项证据不足，未写入。")
        self._append_task_log(f"全书分析结果已写入作品库。处理 {processed or len(chapter_items)} 章。")
        self._switch_drawer_tab("overview")
        self._update_word_count()

    def _refresh_structured_markdown_views_after_analysis(self, book_id: int) -> None:
        if self.state.selected_book_id != book_id:
            return
        if self.state.editor_scope_kind == "book":
            book = self.database.get_book(book_id)
            if book:
                volume_specs = self._book_outline_volume_specs(book_id)
                outline_data = self.database.get_book_outline_form(book_id)
                outline = book_outline_form_to_markdown(outline_data, str(book["title"] or self.state.selected_book_title), volume_specs)
                self.loading_editor = True
                try:
                    self.book_outline_form.set_data(outline_data, str(book["title"] or self.state.selected_book_title), volume_specs)
                    self.outline_editor.setPlainText(outline)
                finally:
                    self.loading_editor = False
                self._refresh_outline_markdown_preview()
                if self.outline_markdown_stack.currentWidget() is self.outline_markdown_preview:
                    self._set_markdown_view_mode("outline", True)
                self._mark_clean()
        if self.selected_character_id:
            self._load_character_form(self.selected_character_id)
            if self.character_markdown_stack.currentWidget() is self.character_profile_preview:
                self._set_markdown_view_mode("character", True)
        if hasattr(self, "world_form"):
            self._load_world_form_for_book(book_id)
            if self.world_markdown_stack.currentWidget() is self.world_content_preview:
                self._set_markdown_view_mode("world", True)

    def _switch_drawer_tab(self, tab_key: str) -> None:
        index = self.drawer_tab_indexes.get(tab_key)
        if index is not None and hasattr(self, "drawer_tabs"):
            self.drawer_tabs.setCurrentIndex(index)

    def _review_current_scope(self) -> None:
        if not self.state.selected_book_id:
            info(self, "长篇校验", "请先选择一本书、一个卷或一个章节。")
            return
        ai_service = self._require_ai_service("review", "长篇校验")
        if ai_service is None:
            return
        review_ai_service = ai_service.clone()
        if self._find_active_background_job("长篇校验"):
            info(self, "任务运行中", "长篇校验已经在运行。")
            return
        if self.state.editor_dirty and not self._save_current():
            return
        targets = self._collect_review_targets()
        if not targets:
            info(self, "长篇校验", "当前范围没有可审查的章节。")
            return
        active_targets = [int(item["id"]) for item in targets if self._has_active_write_job(int(item["id"]))]
        if active_targets:
            info(self, "任务运行中", "审查范围内已有写入型 AI 任务在运行，请先停止或等待完成。")
            return

        self._open_drawer("tasks")
        self._append_task_log(f"长篇校验准备：{len(targets)} 章")
        prepared_targets: list[dict[str, Any]] = []
        for chapter in targets:
            chapter_id = int(chapter["id"])
            base_outline = str(chapter["outline"] or "")
            base_content = str(chapter["content"] or "")
            snapshot_id = self.database.create_snapshot(
                chapter_id,
                f"长篇校验前 · {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                base_outline,
                base_content,
            )
            prepared_targets.append(
                {
                    "chapter_id": chapter_id,
                    "book_id": int(chapter["book_id"]),
                    "title": str(chapter["title"]),
                    "outline": base_outline,
                    "content": base_content,
                    "snapshot_id": snapshot_id,
                    "truth_file": self._build_truth_context(int(chapter["book_id"]), chapter_id),
                    "template_context": self._build_template_context(int(chapter["book_id"]), chapter_id),
                    "skills": self._get_review_skill_payload(chapter_id),
                }
            )

        scope_type = self._review_scope_type()
        book_title = self.state.selected_book_title
        relationship_snapshot = self._prepare_relationship_analysis_snapshot(int(self.state.selected_book_id))

        def work(log: Callable[[str], None], cancel_event: Event | None) -> dict[str, Any]:
            results: list[dict[str, Any]] = []
            cancelled_targets: list[dict[str, Any]] = []
            relationship_analysis: dict[str, Any] | None = None
            for index, target in enumerate(prepared_targets, start=1):
                if cancel_event and cancel_event.is_set():
                    cancelled_targets.extend(prepared_targets[index - 1 :])
                    log("长篇校验已停止，后续章节不会写入。")
                    break
                log(f"长篇校验 {index}/{len(prepared_targets)}：{target['title']}")
                review = review_ai_service.multi_agent_review(
                    book_title=book_title,
                    chapter_title=str(target["title"]),
                    outline=str(target["outline"]),
                    content=str(target["content"]),
                    truth_file=str(target["truth_file"]),
                    selected_skills=list(target["skills"]),
                    require_remote=True,
                    cancel_event=cancel_event,
                    template_content=str(target["template_context"]),
                )
                if cancel_event and cancel_event.is_set():
                    cancelled_targets.extend(prepared_targets[index - 1 :])
                    log("长篇校验已停止，当前结果已丢弃。")
                    break
                if index > 1:
                    previous = prepared_targets[: index - 1]
                    prev_chapters_for_check = [
                        {
                            "title": str(p["title"]),
                            "content": str(p["content"]),
                            "outline": str(p["outline"]),
                        }
                        for p in previous
                    ]
                    cross_result = review_ai_service.cross_chapter_check(
                        current_chapter_title=str(target["title"]),
                        current_chapter_content=str(target["content"]),
                        previous_chapters=prev_chapters_for_check,
                        truth_file=str(target["truth_file"]),
                        require_remote=True,
                        cancel_event=cancel_event,
                    )
                    cross_findings = cross_result.get("cross_findings", [])
                    if cross_findings:
                        for f in cross_findings:
                            f["is_cross_chapter"] = True
                        log(f"跨章节一致性检查：发现 {len(cross_findings)} 个潜在不一致")
                    else:
                        log("跨章节一致性检查：未发现问题")
                    all_findings = list(review.get("findings", []))
                    all_findings.extend(cross_findings)
                    review["findings"] = all_findings
                    cross_summary = cross_result.get("cross_summary", "")
                    if cross_summary:
                        existing = str(review.get("summary", ""))
                        review["summary"] = f"{existing}\n\n跨章节一致性：{cross_summary}".strip()
                    if cross_summary:
                        review["cross_summary"] = cross_summary
                results.append({**target, "review": review})
                log(f"长篇校验进度：已完成 {len(results)}/{len(prepared_targets)}")
            if not (cancel_event and cancel_event.is_set()) and relationship_snapshot:
                characters = relationship_snapshot.get("characters", [])
                chapters = relationship_snapshot.get("chapters", [])
                if isinstance(characters, list) and len(characters) >= 2 and isinstance(chapters, list):
                    log("书籍级人际关系分析：正在刷新星图连线")
                    try:
                        relationships = review_ai_service.extract_character_relationships(
                            book_title=str(relationship_snapshot.get("book_title", "") or ""),
                            characters=characters,
                            chapters=chapters,
                            require_remote=False,
                            cancel_event=cancel_event,
                        )
                        relationship_analysis = {
                            "book_id": int(relationship_snapshot.get("book_id", 0) or 0),
                            "relationships": relationships,
                            "character_count": len(characters),
                        }
                        log(f"书籍级人际关系分析：识别 {len(relationships)} 条最终关系")
                    except Exception as exc:  # noqa: BLE001
                        if cancel_event and cancel_event.is_set():
                            log("书籍级人际关系分析已停止。")
                        else:
                            log(f"书籍级人际关系分析跳过：{str(exc)[:180]}")
            return {
                "scope_type": scope_type,
                "results": results,
                "cancelled": bool(cancel_event and cancel_event.is_set()),
                "cancelled_targets": cancelled_targets,
                "total": len(prepared_targets),
                "relationship_analysis": relationship_analysis,
            }

        def done(payload: Any) -> None:
            self._apply_multi_agent_review_result(payload)

        self._run_background_ai_job(
            label="长篇校验",
            callback=work,
            on_done=done,
            callback_uses_log=True,
            callback_uses_cancel=True,
            deliver_result_on_cancel=True,
        )

    def _review_scope_type(self) -> str:
        if self.state.current_node_kind == "chapter":
            return "chapter"
        if self.state.current_node_kind == "volume":
            return "volume"
        if self.state.current_node_kind == "group":
            return "group"
        return "book"

    def _collect_review_targets(self) -> list[dict[str, Any]]:
        book_id = self.state.selected_book_id
        if not book_id:
            return []
        if self.state.current_node_kind == "chapter" and self.state.selected_chapter_id:
            chapter = self.database.get_chapter(self.state.selected_chapter_id)
            return [dict(chapter)] if chapter else []
        if self.state.current_node_kind == "volume" and self.state.selected_volume_id:
            rows = self.database.list_chapters(book_id, self.state.selected_volume_id)
        elif self.state.current_node_kind == "group":
            rows = self.database.list_chapters(book_id, UNASSIGNED_VOLUMES)
        else:
            rows = self.database.list_chapters(book_id)
        targets: list[dict[str, Any]] = []
        for row in rows:
            chapter = self.database.get_chapter(int(row["id"]))
            if chapter:
                targets.append(dict(chapter))
        return targets

    def _prepare_relationship_analysis_snapshot(self, book_id: int) -> dict[str, Any] | None:
        characters = self._named_book_characters(book_id)
        if len(characters) < 2:
            return None
        chapters: list[dict[str, Any]] = []
        for row in self.database.list_chapters(book_id):
            chapter = self.database.get_chapter(int(row["id"]))
            if chapter:
                chapters.append(
                    {
                        "id": int(chapter["id"]),
                        "title": str(chapter["title"] or ""),
                        "outline": str(chapter["outline"] or ""),
                        "content": str(chapter["content"] or ""),
                    }
                )
        if not chapters:
            return None
        return {
            "book_id": book_id,
            "book_title": self.state.selected_book_title or "",
            "characters": [
                {
                    "id": int(item["id"]),
                    "name": str(item.get("name", "") or ""),
                    "role": str(item.get("role", "") or ""),
                    "profile_text": str(item.get("profile_text", "") or ""),
                }
                for item in characters
            ],
            "chapters": chapters,
        }

    def _get_review_skill_payload(self, chapter_id: int) -> list[dict[str, Any]]:
        seen: set[int] = set()
        payload: list[dict[str, Any]] = []
        for item in self.database.list_bound_skills_for_global():
            skill = dict(item)
            skill_id = int(skill["id"])
            if skill_id not in seen:
                payload.append(skill)
                seen.add(skill_id)
        if self.state.selected_book_id:
            for item in self.database.list_bound_skills_for_book(self.state.selected_book_id):
                skill = dict(item)
                skill_id = int(skill["id"])
                if skill_id not in seen:
                    payload.append(skill)
                    seen.add(skill_id)
        for item in self.database.list_bound_skills_for_chapter(chapter_id):
            skill = dict(item)
            skill_id = int(skill["id"])
            if skill_id not in seen:
                payload.append(skill)
                seen.add(skill_id)
        return payload

    def _build_template_context(self, book_id: int, exclude_chapter_id: int | None = None) -> str:
        templates = self.database.list_template_chapters(book_id)
        lines: list[str] = []
        for row in templates:
            chapter_id = int(row["id"])
            if exclude_chapter_id is not None and chapter_id == exclude_chapter_id:
                continue
            content = str(row["content"] or "")
            stats = compute_template_stats(content)
            lines.extend(
                [
                    f"--- 模板：{row['title']} ---",
                    f"卷：{row['volume_title']}",
                    f"字数：{stats['word_count']}",
                    f"对话密度：{stats['dialogue_density']}%",
                    f"描写比例：{stats['description_ratio']}%",
                    f"正文节选：{content[:700]}",
                ]
            )
        return "\n".join(lines)[:3000]

    def _build_review_truth_file(self, chapter: dict[str, Any], scope_chapters: list[dict[str, Any]]) -> str:
        book_id = int(chapter["book_id"])
        try:
            export_data = self.database.get_book_export_data(book_id)
        except Exception:  # noqa: BLE001
            export_data = {}
        book = export_data.get("book", {}) if isinstance(export_data, dict) else {}
        lines = [
            "Truth File / 长篇生成校验上下文",
            f"书名：{book.get('title') or self.state.selected_book_title or ''}",
            f"当前章节：{chapter.get('title', '')}",
            f"当前卷：{chapter.get('volume_title', '')}",
        ]
        book_outline = str(book.get("outline_text", "") or "").strip()
        if book_outline:
            lines.append(f"书籍大纲：{book_outline[:1200]}")
        chapter_outline = str(chapter.get("outline", "") or "").strip()
        if chapter_outline:
            lines.append(f"章节大纲：{chapter_outline[:1200]}")

        chapters_all = export_data.get("chapters", []) if isinstance(export_data, dict) else []
        chapter_ids = [int(item.get("id", -1)) for item in chapters_all if isinstance(item, dict)]
        try:
            current_index = chapter_ids.index(int(chapter["id"]))
        except ValueError:
            current_index = -1
        if current_index >= 0:
            nearby = chapters_all[max(0, current_index - 2) : current_index] + chapters_all[current_index + 1 : current_index + 3]
            if nearby:
                lines.append("相邻章节：")
                for item in nearby:
                    title = str(item.get("title", "")).strip()
                    outline = str(item.get("outline", "") or item.get("summary_text", "")).strip()
                    content = str(item.get("content", "")).strip()
                    snippet = outline or content[:180]
                    lines.append(f"- {title}：{snippet[:260]}")

        characters = export_data.get("characters", []) if isinstance(export_data, dict) else []
        if characters:
            lines.append("人物卡：")
            for item in characters[:40]:
                name = str(item.get("name", "")).strip()
                role = str(item.get("role", "")).strip()
                profile = str(item.get("profile_text", "")).strip()
                lines.append(f"- {name} / {role}：{profile[:260]}")

        relationships = [dict(item) for item in self.database.list_relationships(book_id)]
        if relationships:
            lines.append("人物关系：")
            for item in relationships[:60]:
                source = str(item.get("source_name", "")).strip()
                target = str(item.get("target_name", "")).strip()
                relation = str(item.get("relationship_type", "")).strip() or "关系"
                description = str(item.get("description", "")).strip()
                lines.append(f"- {source} -> {target}：{relation} {description[:160]}")

        world_entries = export_data.get("world_entries", []) if isinstance(export_data, dict) else []
        if world_entries:
            lines.append("世界观：")
            for item in world_entries[:40]:
                name = str(item.get("name", "")).strip()
                category = str(item.get("category", "")).strip()
                content = str(item.get("content_text", "")).strip()
                lines.append(f"- {name} [{category}]：{content[:260]}")

        events = [dict(item) for item in self.database.list_chapter_events(book_id, int(chapter["id"]))]
        if events:
            lines.append("本章关键事件：")
            for item in events[:20]:
                label = str(item.get("label", "")).strip()
                description = str(item.get("description", "")).strip()
                lines.append(f"- {label}：{description[:180]}")

        skills = self._get_review_skill_payload(int(chapter["id"]))
        if skills:
            lines.append("启用 Skills：")
            for item in skills[:16]:
                lines.append(f"- {item.get('name', '')}：{str(item.get('instruction_text', '')).strip()[:220]}")

        spans = self.database.list_chapter_ai_spans(int(chapter["id"]))
        total = max(1, len(str(chapter.get("content", "") or "")))
        marked = sum(max(0, int(span["end_offset"]) - int(span["start_offset"])) for span in spans)
        lines.append(f"本章 AI 标记占比：{int(round(marked * 100 / total))}%")
        template_context = self._build_template_context(book_id, int(chapter["id"]))
        if template_context:
            lines.append("【对照模板章节】")
            lines.append(template_context)
        lines.append("生成校验机制来源：借鉴多智能体长篇写作工作流，仅做机制级学习，不复用外部源码。")
        return "\n".join(line for line in lines if str(line).strip())[:14000]

    def _build_truth_context(
        self,
        book_id: int,
        chapter_id: int | None = None,
        max_chars: int = 14000,
    ) -> str:
        try:
            return self.truth_manager.assemble_context(book_id, chapter_id, max_chars=max_chars)
        except Exception:
            return ""

    def _sync_truth_files(self, book_id: int, chapter_id: int | None = None) -> None:
        try:
            self.truth_manager.build_all(book_id)
        except Exception:
            pass

    def _apply_multi_agent_review_result(self, payload: Any) -> None:
        if not isinstance(payload, dict):
            return
        results = payload.get("results", [])
        if not isinstance(results, list):
            return
        source_label = str(payload.get("source_label", "长篇生成") or "长篇生成")
        cancelled_targets = payload.get("cancelled_targets", [])
        if not isinstance(cancelled_targets, list):
            cancelled_targets = []
        scope_type = str(payload.get("scope_type", "chapter") or "chapter")
        applied = 0
        rejected = 0
        skipped = 0
        cancelled = 0
        selected_chapter_id = self.state.selected_chapter_id
        recorded_chapter_ids: set[int] = set()
        for item in results:
            if not isinstance(item, dict):
                continue
            chapter_id = int(item.get("chapter_id", 0) or 0)
            book_id = int(item.get("book_id", self.state.selected_book_id or 0) or 0)
            review = item.get("review", {})
            if not chapter_id or not book_id or not isinstance(review, dict):
                continue
            recorded_chapter_ids.add(chapter_id)
            final_verdict = str(review.get("final_verdict", "reject")).strip().lower()
            revised_content = str(review.get("revised_content", "") or "")
            revised_outline = str(review.get("revised_outline", "") or item.get("outline", "") or "")
            base_content = str(item.get("content", "") or "")
            unsafe_reason = ""
            if final_verdict == "approve":
                if not revised_content.strip():
                    unsafe_reason = "模型返回的修订正文为空"
                elif len(base_content.strip()) >= 200 and len(revised_content.strip()) < int(len(base_content.strip()) * 0.4):
                    unsafe_reason = "模型返回的修订正文长度异常缩短"
            current = self.database.get_chapter(chapter_id)
            stale = True
            if current:
                stale = (
                    str(current["content"] or "") != str(item.get("content", ""))
                    or str(current["outline"] or "") != str(item.get("outline", ""))
                    or (selected_chapter_id == chapter_id and self.state.editor_dirty)
                )
            status = "rejected"
            if final_verdict == "approve":
                if unsafe_reason:
                    status = "unsafe_skipped"
                elif stale:
                    status = "stale_skipped"
                else:
                    status = "applied"
            findings = review.get("findings", [])
            findings_list = findings if isinstance(findings, list) else []
            if unsafe_reason:
                findings_list = [
                    {
                        "agent": "safety_gate",
                        "severity": "high",
                        "category": "自动覆盖护栏",
                        "location_hint": "revised_content",
                        "quote": "",
                        "issue": unsafe_reason,
                        "suggestion": "本次审查只保留记录，不自动覆盖正文；请人工比对后再决定是否恢复或手动应用。",
                    },
                    *findings_list,
                ]
            run_id = self.database.create_review_run(
                book_id=book_id,
                chapter_id=chapter_id,
                scope_type=scope_type,
                status=status,
                truth_snapshot=str(item.get("truth_file", "")),
                summary=str(review.get("summary", "")),
                overall_score=int(review.get("overall_score", 0) or 0),
                snapshot_id=int(item.get("snapshot_id", 0) or 0) or None,
                revised_content=revised_content,
                revised_outline=revised_outline,
                final_verdict=final_verdict,
                template_comparison=str(review.get("template_comparison", "")),
                risk_notes=str(review.get("risk_notes", "")),
                cross_summary=str(review.get("cross_summary", "")),
                applied=status == "applied",
            )
            self.database.replace_review_findings(run_id, findings_list)
            if final_verdict != "approve":
                rejected += 1
                self._append_task_log(f"章节 {item.get('title', chapter_id)} 审查未通过，未覆盖正文。")
                continue
            if unsafe_reason:
                skipped += 1
                self._append_task_log(f"章节 {item.get('title', chapter_id)} 未覆盖：{unsafe_reason}。")
                continue
            if stale:
                skipped += 1
                self._append_task_log(f"章节 {item.get('title', chapter_id)} 已被修改，自动覆盖已跳过。")
                continue
            self.database.update_chapter(chapter_id, outline=revised_outline, content=revised_content, summary_text="")
            self.database.delete_chapter_ai_spans(chapter_id)
            self.database.add_chapter_ai_span(chapter_id, 0, len(revised_content), f"review-{run_id}")
            self.database.update_chapter_ai_probability(chapter_id, 100, "certain")
            self._sync_truth_files(book_id, chapter_id)
            self._index_chapter_for_rag(book_id, chapter_id)
            self._record_input_stat("ai", len(revised_content.strip()))
            applied += 1
            if selected_chapter_id == chapter_id:
                self.loading_editor = True
                self.suppress_ai_span_clear = True
                try:
                    self.chapter_outline_editor.setPlainText(revised_outline)
                    self.content_editor.setPlainText(revised_content)
                finally:
                    self.suppress_ai_span_clear = False
                    self.loading_editor = False
                self._refresh_ai_spans_for_current_chapter()
                self._update_ai_probability_badge("certain", 100)
                self._mark_clean()
        for item in cancelled_targets:
            if not isinstance(item, dict):
                continue
            chapter_id = int(item.get("chapter_id", 0) or 0)
            book_id = int(item.get("book_id", self.state.selected_book_id or 0) or 0)
            if not chapter_id or not book_id or chapter_id in recorded_chapter_ids:
                continue
            self.database.create_review_run(
                book_id=book_id,
                chapter_id=chapter_id,
                scope_type=scope_type,
                status="cancelled",
                truth_snapshot=str(item.get("truth_file", "")),
                summary="长篇生成任务已取消，本章未进行自动修订。",
                overall_score=0,
                snapshot_id=int(item.get("snapshot_id", 0) or 0) or None,
                revised_content=str(item.get("content", "") or ""),
                revised_outline=str(item.get("outline", "") or ""),
                final_verdict="cancelled",
                template_comparison="",
                risk_notes="任务取消，正文未覆盖。",
                applied=False,
            )
            cancelled += 1
        self._apply_ai_relationship_analysis(payload.get("relationship_analysis"))
        activate_chapter_id = int(payload.get("activate_chapter_id", 0) or 0)
        if activate_chapter_id:
            self._load_library_tree(("chapter", activate_chapter_id, None))
            self._activate_tree_key(("chapter", activate_chapter_id, None))
        else:
            self._load_library_tree(self.current_tree_key)
        self._refresh_drawer()
        if results or cancelled:
            self._open_drawer("review")
        cancel_note = f"，取消 {cancelled} 章" if cancelled else ""
        self._append_task_log(f"{source_label}写回完成：覆盖 {applied} 章，拒绝 {rejected} 章，跳过 {skipped} 章{cancel_note}。")

    def _create_chapter_after_current(
        self,
        book_id: int,
        after_chapter_id: int,
        title: str,
        content: str = "",
        outline: str = "",
    ) -> int | None:
        try:
            chapter_id = self.database.create_chapter_after(
                after_chapter_id=after_chapter_id,
                title=title,
                outline=outline,
                content=content,
            )
        except Exception as exc:  # noqa: BLE001
            self._append_task_log(f"创建插入章节失败：{exc}")
            return None
        self._load_library_tree(("chapter", chapter_id, None))
        self._activate_tree_key(("chapter", chapter_id, None))
        return chapter_id

    def _open_ai_correction_dialog(self, ai_output: str, context_prompt: str = "") -> None:
        payload = {
            "ai_output": ai_output,
            "context_prompt": context_prompt,
        }
        self._open_drawer("correction", payload=payload)

    def _apply_ai_relationship_analysis(self, payload: Any) -> None:
        if not isinstance(payload, dict):
            return
        book_id = int(payload.get("book_id", self.state.selected_book_id or 0) or 0)
        raw_relationships = payload.get("relationships", [])
        if not book_id or not isinstance(raw_relationships, list):
            return
        named_characters = self._named_book_characters(book_id)
        if len(named_characters) < 2:
            return
        id_by_name = {
            self._relationship_lookup_key(item.get("name", "")): int(item["id"])
            for item in named_characters
            if str(item.get("name", "") or "").strip()
        }
        normalized: list[dict[str, Any]] = []
        seen_pairs: set[tuple[int, int]] = set()
        for item in raw_relationships:
            if not isinstance(item, dict):
                continue
            source_id = id_by_name.get(self._relationship_lookup_key(item.get("source_name") or item.get("source") or ""))
            target_id = id_by_name.get(self._relationship_lookup_key(item.get("target_name") or item.get("target") or ""))
            if not source_id or not target_id or source_id == target_id:
                continue
            pair = tuple(sorted((source_id, target_id)))
            if pair in seen_pairs:
                continue
            seen_pairs.add(pair)
            normalized.append(
                {
                    "source_character_id": source_id,
                    "target_character_id": target_id,
                    "relationship_type": str(
                        item.get("relationship_type") or item.get("type") or item.get("relation") or "关联"
                    ).strip()[:32],
                    "description": str(item.get("description") or item.get("evidence") or "").strip()[:240],
                }
            )
        try:
            self.database.replace_book_relationships(book_id, normalized)
        except Exception as exc:  # noqa: BLE001
            self._append_task_log(f"星图关系写入失败：{str(exc)[:180]}")
            return
        self._append_task_log(
            f"星图关系已刷新：{len(normalized)} 条，使用 {len(named_characters)} 个具名人物。"
        )

    @staticmethod
    def _relationship_lookup_key(value: Any) -> str:
        text = str(value or "").strip()
        text = re.sub(r"[（(].*?[）)]", "", text)
        text = re.sub(r"\s+", "", text)
        return text.casefold()

    def _merge_analysis_outline(self, existing: str, analyzed: str) -> str:
        analyzed = analyzed.strip()
        if not analyzed:
            return existing.strip()
        existing = existing.strip()
        if not existing:
            return analyzed
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        return f"{existing}\n\n【AI 全书分析 · {stamp}】\n{analyzed}"

    def _run_ai_worker(
        self,
        ai_service: SimpleAIService,
        mode: str,
        target_chapter_id: int | None,
        target: str,
        stream_mode: str,
        chapter_title: str,
        outline: str,
        current_content: str,
        summary_text: str,
        base_content: str = "",
        base_outline: str = "",
        buffer_prefix: str = "",
    ) -> None:
        job_id = uuid.uuid4().hex
        cancel_event = Event()
        thread = QThread(self)
        worker = AiWorker(
            ai_service,
            mode=mode,
            book_title=self.state.selected_book_title,
            chapter_title=chapter_title,
            outline=outline,
            current_content=current_content,
            summary_text=summary_text,
            selected_skills=self._get_bound_skill_payload(),
            sourcebook_context=self._build_sourcebook_context(),
            cancel_event=cancel_event,
        )
        self.ai_jobs[job_id] = {
            "id": job_id,
            "thread": thread,
            "worker": worker,
            "cancel_event": cancel_event,
            "mode": mode,
            "target": target,
            "target_chapter_id": target_chapter_id,
            "result_text": "",
            "base_content": base_content,
            "base_outline": base_outline,
            "buffer_prefix": buffer_prefix,
            "prefix_pending": buffer_prefix,
            "stream_mode": stream_mode,
            "label": mode,
            "cancel_requested": False,
            "started_at": datetime.now(),
        }
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.log.connect(lambda message, jid=job_id: self._on_ai_log(jid, message))
        worker.chunk.connect(lambda chunk, jid=job_id: self._on_ai_chunk(jid, chunk))
        worker.done.connect(lambda message, jid=job_id: self._on_ai_done(jid, message))
        worker.cancelled.connect(lambda message, jid=job_id: self._on_ai_done(jid, message))
        worker.error.connect(lambda message, jid=job_id: self._on_ai_error(jid, message))
        worker.done.connect(thread.quit)
        worker.cancelled.connect(thread.quit)
        worker.error.connect(thread.quit)
        thread.finished.connect(worker.deleteLater)
        thread.finished.connect(lambda jid=job_id: self._cleanup_ai_thread(jid))
        self._update_task_activity_state()
        thread.start()

    def _has_active_write_job(self, chapter_id: int, targets: set[str] | None = None) -> bool:
        active_targets = targets or {"content"}
        for job in self.ai_jobs.values():
            if (
                not job.get("cancel_requested")
                and job.get("target_chapter_id") == chapter_id
                and job.get("target") in active_targets
            ):
                return True
        return False

    def _on_ai_log(self, job_id: str, message: str) -> None:
        if job_id not in self.ai_jobs:
            return
        self._append_task_log(message)

    def _on_ai_chunk(self, job_id: str, chunk: str) -> None:
        job = self.ai_jobs.get(job_id)
        if not job:
            return
        target = str(job.get("target", "log"))
        job["result_text"] = str(job.get("result_text", "")) + chunk
        if target == "content":
            if self._is_ai_target_selected(job):
                self.ai_writing = True
                self.suppress_ai_span_clear = True
                try:
                    prefix = str(job.get("prefix_pending", ""))
                    if prefix:
                        self.content_editor.appendPlainText(prefix.strip("\n"))
                        job["prefix_pending"] = ""
                    cursor = self.content_editor.textCursor()
                    cursor.movePosition(cursor.MoveOperation.End)
                    cursor.insertText(chunk)
                    self.content_editor.setTextCursor(cursor)
                finally:
                    self.ai_writing = False
                    self.suppress_ai_span_clear = False
                self.state.editor_dirty = True
                self.save_label.setText("未保存")
            else:
                self._record_input_stat("ai", len(chunk.strip()))
        elif target == "outline":
            self.loading_editor = True
            try:
                if self._is_ai_target_selected(job):
                    self.chapter_outline_editor.setPlainText(str(job.get("result_text", "")))
            finally:
                self.loading_editor = False
            if self._is_ai_target_selected(job):
                self.editor_text_lengths[id(self.chapter_outline_editor)] = len(self.chapter_outline_editor.toPlainText())
                self.state.editor_dirty = True
                self.save_label.setText("未保存")
            self._record_input_stat("ai", len(chunk.strip()))
        else:
            self._record_input_stat("ai", len(chunk.strip()))
            self._append_task_log(chunk, inline=True)

    def _append_task_log(self, text: str, inline: bool = False) -> None:
        text = str(text)
        if inline:
            cursor = self.task_log.textCursor()
            cursor.movePosition(cursor.MoveOperation.End)
            if not self.task_log_streaming and self.task_log.toPlainText().strip():
                cursor.insertText("\n\n")
            cursor.insertText(text)
            self.task_log.setTextCursor(cursor)
            self.task_log.ensureCursorVisible()
            self.task_log_streaming = True
            return
        if text == self.last_task_log_line:
            return
        self.task_log.appendPlainText(text)
        self.last_task_log_line = text
        self.task_log_streaming = False

    def _on_ai_done(self, job_id: str, message: str) -> None:
        job = self.ai_jobs.get(job_id)
        if not job:
            return
        if job.get("cancel_requested"):
            if str(job.get("target", "")) in {"content", "outline"} and self._is_ai_target_selected(job):
                self.state.editor_dirty = True
                self.save_label.setText("未保存")
            self._append_task_log(message)
            self._update_task_activity_state()
            self._set_status("AI 任务已停止")
            return
        mode = str(job.get("mode", ""))
        if mode == "polish":
            self._show_polish_diff_dialog(job)
            self._append_task_log(message)
            self._update_task_activity_state()
            return
        self._persist_ai_result_to_target(job)
        self._append_task_log(message)
        self._update_task_activity_state()
        self._set_status(message)
        target = str(job.get("target", ""))
        if target == "outline":
            self._switch_drawer_tab("outline")

    def _is_ai_target_selected(self, job: dict[str, Any]) -> bool:
        target_chapter_id = job.get("target_chapter_id")
        return bool(target_chapter_id and self.state.selected_chapter_id == target_chapter_id)

    def _persist_ai_result_to_target(self, job: dict[str, Any]) -> None:
        target = str(job.get("target", "log"))
        target_chapter_id = job.get("target_chapter_id")
        if target not in {"content", "outline"} or not target_chapter_id:
            return
        chapter = self.database.get_chapter(int(target_chapter_id))
        if not chapter:
            return
        selected = self._is_ai_target_selected(job)
        result_text = str(job.get("result_text", ""))
        if target == "outline":
            new_outline = self.chapter_outline_editor.toPlainText() if selected else result_text
            if not new_outline.strip():
                return
            self.database.update_chapter(
                int(target_chapter_id),
                outline=new_outline,
                content=str(chapter["content"] or ""),
                summary_text="",
            )
        else:
            if selected:
                new_content = self.content_editor.toPlainText()
            elif str(job.get("stream_mode", "")) == "append":
                new_content = f"{job.get('base_content', '')}{job.get('buffer_prefix', '')}{result_text}"
            else:
                new_content = result_text
            self.database.update_chapter(
                int(target_chapter_id),
                outline=str(chapter["outline"] or ""),
                content=new_content,
                summary_text="",
            )
            self.database.update_chapter_ai_probability(int(target_chapter_id), 100, "certain")
            if str(job.get("stream_mode", "")) == "append":
                span_start = len(str(job.get("base_content", "")) + str(job.get("buffer_prefix", "")))
            else:
                self.database.delete_chapter_ai_spans(int(target_chapter_id))
                span_start = 0
            self.database.add_chapter_ai_span(
                int(target_chapter_id),
                span_start,
                len(new_content),
                str(job.get("id", "")),
            )
        if selected:
            if target == "outline":
                self.loaded_chapter_summary_text = ""
            elif target == "content":
                self.loaded_chapter_content = self.content_editor.toPlainText()
                self.loaded_chapter_summary_text = ""
            if target == "content":
                self._update_ai_probability_badge("certain", 100)
                self._load_library_tree(self.current_tree_key)
                self._refresh_ai_spans_for_current_chapter()
            self._mark_clean()
        else:
            self._append_task_log(f"结果已写回原章节 ID {target_chapter_id}。")
            self._load_library_tree(self.current_tree_key)

    def _on_ai_error(self, job_id: str, message: str) -> None:
        job = self.ai_jobs.get(job_id)
        if not job:
            return
        if job.get("cancel_requested"):
            self._append_task_log("AI 任务已取消，后台连接正在退出。")
            self._update_task_activity_state()
            self._set_status("AI 任务已停止")
            return
        self._append_task_log(self._friendly_ai_error(message))
        self._update_task_activity_state()
        self._set_status("AI 任务失败")

    def _cleanup_ai_thread(self, job_id: str) -> None:
        self.ai_jobs.pop(job_id, None)
        self.detached_cancelled_jobs.pop(job_id, None)
        self._update_task_activity_state()
        self._complete_close_when_tasks_stop()

    def _cancel_task(self) -> None:
        if not self.ai_jobs:
            self._update_task_activity_state()
            self._set_status("当前没有正在运行的任务")
            return
        requested = 0
        already_stopping = 0
        for job in list(self.ai_jobs.values()):
            if job.get("cancel_requested"):
                already_stopping += 1
            else:
                requested += 1
                job["cancel_requested_at"] = datetime.now()
            job["cancel_requested"] = True
            cancel_event = job.get("cancel_event")
            if hasattr(cancel_event, "set"):
                cancel_event.set()
            thread = job.get("thread")
            if isinstance(thread, QThread):
                thread.requestInterruption()
                thread.quit()
        if requested:
            self._append_task_log(f"已请求停止 {requested} 个任务，正在等待后台线程安全退出。")
        elif already_stopping:
            self._append_task_log("停止请求已发送，后台任务仍在退出中。")
        self._update_task_activity_state()
        self.cancel_grace_timer.start(5000)
        self._set_status("正在停止任务...")

    def _update_task_activity_state(self) -> None:
        active = len(self.ai_jobs)
        stopping = sum(1 for job in self.ai_jobs.values() if job.get("cancel_requested"))
        running = active - stopping
        self.state.task_running = active > 0
        if active:
            if stopping and running:
                self.task_state.setText(f"运行中 · {running} / 停止中 · {stopping}")
            elif stopping:
                self.task_state.setText(f"正在停止 · {stopping}")
            else:
                self.task_state.setText(f"运行中 · {active}")
            self.task_progress.setRange(0, 0)
        else:
            self.task_state.setText("空闲")
            self.task_progress.setRange(0, 1)
            self.task_progress.setValue(0)
            if self.cancel_grace_timer.isActive():
                self.cancel_grace_timer.stop()
        if hasattr(self, "cancel_task_button"):
            self.cancel_task_button.setEnabled(running > 0)
            self.cancel_task_button.setText("停止中..." if stopping and not running else "停止任务")
        self._set_task_summary_expanded(False, animated=False)

    def _render_task_summary(self) -> None:
        if not hasattr(self, "task_summary_badge"):
            return
        active_jobs = list(self.ai_jobs.values())
        stopping = sum(1 for job in active_jobs if job.get("cancel_requested"))
        running = len(active_jobs) - stopping
        if not active_jobs:
            self.task_summary_badge.setText("空闲")
            self.task_summary_detail.setText("暂无运行任务")
            return
        if stopping and running:
            self.task_summary_badge.setText(f"{running} 运行 / {stopping} 停止中")
        elif stopping:
            self.task_summary_badge.setText(f"{stopping} 停止中")
        else:
            self.task_summary_badge.setText(f"{running} 运行中")
        lines: list[str] = []
        for job in active_jobs[:4]:
            state = "停止中" if job.get("cancel_requested") else "运行中"
            label = str(job.get("label") or job.get("mode") or "AI 任务")
            lines.append(f"{state} · {label}")
        if len(active_jobs) > 4:
            lines.append(f"还有 {len(active_jobs) - 4} 个任务")
        self.task_summary_detail.setText("\n".join(lines))
        if self.task_summary_expanded and self.task_summary_animation is None:
            self.task_summary_panel.setMaximumHeight(self._task_summary_target_height())

    def _task_summary_target_height(self) -> int:
        if not hasattr(self, "task_summary_panel"):
            return 0
        return max(70, min(150, self.task_summary_panel.sizeHint().height()))

    def _set_task_summary_expanded(self, expanded: bool, *, animated: bool) -> None:
        if not hasattr(self, "task_summary_panel"):
            return
        panel = self.task_summary_panel
        target_height = self._task_summary_target_height() if expanded else 0
        if self.task_summary_animation is not None:
            self.task_summary_animation.stop()
            self.task_summary_animation = None
        if self.task_summary_expanded == expanded and panel.isVisible() == expanded:
            if expanded:
                panel.setMaximumHeight(target_height)
            return
        self.task_summary_expanded = expanded
        if expanded:
            panel.setVisible(True)
        if not animated:
            panel.setMaximumHeight(target_height)
            panel.setVisible(expanded)
            return
        animation = QPropertyAnimation(panel, b"maximumHeight", self)
        animation.setDuration(220)
        animation.setEasingCurve(QEasingCurve.Type.OutCubic if expanded else QEasingCurve.Type.InCubic)
        animation.setStartValue(panel.maximumHeight())
        animation.setEndValue(target_height)

        def finish() -> None:
            if not self.task_summary_expanded:
                panel.setVisible(False)
            if self.task_summary_animation is animation:
                self.task_summary_animation = None

        animation.finished.connect(finish)
        self.task_summary_animation = animation
        animation.start()

    def _finalize_stale_cancelled_jobs(self) -> None:
        stale_jobs = [
            (job_id, job)
            for job_id, job in list(self.ai_jobs.items())
            if job.get("cancel_requested")
        ]
        if not stale_jobs:
            self._update_task_activity_state()
            return
        detached = 0
        for job_id, job in stale_jobs:
            thread = job.get("thread")
            if isinstance(thread, QThread) and thread.isRunning():
                self.detached_cancelled_jobs[job_id] = job
                thread.finished.connect(lambda jid=job_id: self._cleanup_detached_ai_thread(jid))
                thread.requestInterruption()
                thread.quit()
                detached += 1
            self.ai_jobs.pop(job_id, None)
        if detached:
            self._append_task_log(
                f"{detached} 个停止中的任务退出较慢，已从任务面板移出；后台返回结果将被忽略。"
            )
        self._update_task_activity_state()
        self._set_status("任务已停止")

    def _cleanup_detached_ai_thread(self, job_id: str) -> None:
        self.detached_cancelled_jobs.pop(job_id, None)
        self._complete_close_when_tasks_stop()

    def _running_ai_threads(self) -> list[QThread]:
        threads: list[QThread] = []
        seen: set[int] = set()
        for job in [*self.ai_jobs.values(), *self.detached_cancelled_jobs.values()]:
            thread = job.get("thread")
            if isinstance(thread, QThread) and thread.isRunning() and id(thread) not in seen:
                threads.append(thread)
                seen.add(id(thread))
        return threads

    def _request_all_ai_threads_stop(self) -> None:
        for job in [*self.ai_jobs.values(), *self.detached_cancelled_jobs.values()]:
            job["cancel_requested"] = True
            cancel_event = job.get("cancel_event")
            if hasattr(cancel_event, "set"):
                cancel_event.set()
            thread = job.get("thread")
            if isinstance(thread, QThread):
                thread.requestInterruption()
                thread.quit()

    def _complete_close_when_tasks_stop(self) -> None:
        if not self.close_after_task_stop:
            return
        self._request_all_ai_threads_stop()
        if self._running_ai_threads():
            self._set_status("正在等待后台 AI 任务退出，完成后将关闭。")
            if not self.close_poll_timer.isActive():
                self.close_poll_timer.start()
            return
        self.close_poll_timer.stop()
        self.close_after_task_stop = False
        self._set_status("后台任务已退出，正在关闭。")
        self.close()

    def _find_active_background_job(self, label: str) -> str | None:
        for job_id, job in self.ai_jobs.items():
            if (
                not job.get("cancel_requested")
                and job.get("target") == "background"
                and job.get("label") == label
            ):
                return job_id
        return None

    def _friendly_ai_error(self, message: str) -> str:
        text = str(message).strip()
        if "WinError 10061" in text:
            return "AI 接口连接被拒绝。请检查本地模型/API 服务是否已启动，或在 AI 设置里换成可用接口。"
        if "WinError 10054" in text:
            return "AI 接口连接被远程主机关闭。请稍后重试，或检查模型服务是否稳定。"
        if "timed out" in text.lower() or "timeout" in text.lower():
            return "AI 接口响应超时。请稍后重试，或检查网络/模型服务。"
        return text[:800]

    def _run_background_ai_job(
        self,
        *,
        label: str,
        callback: Callable[..., Any],
        on_done: Callable[[Any], None],
        on_error: Callable[[str], None] | None = None,
        callback_uses_log: bool = False,
        callback_uses_cancel: bool = False,
        deliver_result_on_cancel: bool = False,
    ) -> str:
        existing_job_id = self._find_active_background_job(label)
        if existing_job_id:
            self._append_task_log(f"{label}已在运行，忽略重复启动。")
            return existing_job_id
        job_id = uuid.uuid4().hex
        cancel_event = Event()
        thread = QThread(self)
        worker = FunctionWorker(
            callback,
            pass_logger=callback_uses_log,
            pass_cancel=callback_uses_cancel,
            cancel_event=cancel_event,
            deliver_result_on_cancel=deliver_result_on_cancel,
        )
        self.ai_jobs[job_id] = {
            "id": job_id,
            "thread": thread,
            "worker": worker,
            "cancel_event": cancel_event,
            "target": "background",
            "label": label,
            "cancel_requested": False,
            "started_at": datetime.now(),
        }

        def safe_on_done(payload: Any) -> bool:
            try:
                on_done(payload)
                return True
            except Exception as exc:  # noqa: BLE001
                friendly_message = self._friendly_ai_error(str(exc))
                self._append_task_log(f"{label}写回失败：{friendly_message}")
                self._set_status(f"{label}写回失败")
                if on_error is not None:
                    on_error(friendly_message)
                return False

        def handle_done(payload: Any, jid: str = job_id) -> None:
            job = self.ai_jobs.get(jid)
            if job is None:
                return
            if cancel_event.is_set() or (job and job.get("cancel_requested")):
                if deliver_result_on_cancel:
                    self._append_task_log(f"{label}已取消，正在写入已完成记录与取消记录。")
                    if safe_on_done(payload):
                        self._set_status(f"{label}已取消")
                    return
                self._append_task_log(f"{label}已取消，结果未写入。")
                self._set_status(f"{label}已取消")
                return
            self._append_task_log(f"{label}完成")
            if safe_on_done(payload):
                self._set_status(f"{label}完成")

        def handle_cancelled(message: str, jid: str = job_id) -> None:
            if jid not in self.ai_jobs:
                return
            self._append_task_log(str(message or f"{label}已取消。"))
            self._set_status(f"{label}已取消")

        def handle_error(message: str, jid: str = job_id) -> None:
            job = self.ai_jobs.get(jid)
            if job is None:
                return
            if job.get("cancel_requested"):
                self._append_task_log(f"{label}已取消，后台连接正在退出。")
                self._set_status(f"{label}已取消")
                return
            friendly_message = self._friendly_ai_error(message)
            self._append_task_log(f"{label}失败：{friendly_message}")
            self._set_status(f"{label}失败")
            if on_error is not None:
                on_error(friendly_message)

        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.log.connect(lambda message, jid=job_id: self._on_ai_log(jid, message))
        worker.done.connect(handle_done)
        worker.cancelled.connect(handle_cancelled)
        worker.error.connect(handle_error)
        worker.done.connect(thread.quit)
        worker.cancelled.connect(thread.quit)
        worker.error.connect(thread.quit)
        thread.finished.connect(worker.deleteLater)
        thread.finished.connect(lambda jid=job_id: self._cleanup_ai_thread(jid))
        self._append_task_log(f"{label}启动")
        self._update_task_activity_state()
        thread.start()
        return job_id

    def _create_auto_snapshot(self, label_prefix: str) -> None:
        if not self.state.selected_chapter_id:
            return
        label = f"{label_prefix} · {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        self.database.create_snapshot(
            self.state.selected_chapter_id,
            label,
            self.chapter_outline_editor.toPlainText(),
            self.content_editor.toPlainText(),
        )

    def _open_ai_settings(self) -> None:
        dialog = AiSettingsDialog(
            self,
            profiles=self._ai_profiles(),
            purpose_labels=AI_PURPOSES,
            request_retries=getattr(self, "ai_request_retries", DEFAULT_AI_REQUEST_RETRIES),
            book_analysis_concurrency=getattr(self, "book_analysis_concurrency", DEFAULT_BOOK_ANALYSIS_CONCURRENCY),
            book_analysis_fast_extra_body=getattr(
                self,
                "book_analysis_fast_extra_body",
                DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON,
            ),
            book_analysis_final_extra_body=getattr(
                self,
                "book_analysis_final_extra_body",
                DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON,
            ),
        )
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self._save_ai_settings(dialog.values())

    def _open_prompt_settings(self) -> None:
        values = normalize_prompt_overrides(getattr(self, "prompt_overrides", {}), PROMPT_SETTING_ENTRIES)
        dialog = PromptSettingsDialog(self, entries=PROMPT_SETTING_ENTRIES, values=values)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self._save_prompt_settings(dialog.result_values())

    def _save_prompt_settings(self, values: dict[str, str]) -> None:
        prompt_overrides = normalize_prompt_overrides(values, PROMPT_SETTING_ENTRIES)
        self.prompt_overrides = prompt_overrides
        if isinstance(getattr(self, "ai_settings", None), dict):
            self.ai_settings["prompt_settings_version"] = PROMPT_SETTINGS_VERSION
            self.ai_settings["prompt_overrides"] = prompt_overrides
        self.ai_service.set_prompt_overrides(prompt_overrides)
        profiles = self._ai_profiles()
        try:
            self._write_json_file(
                self.ai_config_path,
                self._build_stored_ai_settings(
                    profiles,
                    getattr(self, "ai_request_retries", DEFAULT_AI_REQUEST_RETRIES),
                    getattr(self, "book_analysis_concurrency", DEFAULT_BOOK_ANALYSIS_CONCURRENCY),
                    getattr(self, "book_analysis_fast_extra_body", DEFAULT_BOOK_ANALYSIS_FAST_EXTRA_BODY_JSON),
                    getattr(self, "book_analysis_final_extra_body", DEFAULT_BOOK_ANALYSIS_FINAL_EXTRA_BODY_JSON),
                    prompt_overrides,
                ),
            )
        except (OSError, ValueError) as exc:
            error(self, "提示词配置保存失败", str(exc))
            return
        self._set_status("Agent 与提示词配置已保存")

    def _prompt_text(self, key: str, default: str = "") -> str:
        return prompt_value(getattr(self, "prompt_overrides", {}), key, default)

    def _open_book_settings(self, book_id: int | None = None) -> None:
        book_id = book_id or self.state.selected_book_id
        dialog = QDialog(self)
        dialog.setWindowTitle("Skills 分层管理")
        dialog.resize(980, 660)
        layout = QVBoxLayout(dialog)
        if book_id:
            book = self.database.get_book(book_id)
            title_row = QHBoxLayout()
            title_field = QLineEdit(str(book["title"]) if book else "")
            save_title = QPushButton("保存书名")
            save_title.clicked.connect(lambda: self._save_book_title_from_dialog(book_id, title_field.text(), dialog))
            title_row.addWidget(QLabel("书名"))
            title_row.addWidget(title_field, 1)
            title_row.addWidget(save_title)
            layout.addLayout(title_row)
        else:
            layout.addWidget(QLabel("当前未选择书籍，只管理全局 Skills。"))

        scope_row = QHBoxLayout()
        scope_combo = QComboBox()
        scope_combo.addItem("全局 Skills", "global:0")
        if book_id:
            scope_combo.addItem("书籍 Skills", f"book:{book_id}")
        if self.state.selected_chapter_id:
            scope_combo.addItem("章节 Skills", f"chapter:{self.state.selected_chapter_id}")
        if self.state.selected_chapter_id:
            scope_combo.setCurrentIndex(scope_combo.findData(f"chapter:{self.state.selected_chapter_id}"))
        elif book_id:
            scope_combo.setCurrentIndex(scope_combo.findData(f"book:{book_id}"))
        scope_row.addWidget(QLabel("启用层级"))
        scope_row.addWidget(scope_combo, 1)
        layout.addLayout(scope_row)

        lists = QSplitter(Qt.Orientation.Horizontal)
        sources_list = QListWidget()
        skills_list = QListWidget()
        enabled_list = QListWidget()
        lists.addWidget(sources_list)
        lists.addWidget(skills_list)
        lists.addWidget(enabled_list)
        layout.addWidget(lists, 1)

        preview = QPlainTextEdit()
        preview.setReadOnly(True)
        preview.setMaximumHeight(150)
        layout.addWidget(preview)

        actions = QHBoxLayout()
        import_ref = QPushButton("导入参考书")
        distill = QPushButton("提炼 Skills")
        bind = QPushButton("启用到当前层")
        unbind = QPushButton("停用")
        delete_skill = QPushButton("删除 Skill")
        delete_skill.setObjectName("DangerButton")
        delete_source = QPushButton("删除参考书")
        delete_source.setObjectName("DangerButton")
        close = QPushButton("关闭")
        actions.addWidget(import_ref)
        actions.addWidget(distill)
        actions.addWidget(bind)
        actions.addWidget(unbind)
        actions.addWidget(delete_skill)
        actions.addWidget(delete_source)
        actions.addStretch(1)
        actions.addWidget(close)
        layout.addLayout(actions)

        state: dict[str, list[dict[str, Any]]] = {"sources": [], "skills": [], "enabled": []}

        def current_scope() -> tuple[str, int]:
            raw = str(scope_combo.currentData() or "global:0")
            scope, _, sid = raw.partition(":")
            return scope, int(sid or 0)

        def load_sources() -> None:
            state["sources"] = [dict(item) for item in self.database.list_reference_sources()]
            sources_list.clear()
            for source in state["sources"]:
                prefix = "GitHub" if source.get("source_type") == "github_project" else "参考书"
                item = QListWidgetItem(f"{prefix} · {source.get('title', '')}")
                item.setData(Qt.ItemDataRole.UserRole, source)
                sources_list.addItem(item)
            load_enabled()

        def load_skills(source_id: int | None = None) -> None:
            state["skills"] = [dict(item) for item in self.database.list_skills(source_id)]
            skills_list.clear()
            for skill in state["skills"]:
                item = QListWidgetItem(f"{skill.get('name', '')} · {skill.get('category', '')}")
                item.setData(Qt.ItemDataRole.UserRole, skill)
                skills_list.addItem(item)

        def load_enabled() -> None:
            scope, scope_id = current_scope()
            if scope == "global":
                state["enabled"] = [dict(item) for item in self.database.list_bound_skills_for_global()]
            elif scope == "book" and scope_id:
                state["enabled"] = [dict(item) for item in self.database.list_bound_skills_for_book(scope_id)]
            elif scope == "chapter" and scope_id:
                state["enabled"] = [dict(item) for item in self.database.list_bound_skills_for_chapter(scope_id)]
            else:
                state["enabled"] = []
            enabled_list.clear()
            for skill in state["enabled"]:
                item = QListWidgetItem(str(skill.get("name", "")))
                item.setData(Qt.ItemDataRole.UserRole, skill)
                enabled_list.addItem(item)

        def on_source_changed(current: QListWidgetItem | None, _previous: QListWidgetItem | None) -> None:
            source = current.data(Qt.ItemDataRole.UserRole) if current else None
            load_skills(int(source["id"]) if source else None)

        def on_skill_changed(current: QListWidgetItem | None, _previous: QListWidgetItem | None) -> None:
            skill = current.data(Qt.ItemDataRole.UserRole) if current else None
            if not skill:
                preview.clear()
                return
            preview.setPlainText(f"{skill.get('summary', '')}\n\n{skill.get('instruction_text', '')}")

        def import_reference() -> None:
            file_path, _ = QFileDialog.getOpenFileName(dialog, "选择参考书文本", "", "Text Files (*.txt *.md);;All Files (*.*)")
            if not file_path:
                return
            original = Path(file_path)
            target = self.database.references_dir / f"{original.stem}_{uuid.uuid4().hex[:8]}{original.suffix}"
            shutil.copy2(original, target)
            self.database.create_reference_source(title=original.stem, source_path=str(target), rights_note="user_imported")
            load_sources()

        def distill_selected() -> None:
            current = sources_list.currentItem()
            if not current:
                info(dialog, "提示", "请先选择学习来源。")
                return
            source = current.data(Qt.ItemDataRole.UserRole)
            source_path = Path(str(source["source_path"]))
            if not source_path.exists():
                error(dialog, "错误", "来源文件不存在。")
                return
            text = source_path.read_text(encoding="utf-8", errors="ignore")
            ai_service = self._require_ai_service("skills", "Skills 提炼")
            if ai_service is None:
                return
            source_id = int(source["id"])

            def work() -> list[dict[str, str]]:
                if str(source.get("source_type") or "book") == "github_project":
                    return ai_service.distill_project_skills(
                        source_title=str(source["title"]),
                        source_text=text,
                        source_url=str(source.get("source_url") or ""),
                        license_note=str(source.get("source_license") or "待确认"),
                        reusable_level=str(source.get("reusable_level") or "pattern_only"),
                        require_remote=True,
                    )
                return ai_service.distill_skills(str(source["title"]), text, require_remote=True)

            def done(skills: Any) -> None:
                skill_list = skills if isinstance(skills, list) else []
                self.database.replace_skills_for_source(source_id, skill_list)
                load_skills(source_id)
                self._switch_drawer_tab("skills")

            self._run_background_ai_job(label="Skills 提炼", callback=work, on_done=done)

        def bind_selected() -> None:
            current = skills_list.currentItem()
            if not current:
                return
            skill = current.data(Qt.ItemDataRole.UserRole)
            scope, scope_id = current_scope()
            if scope == "global":
                self.database.bind_skill_to_global(int(skill["id"]))
            elif scope == "book" and scope_id:
                self.database.bind_skill_to_book(scope_id, int(skill["id"]))
            elif scope == "chapter" and scope_id:
                self.database.bind_skill_to_chapter(scope_id, int(skill["id"]))
            load_enabled()
            self._refresh_drawer()

        def unbind_selected() -> None:
            current = enabled_list.currentItem()
            if not current:
                return
            skill = current.data(Qt.ItemDataRole.UserRole)
            scope, scope_id = current_scope()
            if scope == "global":
                self.database.unbind_skill_from_global(int(skill["id"]))
            elif scope == "book" and scope_id:
                self.database.unbind_skill_from_book(scope_id, int(skill["id"]))
            elif scope == "chapter" and scope_id:
                self.database.unbind_skill_from_chapter(scope_id, int(skill["id"]))
            load_enabled()
            self._refresh_drawer()

        def delete_selected_skill() -> None:
            current = skills_list.currentItem()
            if not current:
                info(dialog, "提示", "请先选择要删除的 Skill。")
                return
            skill = current.data(Qt.ItemDataRole.UserRole)
            skill_name = str(skill.get("name", "") or "未命名 Skill")
            if not confirm(dialog, "删除 Skill", f"确定删除「{skill_name}」吗？所有层级绑定也会一并删除。", danger=True):
                return
            source_id = int(skill.get("source_id", 0) or 0)
            self.database.delete_skill(int(skill["id"]))
            load_skills(source_id or None)
            load_enabled()
            preview.clear()
            self._refresh_drawer()

        def delete_selected_source() -> None:
            current = sources_list.currentItem()
            if not current:
                info(dialog, "提示", "请先选择要删除的参考书或调研来源。")
                return
            source = current.data(Qt.ItemDataRole.UserRole)
            title = str(source.get("title", "") or "未命名来源")
            if not confirm(dialog, "删除参考来源", f"确定删除「{title}」吗？其提炼出的 Skills 和绑定也会一并删除。", danger=True):
                return
            source_path = Path(str(source.get("source_path", "") or ""))
            self.database.delete_reference_source(int(source["id"]))
            try:
                if source_path.exists() and source_path.is_file():
                    data_root = self.database.data_dir.resolve()
                    if data_root in source_path.resolve().parents:
                        source_path.unlink()
            except OSError:
                pass
            load_sources()
            load_skills(None)
            preview.clear()
            self._refresh_drawer()

        sources_list.currentItemChanged.connect(on_source_changed)
        skills_list.currentItemChanged.connect(on_skill_changed)
        import_ref.clicked.connect(import_reference)
        distill.clicked.connect(distill_selected)
        bind.clicked.connect(bind_selected)
        unbind.clicked.connect(unbind_selected)
        delete_skill.clicked.connect(delete_selected_skill)
        delete_source.clicked.connect(delete_selected_source)
        scope_combo.currentIndexChanged.connect(lambda _index: load_enabled())
        close.clicked.connect(dialog.accept)
        load_sources()
        load_skills(None)
        dialog.exec()

    def _save_book_title_from_dialog(self, book_id: int, title: str, dialog: QDialog) -> None:
        title = title.strip()
        if not title:
            return
        self.database.rename_book(book_id, title)
        self._load_library_tree(("book", book_id, None))
        self._activate_tree_key(("book", book_id, None))
        self._refresh_drawer()
        info(dialog, "完成", "书名已保存。")

    def _open_history_dialog(self, chapter_id: int | None = None) -> None:
        chapter_id = chapter_id or self.state.selected_chapter_id
        if not chapter_id:
            info(self, "提示", "请先选择一个章节。")
            return
        dialog = QDialog(self)
        dialog.setWindowTitle("章节历史")
        dialog.resize(760, 560)
        layout = QVBoxLayout(dialog)
        snapshots_list = QListWidget()
        preview = QPlainTextEdit()
        preview.setReadOnly(True)
        layout.addWidget(snapshots_list)
        layout.addWidget(preview, 1)
        buttons = QHBoxLayout()
        restore = QPushButton("恢复快照")
        close = QPushButton("关闭")
        buttons.addWidget(restore)
        buttons.addStretch(1)
        buttons.addWidget(close)
        layout.addLayout(buttons)
        snapshots = [dict(item) for item in self.database.list_snapshots(chapter_id)]
        for snap in snapshots:
            item = QListWidgetItem(f"{snap.get('label', '')} · {snap.get('created_at', '')}")
            item.setData(Qt.ItemDataRole.UserRole, snap)
            snapshots_list.addItem(item)

        def on_snapshot(current: QListWidgetItem | None, _previous: QListWidgetItem | None) -> None:
            if not current:
                preview.clear()
                return
            snap = current.data(Qt.ItemDataRole.UserRole)
            detail = self.database.get_snapshot(int(snap["id"]))
            if detail:
                preview.setPlainText(f"大纲：\n{detail['outline']}\n\n正文：\n{detail['content']}")

        def restore_snapshot() -> None:
            current = snapshots_list.currentItem()
            if not current:
                return
            snap = current.data(Qt.ItemDataRole.UserRole)
            detail = self.database.get_snapshot(int(snap["id"]))
            if not detail:
                return
            if not confirm(dialog, "恢复快照", "确定将当前编辑区恢复到该快照吗？", danger=True):
                return
            self.chapter_outline_editor.setPlainText(str(detail["outline"]))
            self.content_editor.setPlainText(str(detail["content"]))
            self.state.editor_dirty = True
            self.save_label.setText("未保存")
            dialog.accept()

        snapshots_list.currentItemChanged.connect(on_snapshot)
        restore.clicked.connect(restore_snapshot)
        close.clicked.connect(dialog.accept)
        dialog.exec()

    def _export_current(self, export_type: str) -> None:
        if self.state.current_node_kind == "chapter" and self.state.selected_chapter_id:
            self._export_chapter(export_type, self.state.selected_chapter_id)
            return
        if self.state.selected_book_id:
            self._export_book(export_type, self.state.selected_book_id)
            return
        info(self, "提示", "请先选择一本书或一个章节。")

    def _export_book(self, export_type: str, book_id: int | None = None) -> None:
        book_id = book_id or self.state.selected_book_id
        if not book_id:
            info(self, "提示", "请先选择一本书。")
            return
        if self.state.editor_dirty and self.state.selected_book_id == book_id and not self._save_current():
            return
        book = self.database.get_book(book_id)
        suffix, filter_text = self._export_format_meta(export_type)
        path, _ = QFileDialog.getSaveFileName(
            self,
            "导出作品",
            f"{self._safe_filename(str(book['title']) if book else 'novel')}.{suffix}",
            filter_text,
        )
        if not path:
            return
        output = self._ensure_output_suffix(Path(path), suffix)
        try:
            if export_type == "txt":
                self.exporter.export_txt(book_id, output)
            elif export_type == "docx":
                self.exporter.export_docx(book_id, output)
            else:
                self.exporter.export_markdown(book_id, output)
        except Exception as exc:  # noqa: BLE001
            error(self, "导出失败", str(exc))
            return
        self._set_status(f"已导出书籍：{output}")

    def _export_book_part(self, part: str, book_id: int | None = None) -> None:
        book_id = book_id or self.state.selected_book_id
        if not book_id:
            info(self, "提示", "请先选择一本书。")
            return
        if self.state.editor_dirty and self.state.selected_book_id == book_id and not self._save_current():
            return
        book = self.database.get_book(book_id)
        title = self._safe_filename(str(book["title"]) if book else "novel")
        path, _ = QFileDialog.getSaveFileName(
            self,
            "导出作品",
            f"{title}_{part}.txt",
            "TXT 文件 (*.txt)",
        )
        if not path:
            return
        output = self._ensure_output_suffix(Path(path), "txt")
        try:
            if part == "content":
                self.exporter.export_content_only_txt(book_id, output)
            elif part == "outline":
                self.exporter.export_outline_only_txt(book_id, output)
            elif part == "characters":
                self.exporter.export_characters_only_txt(book_id, output)
            elif part == "world":
                self.exporter.export_world_only_txt(book_id, output)
            elif part == "per_volume":
                paths = self.exporter.export_per_volume_txt(book_id, output)
                self._set_status(f"已按卷导出 {len(paths)} 个文件：{output.parent}")
                return
            else:
                return
        except Exception as exc:
            error(self, "导出失败", str(exc))
            return
        self._set_status(f"已导出：{output}")

    def _export_chapter(self, export_type: str, chapter_id: int | None = None) -> None:
        chapter_id = chapter_id or self.state.selected_chapter_id
        if not chapter_id:
            info(self, "提示", "请先选择一个章节。")
            return
        if self.state.editor_dirty and self.state.selected_chapter_id == chapter_id and not self._save_current():
            return
        chapter = self.database.get_chapter(chapter_id)
        if not chapter:
            return
        suffix, filter_text = self._export_format_meta(export_type)
        path, _ = QFileDialog.getSaveFileName(
            self,
            "导出章节",
            f"{self._safe_filename(str(chapter['title']))}.{suffix}",
            filter_text,
        )
        if not path:
            return
        output = self._ensure_output_suffix(Path(path), suffix)
        try:
            if export_type == "txt":
                self.exporter.export_chapter_txt(chapter_id, output)
            elif export_type == "docx":
                self.exporter.export_chapter_docx(chapter_id, output)
            else:
                self.exporter.export_chapter_markdown(chapter_id, output)
        except Exception as exc:  # noqa: BLE001
            error(self, "导出失败", str(exc))
            return
        self._set_status(f"已导出章节：{output}")

    def _export_format_meta(self, export_type: str) -> tuple[str, str]:
        if export_type == "txt":
            return "txt", "Text Files (*.txt)"
        if export_type == "docx":
            return "docx", "Word Documents (*.docx)"
        return "md", "Markdown Files (*.md)"

    def _ensure_output_suffix(self, path: Path, suffix: str) -> Path:
        expected = f".{suffix.lower()}"
        if path.suffix.lower() == expected:
            return path
        return path.with_suffix(expected)

    def _safe_filename(self, value: str) -> str:
        cleaned = "".join("_" if char in '<>:"/\\|?*' else char for char in value).strip()
        return cleaned or "novel"

    def resizeEvent(self, event) -> None:  # noqa: N802
        super().resizeEvent(event)
        width = self.width()
        self._apply_responsive_chrome()
        if hasattr(self, "background_video_label") and self.background_video_label:
            self.background_video_label.lower()
        if hasattr(self, "background_video_view") and self.background_video_view:
            self._resize_background_video_item()
            self.background_video_view.lower()
        if hasattr(self, "root_widget") and self.root_widget:
            self.root_widget.raise_()
        if not hasattr(self, "main_splitter"):
            return
        if width < 1100 and not self.library_collapsed:
            self.library_collapsed = True
            self.focus_mode = False
            self._apply_library_collapsed_state()
        elif width > 1280 and not self.library_collapsed and self.main_splitter.sizes()[0] < 140:
            self.main_splitter.setSizes([self.library_width, max(900, self.main_splitter.sizes()[1])])

    def closeEvent(self, event: QCloseEvent) -> None:  # noqa: N802
        if self._running_ai_threads():
            if not self.close_after_task_stop and not confirm(self, "任务运行中", "AI 任务仍在运行或退出中，确定退出吗？", danger=True):
                event.ignore()
                return
            self.close_after_task_stop = True
            self._request_all_ai_threads_stop()
            if not self.cancel_grace_timer.isActive():
                self.cancel_grace_timer.start(5000)
            if not self.close_poll_timer.isActive():
                self.close_poll_timer.start()
            self._set_status("正在停止后台 AI 任务，退出后将自动关闭。")
            event.ignore()
            return
        if not self._confirm_editor_navigation():
            event.ignore()
            return
        if self.close_poll_timer.isActive():
            self.close_poll_timer.stop()
        self._save_ui_settings()
        self.database.close()
        event.accept()


def _resolve_and_set_app_icon(app: QApplication) -> None:
    icon_paths = [
        BASE_DIR / "packaging" / "app_icon.ico",
        Path(getattr(sys, "_MEIPASS", "")) / "packaging" / "app_icon.ico",
    ]
    for candidate in icon_paths:
        try:
            if candidate.exists():
                app.setWindowIcon(QIcon(str(candidate)))
                return
        except Exception:
            continue


def _install_exception_popup(app: QApplication) -> None:
    if getattr(app, "_simple_ai_exception_popup_installed", False):
        return
    setattr(app, "_simple_ai_exception_popup_installed", True)
    previous_hook = sys.excepthook

    def show_exception(exc_type, exc_value, exc_tb) -> None:  # noqa: ANN001
        details = "".join(traceback.format_exception(exc_type, exc_value, exc_tb)).strip()
        try:
            box = QMessageBox()
            box.setIcon(QMessageBox.Icon.Critical)
            box.setWindowTitle("程序异常")
            box.setText("软件遇到未处理异常，错误详情如下。")
            box.setDetailedText(details[-8000:])
            box.exec()
        except Exception:  # noqa: BLE001
            pass
        if previous_hook and previous_hook is not show_exception:
            previous_hook(exc_type, exc_value, exc_tb)

    sys.excepthook = show_exception


def run_qt_app(
    database: Database,
    ai_service: SimpleAIService,
    startup_splash: Any | None = None,
) -> int:
    app = QApplication.instance() or QApplication([])
    app.setApplicationName(APP_DISPLAY_NAME)
    _resolve_and_set_app_icon(app)
    _install_exception_popup(app)
    if startup_splash is not None:
        startup_splash.set_stage("正在构建工作台", "加载书库、编辑器与上下文系统")
    window = NovelQtMainWindow(database, ai_service)
    app_icon = app.windowIcon()
    if not app_icon.isNull():
        window.setWindowIcon(app_icon)
    if os.environ.get("SIMPLE_AI_NOVEL_QT_OFFSCREEN") != "1":
        if startup_splash is not None:
            startup_splash.set_stage("正在显示主窗口", "马上就绪")
        window.show()
        if startup_splash is not None:
            startup_splash.finish(window)
    elif startup_splash is not None:
        startup_splash.close()
    if os.environ.get("SIMPLE_AI_NOVEL_QT_SMOKE") == "1":
        QTimer.singleShot(50, app.quit)
    exit_code = int(app.exec())
    try:
        database.close()
    except Exception:  # noqa: BLE001
        pass
    return exit_code
